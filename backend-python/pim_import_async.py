"""
Chunked async PIM Product Master import: preview/apply in steps to avoid HTTP timeouts.

State: raw_report_uploads.metadata.pim_import_job (lifecycle, cursors, frozen_plan).
"""

from __future__ import annotations

import importlib
import io
import json
import logging
import os
import tempfile
import time
from datetime import datetime, timezone
from typing import Any

import pandas as pd

log = logging.getLogger(__name__)

from pim_product_master import PRODUCT_MASTER_SHEET_NAME
from pim_seed_cleaning import parse_identifier_row


def _pim_import_meta_cancelled(
    meta: dict[str, Any],
    db: Any | None = None,
    organization_id: str | None = None,
    upload_id: str | None = None,
) -> bool:
    """True if the import was cancelled. Checks file_processing_status.cancel_requested_at first,
    then falls back to raw_report_uploads.metadata flags (set by cancel API route)."""
    if db is not None and upload_id:
        try:
            r = (
                db.table("file_processing_status")
                .select("cancel_requested_at")
                .eq("upload_id", upload_id)
                .limit(1)
                .execute()
            )
            row = (r.data or [None])[0]
            if row and row.get("cancel_requested_at"):
                return True
        except Exception:
            log.debug("file_processing_status cancel check skipped", exc_info=True)
    # Metadata fallback — set by cancel API route directly on raw_report_uploads
    if bool(meta.get("pim_import_cancelled")):
        return True
    if str(meta.get("preview_status") or "").strip().lower() == "cancelled":
        return True
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else {}
    return str(job.get("lifecycle") or "").strip().lower() == "cancelled"


def _fps_cancel_merge(meta: dict[str, Any]) -> dict[str, Any]:
    """Return updated metadata dict with cancel flags set (caller must call _persist_meta)."""
    meta2 = dict(meta)
    meta2["pim_import_cancelled"] = True
    meta2["preview_status"] = "cancelled"
    meta2["import_job_status"] = "cancelled"
    job = meta2.get("pim_import_job") if isinstance(meta2.get("pim_import_job"), dict) else {}
    meta2["pim_import_job"] = {
        **job,
        "lifecycle": "cancelled",
        "stage_label": "cancelled",
        "last_error": "cancelled_by_user",
        "preview_phase": None,
        "last_step_at": datetime.now(timezone.utc).isoformat(),
    }
    return meta2


_IGNORE_TAB_SUBSTR = ("readme", "summary", "audit", "instructions")
_REFERENCE_IDENTIFIER_TABS = frozenset(
    {"Identifier_Map", "Identifier_Map_Merged", "Identifier_Map_Audit"}
)
_PRICE_HISTORY_TABS = frozenset({"Price_History", "Price_History_Candidates"})
_VENDOR_REF_TABS = frozenset({"Vendors", "Vendors_Merged"})
_BRAND_MAP_TABS = frozenset({"Brand_By_MFG"})

ROW_CHUNK_DEFAULT = 300
ROW_CHUNK_XLSX_DEFAULT = 300
ROW_CHUNK_CSV_DEFAULT = 300
ROW_CHUNK_MIN = 250
ROW_CHUNK_MAX = 500
# Minimum chunk size allowed after a connection-termination error (overrides ROW_CHUNK_MIN)
ROW_CHUNK_RETRY_MIN = 50
PREVIEW_STEP_TIME_BUDGET_SEC = 42.0
# Apply must process larger chunks per HTTP round-trip; no per-row Amazon SP calls (see skip_amazon_enrichment).
APPLY_STEP_TIME_BUDGET_SEC = 95.0


def _safe_int(v: Any, default: int = 0) -> int:
    try:
        return int(v)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return default


def _count_excel_data_rows(path: str, sheet_name: str, max_scan: int = 2_000_000) -> int:
    """Count non-header rows when ws.max_row is unreliable in read-only mode."""
    import openpyxl

    n = 0
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for _ in ws.iter_rows(min_row=2, max_row=max_scan + 1, values_only=True):
            n += 1
            if n >= max_scan:
                break
    finally:
        wb.close()
    return n


def _pim_public_classifier_signals(raw: Any) -> dict[str, Any]:
    """Small JSON-safe snapshot for API payloads."""
    if not isinstance(raw, dict):
        return {}
    out: dict[str, Any] = {}
    for k, v in list(raw.items())[:40]:
        if isinstance(v, (str, int, float, bool)) or v is None:
            out[str(k)] = v
        elif isinstance(v, list) and len(v) <= 32:
            out[str(k)] = v
    return out


def build_public_preview_payload(
    *,
    upload_id: str,
    quality: dict[str, Any],
    frozen: dict[str, Any],
    job: dict[str, Any],
    seed_session_id: str | None,
) -> dict[str, Any]:
    """Stable JSON shape for API + metadata.pim_preview_result (all counts numeric, no undefined)."""
    mapping = frozen.get("pim_column_map") if isinstance(frozen.get("pim_column_map"), dict) else {}
    rows_scanned = _safe_int(quality.get("rows_total"))
    dirty = _safe_int(quality.get("dirty_rows")) or _safe_int(quality.get("skipped_dirty_row"))
    rows_accepted = _safe_int(quality.get("rows_accepted_estimate"))
    if rows_accepted == 0 and rows_scanned > 0:
        rows_accepted = max(rows_scanned - dirty, 0)
    rejected_sample = quality.get("rejected_sample") if isinstance(quality.get("rejected_sample"), list) else []
    accepted_sample = quality.get("accepted_sample") if isinstance(quality.get("accepted_sample"), list) else []
    pe = quality.get("preview_errors") if isinstance(quality.get("preview_errors"), list) else []
    err_list: list[str] = []
    for item in pe[:200]:
        if isinstance(item, str):
            err_list.append(item)
    if not err_list and isinstance(rejected_sample, list):
        for r in rejected_sample[:200]:
            if isinstance(r, dict) and isinstance(r.get("message"), str):
                err_list.append(r["message"])
    rq = dict(quality)
    pe0 = rq.get("preview_errors")
    if isinstance(pe0, list):
        rq["preview_errors"] = [str(x) for x in pe0 if isinstance(x, str)][:200]
    for k in ("accepted_sample", "rejected_sample"):
        if k in rq and isinstance(rq[k], list):
            rq[k] = rq[k][:100]
    master = job.get("selected_master_sheet")
    selected = [str(master)] if master else []
    ignored_raw = job.get("ignored_sheets")
    ignored: list[Any] = list(ignored_raw) if isinstance(ignored_raw, list) else []
    handles_f = frozen.get("handles") if isinstance(frozen.get("handles"), dict) else {}
    mapped_cat_col = handles_f.get("category")
    raw_cat_dbg = quality.get("category_import_debug")
    cat_dbg_out: dict[str, Any] | None = None
    if isinstance(raw_cat_dbg, dict):
        cat_dbg_out = {**raw_cat_dbg, "mapped_file_column": str(mapped_cat_col) if mapped_cat_col else None}
    elif mapped_cat_col:
        cat_dbg_out = {"mapped_file_column": str(mapped_cat_col), "samples": []}
    q_out = {
        "rows_scanned": rows_scanned,
        "rows_accepted": rows_accepted,
        "rows_rejected": dirty,
        "ambiguous_rows": _safe_int(quality.get("ambiguous_rows")) or _safe_int(quality.get("skipped_ambiguous")),
        "dirty_rate": float(quality.get("dirty_rate") or 0.0),
        "products_would_create": _safe_int(quality.get("products_would_create")),
        "products_would_update": _safe_int(quality.get("products_would_update")),
        "vendors_would_create": _safe_int(quality.get("vendors_would_create")),
        "vendors_reused": _safe_int(quality.get("vendors_reused")),
        "categories_would_create": _safe_int(quality.get("categories_would_create")),
        "categories_reused": _safe_int(quality.get("categories_reused")),
        "identifier_rows_would_insert": _safe_int(quality.get("identifier_rows_would_insert")),
        "identifier_rows_would_update": _safe_int(quality.get("identifier_rows_would_update")),
        "prices_would_insert": _safe_int(quality.get("prices_would_insert")),
        "metadata_attributes_detected": _safe_int(quality.get("metadata_attributes_detected")),
        "duplicate_rows_collapsed": _safe_int(quality.get("duplicate_identifier_tokens_collapsed")),
        "duplicate_identifiers_reused": _safe_int(quality.get("duplicates_reused")),
        "apply_blocked_by_dirty_rate": bool(quality.get("apply_blocked_by_dirty_rate")),
        "apply_blocked_by_conflicts": bool(quality.get("apply_blocked_by_conflicts")),
        "max_dirty_rate": float(quality.get("max_dirty_rate") or 0.0),
        "multi_identifier_rows_allowed": _safe_int(quality.get("multi_identifier_rows_allowed")),
        "multi_identifier_rows_conflicting": _safe_int(quality.get("multi_identifier_rows_conflicting")),
        "canonical_products_from_multi_id_rows": _safe_int(quality.get("canonical_products_from_multi_id_rows")),
        "identifier_tokens_attached": _safe_int(quality.get("identifier_tokens_attached")),
        "conflict_rows_blocked": _safe_int(quality.get("conflict_rows_blocked")),
        "category_import_debug": cat_dbg_out,
        "import_mode": str(frozen.get("import_mode") or job.get("import_mode") or ""),
        "classifier_signals": _pim_public_classifier_signals(
            frozen.get("classifier_signals") if isinstance(frozen.get("classifier_signals"), dict) else job.get("classifier_signals")
        ),
    }
    return {
        "status": "preview_ready",
        "import_job_id": str(upload_id),
        "seed_session_id": seed_session_id,
        "quality": q_out,
        "preview_metrics": q_out,
        "mapping": dict(mapping),
        "category_import_debug": cat_dbg_out,
        "mapping_source": str(frozen.get("mapping_source") or "unknown"),
        "selected_sheets": selected,
        "ignored_sheets": ignored,
        "sheet_classifications": job.get("sheet_classifications") if isinstance(job.get("sheet_classifications"), list) else [],
        "samples": {
            "accepted": list(accepted_sample)[:100],
            "rejected": list(rejected_sample)[:100],
        },
        "errors": err_list[:200],
        "raw_quality": rq,
    }


def _main():
    return importlib.import_module("main")


def _strip_heavy_preview_ready_response(d: dict[str, Any]) -> dict[str, Any]:
    """Keep HTTP + proxy payloads small (samples/errors caps)."""
    out = dict(d)
    pq = out.get("preview_quality")
    if isinstance(pq, dict):
        nq = dict(pq)
        if isinstance(nq.get("preview_errors"), list):
            nq["preview_errors"] = [str(x) for x in nq["preview_errors"] if isinstance(x, str)][:200]
        for k in ("accepted_sample", "rejected_sample"):
            if isinstance(nq.get(k), list):
                nq[k] = nq[k][:100]
        out["preview_quality"] = nq
    samp = out.get("samples")
    if isinstance(samp, dict):
        out["samples"] = {
            "accepted": list(samp.get("accepted") or [])[:100],
            "rejected": list(samp.get("rejected") or [])[:100],
        }
    if isinstance(out.get("errors"), list):
        out["errors"] = out["errors"][:200]
    return out


def classify_sheet_name(name: str) -> tuple[str, float, str]:
    k = str(name or "").strip()
    kl = k.lower().replace(" ", "_")
    if any(x in kl for x in _IGNORE_TAB_SUBSTR):
        return "ignored", 0.95, "readme_or_summary_pattern"
    if k == PRODUCT_MASTER_SHEET_NAME:
        return "product_master", 1.0, "exact_App_Import_Product_Master"
    if k in _REFERENCE_IDENTIFIER_TABS:
        return "identifier_reference", 0.95, "known_identifier_tab"
    if k in _PRICE_HISTORY_TABS:
        return "price_history", 0.9, "known_price_tab"
    if k in _VENDOR_REF_TABS:
        return "vendor_reference", 0.85, "known_vendor_tab"
    if k in _BRAND_MAP_TABS:
        return "brand_map", 0.85, "known_brand_tab"
    return "unknown", 0.0, "no_known_name"


def classify_sheet_by_headers(headers: list[str]) -> tuple[str, float, str]:
    m = _main()
    raw_headers = [str(h).strip() for h in headers if str(h).strip()]
    det = m._pim_try_deterministic_column_map(raw_headers)
    if det:
        id_cols = bool(det.get("seller_sku") or det.get("asin") or det.get("fnsku") or det.get("upc"))
        name_col = bool(det.get("product_name"))
        if id_cols and name_col:
            return "product_master", 0.88, "headers_match_product_master_alias"
        if id_cols and not name_col:
            return "identifier_reference", 0.55, "identifiers_without_product_name"
    hset = {m._normalize_header(str(h)) for h in raw_headers}
    price_hints = {"cost", "unit_cost", "price", "last_cost", "amount"} & hset
    id_hints = {"asin", "sku", "seller_sku", "fnsku", "upc"} & hset
    if price_hints and len(id_hints) >= 1:
        return "price_history", 0.45, "price_like_headers"
    return "unknown", 0.15, "no_confident_match"


def _upload_pim_scan_cache_pickle(db: Any, storage_prefix: str, df: Any) -> str:
    buf = io.BytesIO()
    df.to_pickle(buf, protocol=4)
    data = buf.getvalue()
    dest = f"{storage_prefix.rstrip('/')}/pim_scan_cache.pkl"
    db.storage.from_("raw-reports").upload(
        path=dest,
        file=data,
        file_options={"content-type": "application/octet-stream", "upsert": "true"},
    )
    return dest


def _load_csv_df_from_cache_or_parse(
    db: Any,
    meta: dict[str, Any],
    local_path: str,
    fname: str,
    *,
    organization_id: str,
    upload_id: str,
) -> tuple[Any, dict[str, Any]]:
    """Load CSV as DataFrame from storage pickle cache if present; else parse once and upload cache."""
    prefix = str(meta.get("storage_prefix") or "").strip()
    cache_path = str(meta.get("pim_csv_scan_cache_storage") or "").strip()
    m = _main()
    if cache_path:
        blob = _download_single_storage_path(db, cache_path)
        return pd.read_pickle(io.BytesIO(blob)), meta
    with open(local_path, "rb") as bf:
        raw = bf.read()
    df, _, _ = m._read_text_delimited(raw, fname)
    if prefix:
        dest = _upload_pim_scan_cache_pickle(db, prefix, df)
        meta2 = dict(meta)
        meta2["pim_csv_scan_cache_storage"] = dest
        meta2 = merge_pim_job(
            meta2,
            {
                "stage_label": "csv_cached_for_scan",
                "last_step_at": datetime.now(timezone.utc).isoformat(),
            },
        )
        _persist_meta(db, organization_id, upload_id, meta2, None)
        return df, meta2
    return df, meta


def _ensure_pim_merged_source(
    db: Any,
    organization_id: str,
    upload_id: str,
    meta: dict[str, Any],
    fname: str,
) -> dict[str, Any]:
    """Multi-part uploads: merge chunk parts once into a single storage object so later steps do not re-merge."""
    if str(meta.get("pim_merged_storage_path") or "").strip():
        return meta
    prefix = str(meta.get("storage_prefix") or "").strip()
    parts = int(meta.get("total_parts") or meta.get("upload_chunks_count") or 0)
    if not prefix or parts <= 1:
        return meta
    fl = fname.lower()
    if fl.endswith((".xlsx", ".xlsm")):
        suffix, ctype = ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fl.endswith(".csv"):
        suffix, ctype = ".csv", "text/csv; charset=utf-8"
    elif fl.endswith(".txt"):
        suffix, ctype = ".txt", "text/plain; charset=utf-8"
    else:
        suffix, ctype = ".bin", "application/octet-stream"
    tmp = _merge_upload_parts_to_temp(db, prefix, parts, suffix)
    try:
        dest = f"{prefix.rstrip('/')}/pim_merged_source{suffix}"
        bucket = db.storage.from_("raw-reports")
        with open(tmp, "rb") as bf:
            data = bf.read()
        bucket.upload(
            path=dest,
            file=data,
            file_options={"content-type": ctype, "upsert": "true"},
        )
        meta2 = dict(meta)
        meta2["pim_merged_storage_path"] = dest
        meta2["import_job_status"] = "uploaded"
        meta2 = merge_pim_job(
            meta2,
            {"stage_label": "merged_source_ready", "last_merge_parts_at": time.time()},
        )
        _persist_meta(db, organization_id, upload_id, meta2, None)
        return meta2
    finally:
        try:
            os.unlink(tmp)
        except OSError:
            pass


def _merge_upload_parts_to_temp(db: Any, storage_prefix: str, total_parts: int, suffix: str) -> str:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    tmp_path = tmp.name
    tmp.close()
    bucket = db.storage.from_("raw-reports")
    try:
        with open(tmp_path, "wb") as out:
            for i in range(total_parts):
                path = f"{storage_prefix.rstrip('/')}/part-{i:06d}"
                blob = bucket.download(path)
                if blob:
                    out.write(blob if isinstance(blob, (bytes, bytearray)) else bytes(blob))
        return tmp_path
    except Exception:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        raise


def _download_single_storage_path(db: Any, path: str) -> bytes:
    blob = db.storage.from_("raw-reports").download(path)
    return blob if isinstance(blob, (bytes, bytearray)) else bytes(blob)


def resolve_job_file_path(meta: dict[str, Any]) -> tuple[str, str]:
    merged = (
        (meta.get("pim_merged_storage_path") or meta.get("raw_file_path") or meta.get("storage_path") or "").strip()
    )
    if merged:
        return "single", merged
    prefix = (meta.get("storage_prefix") or "").strip()
    parts = int(meta.get("total_parts") or meta.get("upload_chunks_count") or 0)
    if prefix and parts > 1:
        return "parts", prefix
    if prefix:
        return "single", f"{prefix}/part-000000"
    raise ValueError("No storage path in upload metadata (storage_prefix / merged path)")


def materialize_job_local_file(db: Any, meta: dict[str, Any], filename: str | None) -> tuple[str, bool]:
    fname = (filename or "upload.bin").lower()
    if fname.endswith((".xlsx", ".xlsm")):
        suffix = ".xlsx"
    elif fname.endswith(".csv"):
        suffix = ".csv"
    elif fname.endswith(".txt"):
        suffix = ".txt"
    else:
        suffix = ".bin"

    kind, path = resolve_job_file_path(meta)
    if kind == "single":
        data = _download_single_storage_path(db, path)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
        tmp.write(data)
        tmp.close()
        return tmp.name, True
    parts = int(meta.get("total_parts") or 1)
    return _merge_upload_parts_to_temp(db, path, max(parts, 2), suffix), True


def default_pim_job() -> dict[str, Any]:
    return {
        "lifecycle": "uploaded",
        "progress_pct": 0,
        "stage_label": "uploaded",
        "source_type": "file",
        "preview_phase": None,
        "scan_cursor": None,
        "frozen_plan": None,
        "import_mode": None,
        "import_mode_signals": None,
        "classifier_signals": None,
        "quality_acc": None,
        "preview_quality": None,
        "apply_metrics": None,
        "apply_cursor": None,
        "detected_sheets": [],
        "ignored_sheets": [],
        "selected_master_sheet": None,
        "sheet_classifications": [],
        "last_error": None,
        "elapsed_preview_sec": 0.0,
    }


def merge_pim_job(meta: dict[str, Any], patch: dict[str, Any]) -> dict[str, Any]:
    base = meta.get("pim_import_job")
    if not isinstance(base, dict):
        base = default_pim_job()
    merged_j = {**base, **patch}
    out = dict(meta)
    out["pim_import_job"] = merged_j
    return out


def run_pim_import_retry_preview(
    db: Any,
    organization_id: str,
    store_id: str,
    upload_id: str,
) -> dict[str, Any]:
    """Reuse same raw_report_uploads row: reset scan cursor and quality acc for a new preview pass."""
    _ = store_id
    m = _main()
    res = (
        db.table("raw_report_uploads")
        .select("id,metadata,status,created_by,file_name")
        .eq("id", upload_id)
        .eq("organization_id", organization_id)
        .limit(1)
        .execute()
    )
    rows = res.data or []
    if not rows:
        return {"ok": False, "error": "upload_not_found", "done": True, "terminal": True}
    row0 = rows[0]
    meta = row0.get("metadata") if isinstance(row0.get("metadata"), dict) else {}
    rname = str(row0.get("file_name") or "upload")
    rcreated = str(row0.get("created_by") or "").strip() or None
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()
    frozen = job.get("frozen_plan")
    if not isinstance(frozen, dict):
        return {"ok": False, "error": "missing_frozen_plan", "done": True, "terminal": True}
    sheet = str(frozen.get("sheet_label") or "")
    c0 = job.get("scan_cursor") if isinstance(job.get("scan_cursor"), dict) else {}
    tr = int(c0.get("total_rows") or meta.get("row_count") or 0)
    acc = m._pim_seed_quality_acc()
    frozen2 = dict(frozen)
    patch = {
        "lifecycle": "previewing",
        "preview_phase": "scan",
        "stage_label": "scanning_rows",
        "quality_acc": acc,
        "scan_cursor": {"sheet": sheet, "data_row": 0, "total_rows": tr},
        "preview_quality": None,
        "progress_pct": 25,
        "last_error": None,
        "frozen_plan": frozen2,
        "import_total_rows": 0,
        "apply_metrics": None,
        "apply_cursor": None,
        "apply_effective_row_chunk": None,
        "effective_row_chunk": None,
    }
    meta2 = merge_pim_job(meta, patch)
    meta2.pop("pim_preview_result", None)
    meta2["preview_status"] = "previewing"
    meta2["import_job_status"] = "processing"
    meta2["pim_import_confirmed"] = False
    meta2["pim_import_cancelled"] = False
    _persist_meta(db, organization_id, upload_id, meta2, "processing")
    _fps_update(
        db,
        upload_id,
        organization_id,
        process_pct=25,
        processed_rows=0,
        total_rows=tr or None,
        phase_label="scanning_rows",
    )
    return {
        "ok": True,
        "done": False,
        "import_job_id": str(upload_id),
        "lifecycle": "previewing",
        "stage_label": "scanning_rows",
    }


def _persist_meta(db: Any, organization_id: str, upload_id: str, meta: dict[str, Any], status: str | None) -> None:
    upd: dict[str, Any] = {"metadata": meta, "updated_at": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat()}
    if status:
        upd["status"] = status
    db.table("raw_report_uploads").update(upd).eq("id", upload_id).eq("organization_id", organization_id).execute()


def _fps_update(
    db: Any,
    upload_id: str,
    organization_id: str,
    *,
    process_pct: int,
    processed_rows: int,
    total_rows: int | None,
    phase_label: str,
) -> None:
    try:
        db.table("file_processing_status").upsert(
            {
                "upload_id": upload_id,
                "organization_id": organization_id,
                "status": "processing",
                "process_pct": min(100, max(0, process_pct)),
                "processed_rows": processed_rows,
                "total_rows": total_rows,
                "current_phase": phase_label[:80],
            },
            on_conflict="upload_id",
        ).execute()
    except Exception as e:
        log.warning("file_processing_status update skipped: %s", e)


def _list_xlsx_sheet_names(path: str) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _pick_master_sheet(xlsx_path: str, sheet_names: list[str]) -> tuple[str | None, list[dict[str, Any]], list[str]]:
    classifications: list[dict[str, Any]] = []
    ignored: list[str] = []
    master: str | None = None
    for sn in sheet_names:
        role, conf, reason = classify_sheet_name(sn)
        classifications.append({"name": sn, "role": role, "confidence": conf, "reason": reason})
        if role == "ignored":
            ignored.append(sn)
        if role == "product_master" and master is None:
            master = sn
    if master:
        return master, classifications, ignored
    # fallback: first non-ignored sheet that looks like product master by headers
    import openpyxl

    for sn in sheet_names:
        if sn in ignored:
            continue
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        try:
            ws = wb[sn]
            rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
            header_row = next(rows, None)
            if not header_row:
                continue
            headers = [str(c).strip() if c is not None else "" for c in header_row]
            hr, hconf, _ = classify_sheet_by_headers(headers)
            classifications.append(
                {"name": sn, "role_rescan": hr, "confidence": hconf, "reason": "header_heuristic"}
            )
            if hr == "product_master" and hconf >= 0.5:
                master = sn
                break
        finally:
            wb.close()
        if master:
            break
    if master is None and sheet_names:
        # last resort: first sheet with usable headers
        for sn in sheet_names:
            if sn in ignored:
                continue
            wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
            try:
                ws = wb[sn]
                rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
                header_row = next(rows, None)
                if header_row and any(str(c).strip() for c in header_row):
                    master = sn
                    break
            finally:
                wb.close()
    return master, classifications, ignored


def classify_pim_import_mode(
    *,
    match_src: str,
    sheet_label: str,
    file_name: str,
    meta: dict[str, Any] | None,
    handles: dict[str, Any] | None = None,
    headers: list[str] | None = None,
    sheet_name_role: tuple[str, float, str] | None = None,
) -> tuple[str, dict[str, Any]]:
    """
    product_master: chunked async PIM — DB-only multi-ID relax (no name/vendor heuristics).
    clean_product_master: synthetic row split then strict per synthetic row (no DB-unify on source row).
    generic_raw: strict identifier rules (ETL / non-async paths).
    """
    meta = meta if isinstance(meta, dict) else {}
    handles = handles if isinstance(handles, dict) else {}
    headers = headers if isinstance(headers, list) else []

    for key in ("pim_import_mode", "import_mode_hint", "pim_import_mode_hint"):
        raw = meta.get(key)
        hint = str(raw or "").strip().lower()
        if hint in ("clean", "clean_product_master", "clean_pm"):
            mode = "clean_product_master"
            sig_src = {"source": "metadata", "key": key, "value": hint}
            break
        if hint in ("generic", "generic_raw", "strict"):
            mode = "generic_raw"
            sig_src = {"source": "metadata", "key": key, "value": hint}
            break
        if hint in ("product_master", "pm", "product", "pim_product_master"):
            mode = "product_master"
            sig_src = {"source": "metadata", "key": key, "value": hint}
            break
    else:
        sig_src = {}
        fn = (file_name or "").lower()
        if "clean" in fn and ("master" in fn or "product" in fn):
            mode = "clean_product_master"
            sig_src = {"source": "filename_heuristic", "file_name": file_name}
        elif str(match_src).startswith("pim_async"):
            mode = "product_master"
            sig_src = {"source": "pim_async_default", "match_src": match_src, "sheet": sheet_label}
        else:
            mode = "generic_raw"
            sig_src = {"source": "non_async_default", "match_src": match_src, "sheet": sheet_label}

    mapped = [k for k, v in handles.items() if v]
    hdr_lc = {str(h).strip().lower() for h in headers if str(h).strip()}
    pm_header_hits = sum(
        1
        for needle in ("seller sku", "asin", "fnsku", "product name", "category", "vendor")
        if any(needle in h for h in hdr_lc)
    )
    signals: dict[str, Any] = {
        **sig_src,
        "import_mode": mode,
        "header_count": len(headers),
        "mapped_standard_field_count": len(mapped),
        "mapped_standard_fields": mapped[:24],
        "has_category_handle": bool(handles.get("category")),
        "has_seller_sku_handle": bool(handles.get("seller_sku")),
        "product_master_header_hits": pm_header_hits,
    }
    if sheet_name_role is not None:
        signals["sheet_name_role"] = sheet_name_role[0]
        signals["sheet_name_confidence"] = float(sheet_name_role[1])
        signals["sheet_name_reason"] = str(sheet_name_role[2])
    return mode, signals


def _build_frozen_plan_from_headers(
    headers: list[str],
    openai_key: str | None,
    sheet_label: str,
    match_src: str,
    *,
    file_name: str = "",
    meta: dict[str, Any] | None = None,
) -> dict[str, Any]:
    m = _main()
    column_map, mapping_src = m._pim_resolve_column_map(headers, openai_key)
    m._validate_gpt_column_map(column_map, headers)
    handles = m._pim_handles_from_map(column_map, headers)
    handles = m._pim_augment_handles_from_deterministic(headers, handles)
    sheet_role = classify_sheet_name(sheet_label)
    import_mode, import_signals = classify_pim_import_mode(
        match_src=match_src,
        sheet_label=sheet_label,
        file_name=file_name,
        meta=meta,
        handles=handles,
        headers=headers,
        sheet_name_role=sheet_role,
    )
    return {
        "mapping_source": mapping_src,
        "pim_column_map": m._pim_public_column_map(column_map),
        "handles": handles,
        "headers": headers,
        "sheet_label": sheet_label,
        "match_src": match_src,
        "file_kind": "xlsx",
        # Chunked PIM Product Master preview/apply only — relaxes multi-SKU ambiguity vs generic seed.
        "pim_product_master_seed": str(match_src).startswith("pim_async"),
        "import_mode": import_mode,
        "import_mode_signals": import_signals,
        "classifier_signals": import_signals,
    }


def run_pim_import_preview_step(
    db: Any,
    organization_id: str,
    store_id: str,
    upload_id: str,
    *,
    openai_key: str | None,
    amazon_creds: Any,
    row_chunk: int | None = None,
    scan_data_row_hint: int | None = None,
) -> dict[str, Any]:
    """Single bounded preview step; caller polls until done=True."""
    m = _main()
    t0 = time.monotonic()

    res = (
        db.table("raw_report_uploads")
        .select("id,file_name,metadata,status,created_by")
        .eq("id", upload_id)
        .eq("organization_id", organization_id)
        .limit(1)
        .execute()
    )
    rows = res.data or []
    if not rows:
        return {"ok": False, "error": "upload_not_found", "done": False, "status": "failed", "terminal": True}
    row = rows[0]
    meta = row.get("metadata") if isinstance(row.get("metadata"), dict) else {}
    fname = str(row.get("file_name") or "upload")
    upload_created_by = str(row.get("created_by") or "").strip() or None
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()

    # Resume / idempotent read: return persisted preview contract without re-scanning.
    cached = meta.get("pim_preview_result")
    if isinstance(cached, dict) and str(cached.get("status") or "") == "preview_ready":
        raw_q = cached.get("raw_quality") if isinstance(cached.get("raw_quality"), dict) else None
        pub_q = cached.get("quality") if isinstance(cached.get("quality"), dict) else None
        if raw_q is not None and pub_q is not None and "rows_scanned" in pub_q:
            sid = meta.get("pim_seed_session_id")
            return _strip_heavy_preview_ready_response(
                {
                    "ok": True,
                    "done": True,
                    "terminal": True,
                    "status": "preview_ready",
                    "import_job_id": str(upload_id),
                    "preview_quality": raw_q,
                    "quality": pub_q,
                    "mapping": cached.get("mapping") if isinstance(cached.get("mapping"), dict) else {},
                    "mapping_source": str(cached.get("mapping_source") or ""),
                    "selected_sheets": cached.get("selected_sheets")
                    if isinstance(cached.get("selected_sheets"), list)
                    else [],
                    "ignored_sheets": cached.get("ignored_sheets")
                    if isinstance(cached.get("ignored_sheets"), list)
                    else [],
                    "samples": cached.get("samples")
                    if isinstance(cached.get("samples"), dict)
                    else {"accepted": [], "rejected": []},
                    "errors": cached.get("errors") if isinstance(cached.get("errors"), list) else [],
                    "seed_session_id": str(sid) if sid else None,
                    "progress_pct": 100,
                    "stage_label": "preview_ready",
                    "lifecycle": "waiting_for_confirmation",
                    "cached": True,
                }
            )

    _fps_update(
        db,
        upload_id,
        organization_id,
        process_pct=5,
        processed_rows=0,
        total_rows=None,
        phase_label="reading_workbook",
    )

    meta = _ensure_pim_merged_source(db, organization_id, upload_id, meta, fname)
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()

    # Check cancel signal (file_processing_status or metadata flags)
    if _pim_import_meta_cancelled(meta, db=db, upload_id=upload_id):
        meta = _fps_cancel_merge(meta)
        _persist_meta(db, organization_id, upload_id, meta, "cancelled")
        return {
            "ok": False,
            "done": True,
            "terminal": True,
            "status": "cancelled",
            "import_job_id": str(upload_id),
            "lifecycle": "cancelled",
            "last_error": "cancelled_by_user",
        }

    local_path: str | None = None
    try:
        local_path, _unlink = materialize_job_local_file(db, meta, fname)
    except Exception as e:
        err = str(e)
        meta = merge_pim_job(
            meta,
            {"lifecycle": "failed", "last_error": err, "stage_label": "download_failed", "preview_phase": None},
        )
        meta["preview_status"] = "failed"
        meta["import_job_status"] = "failed"
        _persist_meta(db, organization_id, upload_id, meta, "failed")
        # (pim_import_sessions removed — raw_report_uploads.metadata is the source of truth)
        return {
            "ok": False,
            "error": err,
            "last_error": err,
            "done": True,
            "terminal": True,
            "status": "failed",
            "import_job_id": str(upload_id),
            "stage_label": "download_failed",
        }

    try:
        is_xlsx = fname.lower().endswith((".xlsx", ".xlsm")) or str(meta.get("file_extension") or "").lower() in (
            "xlsx",
            "xlsm",
        )
        if not is_xlsx and os.path.splitext(local_path)[1].lower() in (".xlsx", ".xlsm"):
            is_xlsx = True

        preview_phase = job.get("preview_phase")

        # ── Phase detect ─────────────────────────────────────────────────────
        if preview_phase in (None, "", "detect"):
            meta = merge_pim_job(
                meta,
                {
                    "lifecycle": "previewing",
                    "preview_phase": "detect",
                    "stage_label": "detecting_sheets",
                    "started_at_monotonic": time.monotonic(),
                },
            )
            _persist_meta(db, organization_id, upload_id, meta, "processing")

            if is_xlsx:
                sheet_names = _list_xlsx_sheet_names(local_path)
                master, classifs, ignored = _pick_master_sheet(local_path, sheet_names)
                if not master:
                    meta = merge_pim_job(
                        meta,
                        {
                            "lifecycle": "failed",
                            "last_error": "Could not find a product master sheet",
                            "preview_phase": None,
                            "stage_label": "no_master_sheet",
                        },
                    )
                    meta["preview_status"] = "failed"
                    _persist_meta(db, organization_id, upload_id, meta, "failed")
                    return {
                        "ok": False,
                        "error": "no_master_sheet",
                        "last_error": "Could not find a product master sheet (expected App_Import_Product_Master or matching headers).",
                        "done": True,
                        "terminal": True,
                        "status": "failed",
                        "import_job_id": str(upload_id),
                        "stage_label": "no_master_sheet",
                    }

                import openpyxl

                wb = openpyxl.load_workbook(local_path, read_only=True, data_only=True)
                try:
                    ws = wb[master]
                    hdr_it = ws.iter_rows(min_row=1, max_row=1, values_only=True)
                    header_row = next(hdr_it, None)
                    headers = [str(c).strip() if c is not None else "" for c in (header_row or [])]
                    max_row = ws.max_row or 0
                    total_data_rows = max(0, max_row - 1)
                finally:
                    wb.close()

                if total_data_rows == 0:
                    try:
                        total_data_rows = _count_excel_data_rows(local_path, master)
                    except Exception:
                        total_data_rows = 0

                frozen = _build_frozen_plan_from_headers(
                    headers,
                    openai_key,
                    master,
                    f"pim_async:{master}",
                    file_name=fname,
                    meta=meta,
                )
                acc = m._pim_seed_quality_acc()
                job_patch = {
                    "preview_phase": "scan",
                    "frozen_plan": frozen,
                    "import_mode": frozen.get("import_mode"),
                    "import_mode_signals": frozen.get("import_mode_signals"),
                    "classifier_signals": frozen.get("classifier_signals"),
                    "selected_master_sheet": master,
                    "detected_sheets": sheet_names,
                    "ignored_sheets": ignored,
                    "sheet_classifications": classifs,
                    "quality_acc": acc,
                    "scan_cursor": {"sheet": master, "data_row": 0, "total_rows": total_data_rows},
                    "progress_pct": 25,
                    "stage_label": "scanning_rows",
                }
                meta = merge_pim_job(meta, job_patch)
                meta["row_count"] = total_data_rows
                _persist_meta(db, organization_id, upload_id, meta, "processing")
                return {
                    "ok": True,
                    "done": False,
                    "lifecycle": "previewing",
                    "progress_pct": 25,
                    "stage_label": "scanning_rows",
                    "elapsed_sec": time.monotonic() - t0,
                }
            else:
                # CSV / TXT
                with open(local_path, "rb") as bf:
                    raw = bf.read()
                df, delim2, _unc = m._read_text_delimited(raw, fname)
                if hasattr(df, "attrs"):
                    df.attrs["pim_delimiter_detected"] = delim2
                headers = [str(c).strip() for c in df.columns]
                frozen = _build_frozen_plan_from_headers(
                    headers, openai_key, "(csv)", "pim_async_csv", file_name=fname, meta=meta
                )
                frozen["file_kind"] = "csv"
                total_data_rows = int(df.shape[0])
                acc = m._pim_seed_quality_acc()
                meta = merge_pim_job(
                    meta,
                    {
                        "lifecycle": "previewing",
                        "preview_phase": "scan",
                        "frozen_plan": frozen,
                        "import_mode": frozen.get("import_mode"),
                        "import_mode_signals": frozen.get("import_mode_signals"),
                        "classifier_signals": frozen.get("classifier_signals"),
                        "selected_master_sheet": "(csv)",
                        "detected_sheets": ["(csv)"],
                        "quality_acc": acc,
                        "scan_cursor": {"sheet": "(csv)", "data_row": 0, "total_rows": total_data_rows},
                        "progress_pct": 25,
                        "stage_label": "scanning_rows",
                        "pim_delimiter_detected": str(delim2),
                    },
                )
                meta["row_count"] = total_data_rows
                _persist_meta(db, organization_id, upload_id, meta, "processing")
                return {
                    "ok": True,
                    "done": False,
                    "lifecycle": "previewing",
                    "progress_pct": 25,
                    "stage_label": "scanning_rows",
                    "elapsed_sec": time.monotonic() - t0,
                }

        # ── Phase scan ────────────────────────────────────────────────────────
        frozen = job.get("frozen_plan")
        if not isinstance(frozen, dict):
            meta = merge_pim_job(
                meta,
                {
                    "lifecycle": "failed",
                    "last_error": "missing_frozen_plan",
                    "preview_phase": None,
                    "stage_label": "missing_frozen_plan",
                },
            )
            meta["preview_status"] = "failed"
            _persist_meta(db, organization_id, upload_id, meta, "failed")
            return {
                "ok": False,
                "error": "missing_frozen_plan",
                "last_error": "Preview state lost (frozen_plan). Re-upload and run preview again.",
                "done": True,
                "terminal": True,
                "status": "failed",
                "import_job_id": str(upload_id),
            }

        handles = frozen.get("handles")
        headers = frozen.get("headers")
        if not isinstance(handles, dict) or not isinstance(headers, list):
            meta = merge_pim_job(
                meta,
                {
                    "lifecycle": "failed",
                    "last_error": "invalid_frozen_plan",
                    "preview_phase": None,
                    "stage_label": "invalid_frozen_plan",
                },
            )
            meta["preview_status"] = "failed"
            _persist_meta(db, organization_id, upload_id, meta, "failed")
            return {
                "ok": False,
                "error": "invalid_frozen_plan",
                "last_error": "Invalid column mapping state. Re-upload and run preview again.",
                "done": True,
                "terminal": True,
                "status": "failed",
                "import_job_id": str(upload_id),
            }

        acc = job.get("quality_acc")
        if not isinstance(acc, dict):
            acc = m._pim_seed_quality_acc()
        if "_pim_scan_has_category_handle" not in acc:
            acc["_pim_scan_has_category_handle"] = bool(handles.get("category"))

        vendor_index = m._pim_load_vendor_index(db, organization_id)
        category_index = m._pim_load_category_index(db, organization_id)
        cursor = job.get("scan_cursor") or {}
        data_row_start = int(cursor.get("data_row") or 0)
        total_rows_hint = int(cursor.get("total_rows") or 0)
        if scan_data_row_hint is not None and int(scan_data_row_hint) != int(data_row_start):
            log.warning(
                "pim_import preview-step cursor_hint_mismatch job_id=%s hint=%s db_cursor=%s",
                upload_id,
                scan_data_row_hint,
                data_row_start,
            )
        sheet_label = str(frozen.get("sheet_label") or "")
        file_kind = str(frozen.get("file_kind") or "xlsx")
        pim_mode = str(frozen.get("import_mode") or "").strip().lower()
        if not pim_mode:
            pim_mode = "product_master" if bool(frozen.get("pim_product_master_seed")) else "generic_raw"

        base_default = ROW_CHUNK_CSV_DEFAULT if file_kind == "csv" else ROW_CHUNK_XLSX_DEFAULT
        eff_stored = job.get("effective_row_chunk")
        eff_int = int(eff_stored) if isinstance(eff_stored, (int, float)) and int(eff_stored) > 0 else 0
        req = int(row_chunk) if row_chunk is not None and int(row_chunk) > 0 else 0
        chunk_cap = req or eff_int or base_default
        chunk_cap = max(ROW_CHUNK_MIN, min(chunk_cap, ROW_CHUNK_MAX))
        log.info(
            "pim_import preview-step slice job_id=%s kind=%s cursor=%s chunk_cap=%s total_hint=%s",
            upload_id,
            file_kind,
            data_row_start,
            chunk_cap,
            total_rows_hint,
        )

        processed_in_step = 0
        next_cursor = data_row_start
        df_rows_total = 0
        if file_kind == "csv":
            df, meta = _load_csv_df_from_cache_or_parse(
                db, meta, local_path, fname, organization_id=organization_id, upload_id=upload_id
            )
            df_rows_total = int(df.shape[0])
            end_idx = min(data_row_start + chunk_cap, df_rows_total)
            if data_row_start >= df_rows_total:
                next_cursor = data_row_start
                done_scan = True
                processed_in_step = 0
            else:
                rows_chunk: list[tuple[int, dict[str, Any], int]] = []
                scan_idx = data_row_start
                while scan_idx < end_idx:
                    if time.monotonic() - t0 > PREVIEW_STEP_TIME_BUDGET_SEC * 0.35:
                        break
                    row = df.iloc[scan_idx]
                    row_cells, row_trim = m._pim_row_cells_from_series(row, headers)
                    rows_chunk.append((scan_idx, row_cells, row_trim))
                    scan_idx += 1
                if not rows_chunk and data_row_start < df_rows_total:
                    row = df.iloc[data_row_start]
                    row_cells, row_trim = m._pim_row_cells_from_series(row, headers)
                    rows_chunk.append((data_row_start, row_cells, row_trim))
                all_asin: list[str] = []
                all_upc: list[str] = []
                for _idx, rc, _rt in rows_chunk:
                    pr = parse_identifier_row(rc, handles)
                    all_asin.extend(list(pr.asin.accepted))
                    all_upc.extend(list(pr.upc.accepted))
                asin_pref = m._pim_product_ids_for_values_batch(
                    db, organization_id, store_id, all_asin, "asin"
                )
                upc_pref = m._pim_product_ids_for_values_batch(
                    db, organization_id, store_id, all_upc, "upc_code"
                )
                chunk_db_cache = m._pim_preview_prepare_chunk_db_cache(
                    db,
                    organization_id,
                    store_id,
                    rows_chunk,
                    handles,
                    headers,
                    asin_pref,
                    upc_pref,
                )
                processed_in_step = 0
                for idx, row_cells, row_trim in rows_chunk:
                    if time.monotonic() - t0 > PREVIEW_STEP_TIME_BUDGET_SEC:
                        break
                    row_lbl = idx + 2
                    m._pim_analyze_seed_row_for_quality(
                        db,
                        organization_id,
                        store_id,
                        row_cells,
                        handles,
                        headers,
                        vendor_index,
                        category_index,
                        acc,
                        row_lbl,
                        row_trim,
                        amazon_creds=amazon_creds,
                        skip_amazon=True,
                        asin_prefetch=asin_pref,
                        upc_prefetch=upc_pref,
                        chunk_db_cache=chunk_db_cache,
                        pim_import_mode=pim_mode,
                    )
                    processed_in_step += 1
                next_cursor = data_row_start + processed_in_step
                done_scan = next_cursor >= df_rows_total
        else:
            import openpyxl

            wb = openpyxl.load_workbook(local_path, read_only=True, data_only=True)
            hit_eof = False
            try:
                ws = wb[sheet_label]
                rows_chunk: list[tuple[int, dict[str, Any], int]] = []
                rows_it = ws.iter_rows(min_row=2 + data_row_start, values_only=True)
                while len(rows_chunk) < chunk_cap:
                    if time.monotonic() - t0 > PREVIEW_STEP_TIME_BUDGET_SEC * 0.35:
                        break
                    try:
                        raw_row = next(rows_it)
                    except StopIteration:
                        hit_eof = True
                        break
                    cells = list(raw_row)
                    pad = len(headers) - len(cells)
                    if pad > 0:
                        cells = cells + [None] * pad
                    elif len(cells) > len(headers):
                        cells = cells[: len(headers)]
                    row_cells, row_trim = m._pim_row_cells_from_series(
                        pd.Series({headers[i]: cells[i] for i in range(len(headers))}),
                        headers,
                    )
                    scan_idx = data_row_start + len(rows_chunk)
                    rows_chunk.append((scan_idx, row_cells, row_trim))
                if rows_chunk:
                    all_asin: list[str] = []
                    all_upc: list[str] = []
                    for _idx, rc, _rt in rows_chunk:
                        pr = parse_identifier_row(rc, handles)
                        all_asin.extend(list(pr.asin.accepted))
                        all_upc.extend(list(pr.upc.accepted))
                    asin_pref = m._pim_product_ids_for_values_batch(
                        db, organization_id, store_id, all_asin, "asin"
                    )
                    upc_pref = m._pim_product_ids_for_values_batch(
                        db, organization_id, store_id, all_upc, "upc_code"
                    )
                    chunk_db_cache = m._pim_preview_prepare_chunk_db_cache(
                        db,
                        organization_id,
                        store_id,
                        rows_chunk,
                        handles,
                        headers,
                        asin_pref,
                        upc_pref,
                    )
                    processed_in_step = 0
                    for idx, row_cells, row_trim in rows_chunk:
                        if time.monotonic() - t0 > PREVIEW_STEP_TIME_BUDGET_SEC:
                            break
                        row_lbl = f"{sheet_label}!{idx + 2}"
                        m._pim_analyze_seed_row_for_quality(
                            db,
                            organization_id,
                            store_id,
                            row_cells,
                            handles,
                            headers,
                            vendor_index,
                            category_index,
                            acc,
                            row_lbl,
                            row_trim,
                            amazon_creds=amazon_creds,
                            skip_amazon=True,
                            asin_prefetch=asin_pref,
                            upc_prefetch=upc_pref,
                            chunk_db_cache=chunk_db_cache,
                            pim_import_mode=pim_mode,
                        )
                        processed_in_step += 1
                next_cursor = data_row_start + processed_in_step
                # openpyxl read_only ws.max_row / total_rows_hint are often wrong on large sheets — only EOF ends scan.
                if hit_eof:
                    done_scan = True
                else:
                    done_scan = False
            finally:
                wb.close()

        # Check cancel signal between scan chunks
        if _pim_import_meta_cancelled(meta, db=db, upload_id=upload_id):
            meta = _fps_cancel_merge(meta)
            meta = merge_pim_job(
                meta,
                {
                    "quality_acc": acc,
                    "scan_cursor": {
                        "sheet": sheet_label,
                        "data_row": next_cursor,
                        "total_rows": total_rows_hint,
                    },
                    "progress_pct": min(90, 30 + int(60 * min(next_cursor / float(max(total_rows_hint or next_cursor or 1, 1)), 1.0))),
                    "stage_label": "scanning_rows",
                },
            )
            _persist_meta(db, organization_id, upload_id, meta, "cancelled")
            return {
                "ok": False,
                "done": True,
                "terminal": True,
                "status": "cancelled",
                "import_job_id": str(upload_id),
                "lifecycle": "cancelled",
                "last_error": "cancelled_by_user",
            }

        total_for_pct = df_rows_total if file_kind == "csv" else max(total_rows_hint, next_cursor, 1)
        pct = 30
        if total_for_pct > 0:
            pct = min(90, 30 + int(60 * min(next_cursor / float(total_for_pct), 1.0)))

        new_eff: int | None = None
        if not done_scan and processed_in_step == 0 and (time.monotonic() - t0) >= PREVIEW_STEP_TIME_BUDGET_SEC * 0.95:
            new_eff = max(ROW_CHUNK_MIN, chunk_cap // 2)

        scan_total_rows = total_rows_hint
        if file_kind != "csv":
            scan_total_rows = max(int(total_rows_hint or 0), int(next_cursor), int(next_cursor) + (0 if done_scan else 1))
        patch_scan: dict[str, Any] = {
            "quality_acc": acc,
            "scan_cursor": {
                "sheet": sheet_label,
                "data_row": next_cursor,
                "total_rows": scan_total_rows,
            },
            "progress_pct": pct,
            "stage_label": "scanning_rows",
            "elapsed_preview_sec": float(job.get("elapsed_preview_sec") or 0) + (time.monotonic() - t0),
            "preview_progress": {
                "parsed_rows": next_cursor,
                "mapped_rows": next_cursor,
                "validated_rows": next_cursor,
                "resolved_rows": next_cursor,
                "total_rows": total_for_pct,
                "percent": pct,
                "current_step": "scanning_rows",
                "last_heartbeat": datetime.now(timezone.utc).isoformat(),
            },
        }
        if new_eff is not None:
            patch_scan["effective_row_chunk"] = new_eff
        meta = merge_pim_job(meta, patch_scan)
        _persist_meta(db, organization_id, upload_id, meta, "processing")
        q_part = m._pim_seed_quality_finalize(dict(acc))
        _fps_update(
            db,
            upload_id,
            organization_id,
            process_pct=pct,
            processed_rows=next_cursor,
            total_rows=total_rows_hint or None,
            phase_label="scanning_rows",
        )
        try:
            pp = patch_scan.get("preview_progress") or {}
            approx_out = len(json.dumps({"preview_progress": pp, "rows_total": acc.get("rows_total")}))
        except Exception:
            approx_out = 0
        log.info(
            "pim_import preview-step chunk_persist job_id=%s processed_in_step=%s next_cursor=%s duration_ms=%.0f response_hint_bytes=%s",
            upload_id,
            processed_in_step,
            next_cursor,
            (time.monotonic() - t0) * 1000,
            approx_out,
        )

        if not done_scan:
            return {
                "ok": True,
                "done": False,
                "lifecycle": "previewing",
                "progress_pct": pct,
                "stage_label": "scanning_rows",
                "rows_scanned": next_cursor,
                "total_rows": total_rows_hint,
                "next_cursor": next_cursor,
                "import_job_id": str(upload_id),
                "elapsed_sec": time.monotonic() - t0,
                "preview_progress": patch_scan.get("preview_progress"),
            }

        # ── Finalize preview ───────────────────────────────────────────────────
        quality = m._pim_seed_quality_finalize(acc)
        j_latest = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()
        sig = dict(frozen.get("classifier_signals") if isinstance(frozen.get("classifier_signals"), dict) else {})
        if isinstance(j_latest.get("classifier_signals"), dict):
            sig.update(j_latest["classifier_signals"])
        rt_scan = max(int(quality.get("rows_total") or 1), 1)
        mh = int(quality.get("preview_rows_multi_identifier") or 0)
        sig["multi_identifier_row_hits"] = mh
        sig["multi_identifier_row_fraction"] = round(mh / float(rt_scan), 5)
        sig["clean_pm_physical_rows_split"] = int(quality.get("clean_pm_physical_rows_split") or 0)
        frozen2 = dict(frozen)
        frozen2["classifier_signals"] = sig
        frozen2["preview_plan_version"] = 1
        frozen2["preview_accepted_row_total"] = int(quality.get("rows_total") or 0)
        meta = merge_pim_job(meta, {"frozen_plan": frozen2, "classifier_signals": sig})
        frozen = frozen2

        col_map = frozen.get("pim_column_map") if isinstance(frozen.get("pim_column_map"), dict) else {}
        sheet_key = str(job.get("selected_master_sheet") or frozen.get("sheet_label") or "").strip()
        map_ok = len(col_map) > 0
        sheet_ok = bool(sheet_key)
        rows_ok = "rows_total" in quality
        if not (map_ok and sheet_ok and rows_ok):
            reason = f"preview_incomplete mapping_ok={map_ok} sheet_ok={sheet_ok} rows_ok={rows_ok}"
            meta = merge_pim_job(
                meta,
                {
                    "lifecycle": "failed",
                    "last_error": reason,
                    "preview_phase": None,
                    "stage_label": "preview_incomplete",
                    "preview_quality": quality,
                },
            )
            meta["preview_status"] = "failed"
            _persist_meta(db, organization_id, upload_id, meta, "failed")
            return {
                "ok": False,
                "done": True,
                "terminal": True,
                "status": "failed",
                "error": "Preview incomplete: need a selected sheet, column mapping, and row scan.",
                "last_error": reason,
                "import_job_id": str(upload_id),
                "preview_quality": quality,
                "stage_label": "preview_incomplete",
                "elapsed_sec": time.monotonic() - t0,
            }

        hist_id = m._pim_seed_history_insert_preview(
            db,
            organization_id,
            store_id,
            fname,
            col_map,
            str(frozen.get("mapping_source") or "unknown"),
            quality,
        )

        meta = merge_pim_job(
            meta,
            {
                "lifecycle": "waiting_for_confirmation",
                "preview_phase": "done",
                "preview_quality": quality,
                "import_total_rows": int(quality.get("rows_total") or 0),
                "progress_pct": 100,
                "stage_label": "preview_ready",
                "scan_cursor": None,
            },
        )
        jf = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()
        public_payload = build_public_preview_payload(
            upload_id=str(upload_id),
            quality=quality,
            frozen=frozen,
            job=jf,
            seed_session_id=str(hist_id) if hist_id else None,
        )
        meta["pim_preview_result"] = public_payload
        meta["preview_status"] = "preview_ready"
        meta["import_job_status"] = "preview_ready"
        if hist_id:
            meta["pim_seed_session_id"] = str(hist_id)
        _persist_meta(db, organization_id, upload_id, meta, "mapped")
        _fps_update(
            db,
            upload_id,
            organization_id,
            process_pct=100,
            processed_rows=quality.get("rows_total"),
            total_rows=quality.get("rows_total"),
            phase_label="preview_ready",
        )

        return _strip_heavy_preview_ready_response(
            {
                "ok": True,
                "done": True,
                "terminal": True,
                "status": "preview_ready",
                "import_job_id": str(upload_id),
                "lifecycle": "waiting_for_confirmation",
                "preview_quality": quality,
                "quality": public_payload["quality"],
                "preview_metrics": public_payload["preview_metrics"],
                "mapping": public_payload["mapping"],
                "mapping_source": public_payload["mapping_source"],
                "selected_sheets": public_payload["selected_sheets"],
                "ignored_sheets": public_payload["ignored_sheets"],
                "samples": public_payload["samples"],
                "errors": public_payload["errors"],
                "seed_session_id": hist_id,
                "progress_pct": 100,
                "stage_label": "preview_ready",
                "elapsed_sec": time.monotonic() - t0,
            }
        )
    finally:
        if local_path:
            try:
                os.unlink(local_path)
            except OSError:
                pass


def get_pim_import_preview_status(
    db: Any,
    organization_id: str,
    store_id: str,
    upload_id: str,
) -> dict[str, Any]:
    """Read-only: return persisted preview contract or in-progress job state (same shape as final preview-step when ready)."""
    _ = store_id  # reserved for future store-scoped checks
    res = (
        db.table("raw_report_uploads")
        .select("id,metadata")
        .eq("id", upload_id)
        .eq("organization_id", organization_id)
        .limit(1)
        .execute()
    )
    rows = res.data or []
    if not rows:
        return {"ok": False, "error": "upload_not_found", "terminal": True, "status": "failed"}
    meta = rows[0].get("metadata") if isinstance(rows[0].get("metadata"), dict) else {}
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()
    if _pim_import_meta_cancelled(meta, db, organization_id, upload_id):
        return {
            "ok": True,
            "done": True,
            "terminal": True,
            "status": "cancelled",
            "import_job_id": str(upload_id),
            "import_session_id": str(meta.get("pim_import_session_id") or ""),
            "stage_label": "cancelled",
            "lifecycle": "cancelled",
            "progress_pct": int(job.get("progress_pct") or 0),
            "last_error": job.get("last_error") or "cancelled_by_user",
        }
    cached = meta.get("pim_preview_result")
    if (
        not _pim_import_meta_cancelled(meta, db, organization_id, upload_id)
        and isinstance(cached, dict)
        and str(cached.get("status") or "") == "preview_ready"
    ):
        raw_q = cached.get("raw_quality") if isinstance(cached.get("raw_quality"), dict) else None
        pub_q = cached.get("quality") if isinstance(cached.get("quality"), dict) else None
        if raw_q is not None and pub_q is not None and "rows_scanned" in pub_q:
            sid = meta.get("pim_seed_session_id")
            return _strip_heavy_preview_ready_response(
                {
                    "ok": True,
                    "done": True,
                    "terminal": True,
                    "status": "preview_ready",
                    "import_job_id": str(upload_id),
                    "preview_quality": raw_q,
                    "quality": pub_q,
                    "mapping": cached.get("mapping") if isinstance(cached.get("mapping"), dict) else {},
                    "mapping_source": str(cached.get("mapping_source") or ""),
                    "selected_sheets": cached.get("selected_sheets")
                    if isinstance(cached.get("selected_sheets"), list)
                    else [],
                    "ignored_sheets": cached.get("ignored_sheets")
                    if isinstance(cached.get("ignored_sheets"), list)
                    else [],
                    "samples": cached.get("samples")
                    if isinstance(cached.get("samples"), dict)
                    else {"accepted": [], "rejected": []},
                    "errors": cached.get("errors") if isinstance(cached.get("errors"), list) else [],
                    "seed_session_id": str(sid) if sid else None,
                    "progress_pct": 100,
                    "stage_label": "preview_ready",
                    "lifecycle": job.get("lifecycle") or "waiting_for_confirmation",
                }
            )
    return {
        "ok": True,
        "done": False,
        "terminal": False,
        "status": str(meta.get("preview_status") or job.get("lifecycle") or "previewing"),
        "import_job_id": str(upload_id),
        "stage_label": str(job.get("stage_label") or ""),
        "progress_pct": int(job.get("progress_pct") or 0),
        "lifecycle": str(job.get("lifecycle") or ""),
        "last_error": job.get("last_error"),
    }


def run_pim_import_apply_step(
    db: Any,
    organization_id: str,
    store_id: str,
    upload_id: str,
    *,
    amazon_creds: Any,
    price_source: str,
    row_chunk: int | None = None,
    confirmed_via_api: bool = False,
) -> dict[str, Any]:
    """Chunked apply using frozen_plan from preview."""
    m = _main()
    t0 = time.monotonic()

    res = (
        db.table("raw_report_uploads")
        .select("id,file_name,metadata,created_by")
        .eq("id", upload_id)
        .eq("organization_id", organization_id)
        .limit(1)
        .execute()
    )
    rows = res.data or []
    if not rows:
        return {"ok": False, "error": "upload_not_found", "done": True}
    row_ap = rows[0]
    meta = row_ap.get("metadata") if isinstance(row_ap.get("metadata"), dict) else {}
    fname = str(row_ap.get("file_name") or "upload")
    apply_created_by = str(row_ap.get("created_by") or "").strip() or None
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()

    if _pim_import_meta_cancelled(meta, db, organization_id, upload_id):
        return {
            "ok": False,
            "error": "cancelled",
            "done": True,
            "terminal": True,
            "status": "cancelled",
            "last_error": "cancelled_by_user",
        }

    if not (bool(meta.get("pim_import_confirmed")) or confirmed_via_api):
        return {
            "ok": False,
            "error": "confirm_required",
            "done": False,
            "terminal": False,
            "status": "waiting_for_confirmation",
            "message": "Confirm import in the UI before applying rows.",
        }

    if job.get("lifecycle") == "completed" and isinstance(job.get("apply_metrics"), dict):
        return {"ok": True, "done": True, "metrics": job["apply_metrics"], "cached": True}

    frozen = job.get("frozen_plan")
    if not isinstance(frozen, dict):
        return {
            "ok": False,
            "error": "preview_plan_missing",
            "done": True,
            "terminal": True,
            "message": "Preview plan expired. Re-run preview.",
        }

    pim_mode_apply = str(frozen.get("import_mode") or "").strip().lower()
    if not pim_mode_apply:
        pim_mode_apply = "product_master" if bool(frozen.get("pim_product_master_seed")) else "generic_raw"

    handles = frozen.get("handles")
    headers = frozen.get("headers")
    if not isinstance(handles, dict) or not isinstance(headers, list):
        return {"ok": False, "error": "invalid_frozen_plan", "done": True}

    metrics = job.get("apply_metrics")
    if not isinstance(metrics, dict):
        metrics = m._empty_pim_seed_metrics()

    sheet_label = str(frozen.get("sheet_label") or "")
    file_kind = str(frozen.get("file_kind") or "xlsx")
    cursor = job.get("apply_cursor") or {"data_row": 0}
    data_row_start = int(cursor.get("data_row") or 0)


    meta = _ensure_pim_merged_source(db, organization_id, upload_id, meta, fname)
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else default_pim_job()

    local_path, _ = materialize_job_local_file(db, meta, fname)
    try:
        try:
            vendor_index = m._pim_load_vendor_index(db, organization_id)
            category_index = m._pim_load_category_index(db, organization_id)
            brand_by_mfg_map: dict[str, str] | None = None
            br_raw = frozen.get("brand_by_mfg_map")
            if isinstance(br_raw, dict) and br_raw:
                brand_by_mfg_map = {
                    str(k).strip().lower(): str(v).strip()
                    for k, v in br_raw.items()
                    if k is not None and v is not None and str(v).strip()
                }
            elif file_kind != "csv":
                try:
                    loaded = m._pim_load_brand_by_mfg_from_xlsx(local_path)
                    if loaded:
                        brand_by_mfg_map = loaded
                        frozen_new = dict(frozen)
                        frozen_new["brand_by_mfg_map"] = loaded
                        frozen = frozen_new
                        meta = merge_pim_job(meta, {"frozen_plan": frozen_new})
                        _persist_meta(db, organization_id, upload_id, meta, "processing")
                except Exception:
                    log.debug("brand_by_mfg optional sheet missing or unreadable", exc_info=True)
            processed = 0
            next_a = data_row_start
            total_cap = int(job.get("import_total_rows") or 0)
            total_r = 0
            done = False
            base_default = ROW_CHUNK_CSV_DEFAULT if file_kind == "csv" else ROW_CHUNK_XLSX_DEFAULT
            eff_stored = job.get("apply_effective_row_chunk")
            eff_int = int(eff_stored) if isinstance(eff_stored, (int, float)) and int(eff_stored) > 0 else 0
            req = int(row_chunk) if row_chunk is not None and int(row_chunk) > 0 else 0
            chunk_cap = req or eff_int or base_default
            # Allow sub-ROW_CHUNK_MIN chunks when a previous failure set a smaller effective size
            # (retry mode after ConnectionTerminated or other connection errors).
            effective_min = ROW_CHUNK_RETRY_MIN if (eff_int > 0 and eff_int < ROW_CHUNK_MIN) else ROW_CHUNK_MIN
            chunk_cap = max(effective_min, min(chunk_cap, ROW_CHUNK_MAX))
            if file_kind == "csv":
                df, meta = _load_csv_df_from_cache_or_parse(
                    db, meta, local_path, fname, organization_id=organization_id, upload_id=upload_id
                )
                total_r = int(df.shape[0])
                if total_cap <= 0:
                    total_cap = total_r
                if _pim_import_meta_cancelled(meta, db=db, upload_id=upload_id):
                    meta = _fps_cancel_merge(meta)
                    _persist_meta(db, organization_id, upload_id, meta, "cancelled")
                    return {
                        "ok": False,
                        "error": "cancelled",
                        "done": True,
                        "terminal": True,
                        "status": "cancelled",
                        "metrics_partial": metrics,
                    }
                while processed < chunk_cap and data_row_start + processed < total_r:
                    if time.monotonic() - t0 > APPLY_STEP_TIME_BUDGET_SEC:
                        break
                    idx = data_row_start + processed
                    row = df.iloc[idx]
                    row_cells, row_trim = m._pim_row_cells_from_series(row, headers)
                    row_lbl = idx + 2
                    m._process_pim_seed_row(
                        db,
                        organization_id,
                        store_id,
                        row_cells,
                        handles,
                        headers,
                        amazon_creds,
                        price_source,
                        str(frozen.get("match_src") or "pim_async"),
                        metrics,
                        metrics["errors"],
                        row_lbl,
                        vendor_index,
                        category_index,
                        row_trim,
                        pim_import_mode=pim_mode_apply,
                        skip_amazon_enrichment=True,
                        pim_upload_id=str(upload_id),
                        brand_by_mfg_map=brand_by_mfg_map,
                    )
                    processed += 1
                next_a = data_row_start + processed
                done = next_a >= total_r
            else:
                import openpyxl

                if total_cap <= 0:
                    total_cap = 100000
                if _pim_import_meta_cancelled(meta, db=db, upload_id=upload_id):
                    meta = _fps_cancel_merge(meta)
                    _persist_meta(db, organization_id, upload_id, meta, "cancelled")
                    return {
                        "ok": False,
                        "error": "cancelled",
                        "done": True,
                        "terminal": True,
                        "status": "cancelled",
                        "metrics_partial": metrics,
                    }
                hit_eof = False
                wb = openpyxl.load_workbook(local_path, read_only=True, data_only=True)
                try:
                    ws = wb[sheet_label]
                    rows_it = ws.iter_rows(min_row=2 + data_row_start, values_only=True)
                    while processed < chunk_cap:
                        if time.monotonic() - t0 > APPLY_STEP_TIME_BUDGET_SEC:
                            break
                        try:
                            raw_row = next(rows_it)
                        except StopIteration:
                            hit_eof = True
                            break
                        cells = list(raw_row)
                        pad = len(headers) - len(cells)
                        if pad > 0:
                            cells = cells + [None] * pad
                        elif len(cells) > len(headers):
                            cells = cells[: len(headers)]
                        row_cells, row_trim = m._pim_row_cells_from_series(
                            pd.Series({headers[i]: cells[i] for i in range(len(headers))}),
                            headers,
                        )
                        row_lbl = f"{sheet_label}!{data_row_start + processed + 2}"
                        m._process_pim_seed_row(
                            db,
                            organization_id,
                            store_id,
                            row_cells,
                            handles,
                            headers,
                            amazon_creds,
                            price_source,
                            str(frozen.get("match_src") or "pim_async"),
                            metrics,
                            metrics["errors"],
                            row_lbl,
                            vendor_index,
                            category_index,
                            row_trim,
                            pim_import_mode=pim_mode_apply,
                            skip_amazon_enrichment=True,
                            pim_upload_id=str(upload_id),
                            brand_by_mfg_map=brand_by_mfg_map,
                        )
                        processed += 1
                    next_a = data_row_start + processed
                    # Same as preview scan: do not trust import_total_rows / max_row — only EOF completes xlsx apply.
                    if hit_eof:
                        done = True
                    else:
                        done = False
                finally:
                    wb.close()

            if processed == 0 and (time.monotonic() - t0) >= APPLY_STEP_TIME_BUDGET_SEC * 0.95:
                new_apply_eff = max(ROW_CHUNK_MIN, chunk_cap // 2)
                meta = merge_pim_job(meta, {"apply_effective_row_chunk": new_apply_eff})
                _persist_meta(db, organization_id, upload_id, meta, "processing")

            preview_row_goal = int(frozen.get("preview_accepted_row_total") or job.get("import_total_rows") or 0)
            if file_kind == "csv":
                denom = max(total_r, 1)
            else:
                denom = max(preview_row_goal, total_cap, next_a, 1)
            pct = min(99, int(100 * next_a / float(denom)))

            meta = merge_pim_job(
                meta,
                {
                    "lifecycle": "importing",
                    "apply_metrics": metrics,
                    "apply_cursor": {"data_row": next_a},
                    "progress_pct": pct,
                    "stage_label": "importing",
                },
            )
            _persist_meta(db, organization_id, upload_id, meta, "processing")
            job_mid = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else {}
            pq_snap = job_mid.get("preview_quality") if isinstance(job_mid.get("preview_quality"), dict) else None
            # pim_import_sessions removed — progress is tracked in file_processing_status + raw_report_uploads.metadata

            if not done:
                chunk_total_est = max(1, (preview_row_goal + chunk_cap - 1) // chunk_cap) if preview_row_goal > 0 else None
                chunk_idx = (max(0, data_row_start) // max(chunk_cap, 1)) + 1 if chunk_cap else 1
                msg = (
                    f"Importing rows… processed {next_a:,}"
                    + (f" of ~{preview_row_goal:,} accepted" if preview_row_goal > 0 else "")
                    + (f" (chunk ~{chunk_idx}" + (f"/~{chunk_total_est}" if chunk_total_est else "") + ")")
                )
                return {
                    "ok": True,
                    "done": False,
                    "metrics_partial": metrics,
                    "progress_pct": pct,
                    "stage_label": "importing",
                    "apply_row_cursor": next_a,
                    "apply_row_goal": preview_row_goal or None,
                    "apply_chunk_cap": chunk_cap,
                    "apply_chunk_index": chunk_idx,
                    "apply_chunks_estimate": chunk_total_est,
                    "user_message": msg,
                    "import_job_id": str(upload_id),
                }

            quality_snapshot = (
                job_mid.get("preview_quality") if isinstance(job_mid.get("preview_quality"), dict) else {}
            )
            metrics["reconciliation"] = m._pim_reconcile_apply_vs_preview(metrics, quality_snapshot)
            seed_sid = meta.get("pim_seed_session_id")
            m._pim_seed_history_finalize_apply(
                db,
                str(seed_sid) if seed_sid else None,
                organization_id,
                store_id,
                success=True,
                metrics=metrics,
                quality=quality_snapshot,
                error_message=None,
            )
            meta = merge_pim_job(
                meta,
                {
                    "lifecycle": "completed",
                    "progress_pct": 100,
                    "stage_label": "completed",
                    "apply_cursor": None,
                },
            )
            _persist_meta(db, organization_id, upload_id, meta, "complete")
            return {"ok": True, "done": True, "metrics": metrics}
        except Exception as apply_exc:
            log.exception("pim_import apply-step fatal job_id=%s", upload_id)
            err_txt = str(apply_exc)
            # Detect connection-termination errors — use a much smaller chunk on retry
            is_connection_error = any(
                kw in err_txt.lower()
                for kw in ("connectionterminated", "connection terminated", "econnreset",
                           "connection reset", "broken pipe", "connection lost", "stream error")
            )
            try:
                proc_n = int(processed)
            except Exception:
                proc_n = 0
            try:
                resume_at = int(data_row_start) + proc_n
            except Exception:
                try:
                    resume_at = int(data_row_start)
                except Exception:
                    resume_at = 0
            prev_pct = 0
            try:
                prev_pct = int(
                    (meta.get("pim_import_job") or {}).get("progress_pct")
                    or job.get("progress_pct")
                    or 0
                )
            except Exception:
                prev_pct = 0
            md = metrics if isinstance(metrics, dict) else {}
            partial_written = proc_n > 0 or any(
                int(md.get(k) or 0) > 0
                for k in (
                    "products_created",
                    "products_updated",
                    "identifiers_created",
                    "identifiers_updated",
                    "prices_inserted",
                    "vendors_created",
                    "categories_created",
                    "rows_processed",
                )
            )
            fail_stage = "import_partial_failed" if partial_written else "apply_failed"
            preview_stat = "import_partial_failed" if partial_written else "failed"

            # Reduce chunk size for next retry — connection errors need smaller chunks
            if is_connection_error:
                retry_chunk = max(ROW_CHUNK_RETRY_MIN, chunk_cap // 4)
                log.info(
                    "pim apply ConnectionTerminated job_id=%s chunk_cap=%d → retry_chunk=%d resume_at=%d",
                    upload_id, chunk_cap, retry_chunk, resume_at,
                )
            else:
                retry_chunk = max(ROW_CHUNK_RETRY_MIN, chunk_cap // 2)

            meta = merge_pim_job(
                meta,
                {
                    "lifecycle": "failed",
                    "last_error": err_txt[:2000],
                    "stage_label": fail_stage,
                    "progress_pct": prev_pct,
                    "apply_cursor": {"data_row": resume_at},
                    "apply_metrics": metrics if isinstance(metrics, dict) else job.get("apply_metrics"),
                    "apply_effective_row_chunk": retry_chunk,
                },
            )
            meta["preview_status"] = preview_stat
            meta["import_job_status"] = "failed"
            _persist_meta(db, organization_id, upload_id, meta, "failed")
            um = (
                f"Import paused after writing catalog rows (checkpoint row offset {resume_at}). "
                f"Resume will continue from row {resume_at} without duplicating committed rows."
                + (f" (Retry chunk reduced to {retry_chunk} rows)" if is_connection_error else "")
                + f" Error: {err_txt[:180]}"
                if partial_written
                else f"Import failed at spreadsheet row offset {resume_at} (resume checkpoint saved): {err_txt[:240]}"
            )
            return {
                "ok": False,
                "done": True,
                "terminal": True,
                "error": "apply_failed",
                "last_error": err_txt[:2000],
                "stage_label": fail_stage,
                "apply_failed_at_row": resume_at,
                "apply_retry_chunk": retry_chunk,
                "retryable": True,
                "metrics_partial": metrics if isinstance(metrics, dict) else {},
                "user_message": um,
            }
    finally:
        try:
            os.unlink(local_path)
        except Exception:
            pass
