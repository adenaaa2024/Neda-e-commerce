import csv
import io
import json
import logging
import math
import time
import traceback
import uuid
import hashlib
import os
import re
from collections import defaultdict, deque
from typing import Any, Callable, Iterator, List

import pandas as pd
from dotenv import load_dotenv
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from supabase import Client, create_client

from pim_import_async import (
    get_pim_import_preview_status,
    run_pim_import_apply_step,
    run_pim_import_price_backfill_step,
    run_pim_import_preview_step,
    run_pim_import_retry_preview,
)

# Load environment variables securely from .env file
load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
log = logging.getLogger("etl")

# Initialize FastAPI App
app = FastAPI(title="Logistics AI Agent API", version="1.0")

# --- CORS Middleware ---
# NOTE: allow_credentials=True + allow_origins=["*"] is invalid per the spec and
# makes browsers reject preflight. JWT cookies are not used on these ETL routes.
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Background Task Progress Store ---
task_store: dict[str, dict[str, Any]] = {}

def _update_task(task_id: str, progress: int, message: str, status: str = "running") -> None:
    task_store[task_id] = {
        "task_id": task_id,
        "status": status,
        "progress": progress,
        "message": message,
    }

# Fetch keys from environment
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")

# Initialize Database Connection
supabase: Client | None = None
try:
    if not SUPABASE_URL or not SUPABASE_KEY:
        print("Error: Missing Supabase credentials in .env file!")
    else:
        supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
        print("Connected to Supabase Successfully!")
except Exception as e:
    print(f"Database Connection Error: {e}")

def _require_supabase() -> Client:
    if supabase is None:
        raise HTTPException(
            status_code=503,
            detail="Database not configured. Set SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY.",
        )
    return supabase


# Align with Next.js Phase 2 staging: transient PostgREST / network hiccups (e.g. Nano).
_IMPORT_RETRY_BACKOFF_SEC = (5.0, 15.0, 30.0, 60.0, 120.0)
_IMPORT_OP_MAX_ATTEMPTS = 6


def _is_transient_supabase_error(exc: BaseException) -> bool:
    msg = str(exc).lower()
    needles = (
        "timeout",
        "timed out",
        "statement timeout",
        "deadlock",
        "too many connections",
        "connection reset",
        "econnreset",
        "etimedout",
        "eai_again",
        "fetch failed",
        "network",
        "socket",
        "502",
        "503",
        "504",
        "bad gateway",
        "gateway timeout",
        "service unavailable",
        "internal server error",
        "premature",
        "cloudflare",
    )
    return any(n in msg for n in needles)


def _execute_with_import_retries(label: str, fn: Callable[[], Any], max_attempts: int = _IMPORT_OP_MAX_ATTEMPTS) -> Any:
    last_exc: BaseException | None = None
    for attempt in range(1, max_attempts + 1):
        try:
            return fn()
        except Exception as e:
            last_exc = e
            if not _is_transient_supabase_error(e) or attempt >= max_attempts:
                raise
            idx = min(attempt - 1, len(_IMPORT_RETRY_BACKOFF_SEC) - 1)
            delay = _IMPORT_RETRY_BACKOFF_SEC[idx]
            log.warning(
                "[etl] %s attempt %d/%d failed (%s); sleeping %.1fs before retry",
                label,
                attempt,
                max_attempts,
                e,
                delay,
            )
            time.sleep(delay)
    if last_exc is not None:
        raise last_exc
    raise RuntimeError(f"{label}: retries exhausted with no exception (should not happen)")


def _normalize_header(name: str) -> str:
    return (
        name.replace("\ufeff", "")
        .strip()
        .replace("-", "_")
        .replace(" ", "_")
        .lower()
    )

def _cell_to_json_safe(val: Any) -> Any:
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except TypeError:
        pass
    if isinstance(val, str):
        s = val.strip()
        return s if s else None
    if hasattr(val, "item"):
        try:
            return val.item()
        except Exception:
            pass
    return val

def _safe_int(val: Any) -> int | None:
    if val is None:
        return None
    s = str(val).strip().replace(",", "")
    if not s:
        return None
    try:
        return int(float(s))
    except (TypeError, ValueError):
        return None

def _safe_float(val: Any) -> float | None:
    if val is None:
        return None
    s = str(val).strip().replace(",", "").replace("$", "")
    if not s:
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        return None

def _safe_bool(val: Any) -> bool | None:
    if val is None:
        return None
    s = str(val).strip().lower()
    if s in {"true", "yes", "y", "1"}:
        return True
    if s in {"false", "no", "n", "0"}:
        return False
    return None

def _safe_timestamp(val: Any) -> str | None:
    if val is None:
        return None
    try:
        ts = pd.to_datetime(val, errors="coerce")
        if pd.isna(ts):
            return None
        return ts.isoformat()
    except Exception:
        return None

def _safe_date(val: Any) -> str | None:
    ts = _safe_timestamp(val)
    if not ts:
        return None
    return ts[:10]

def _read_tabular_file(content: bytes) -> pd.DataFrame:
    buf = io.BytesIO(content)
    _csv_kw: dict[str, Any] = {
        "sep": None,
        "engine": "python",
        "encoding": "utf-8-sig",
        "dtype": object,
        "keep_default_na": False,
    }
    try:
        return pd.read_csv(buf, **_csv_kw)
    except Exception:
        buf.seek(0)
        return pd.read_csv(
            buf, sep="\t", encoding="utf-8-sig", dtype=object, keep_default_na=False,
        )


def _read_text_delimited(content: bytes, fname: str | None) -> tuple[pd.DataFrame, str, bool]:
    """
    Read .csv / .txt with delimiter sniffing (comma, tab, semicolon, pipe, or pandas auto).
    Returns (dataframe, human_label, uncertain_single_column).
    """
    best_cols = -1
    best_rows = -1
    best_label = "comma"
    best_df: pd.DataFrame | None = None
    candidates: list[tuple[str | None, str]] = [
        (None, "auto-detected"),
        ("\t", "tab"),
        (",", "comma"),
        (";", "semicolon"),
        ("|", "pipe"),
    ]
    for sep, label in candidates:
        try:
            buf = io.BytesIO(content)
            df = pd.read_csv(
                buf,
                sep=sep,
                engine="python",
                encoding="utf-8-sig",
                dtype=object,
                keep_default_na=False,
            )
            ncols = int(df.shape[1])
            nrows = int(df.shape[0])
            if ncols > best_cols or (ncols == best_cols and nrows > best_rows):
                best_cols, best_rows, best_label, best_df = ncols, nrows, label, df
        except Exception:
            continue
    if best_df is None:
        return _read_tabular_file(content), "fallback", True
    uncertain = best_cols < 2 and len(content) > 80
    return best_df, best_label, uncertain


def _looks_like_xlsx_bytes(content: bytes, filename: str | None) -> bool:
    fn = (filename or "").strip().lower()
    if fn.endswith((".xlsx", ".xlsm")):
        return True
    return len(content) >= 4 and content[:4] == b"PK\x03\x04"


def _iter_seed_product_frames(content: bytes, filename: str | None) -> Iterator[tuple[str | None, pd.DataFrame]]:
    """
    Yield (sheet_name_or_none, dataframe) for seed-products.
    CSV/TSV: one frame (no tab name). Excel (.xlsx): only the tab named App_Import_Product_Master.
    """
    if not _looks_like_xlsx_bytes(content, filename):
        df, delim, uncertain = _read_text_delimited(content, filename)
        if hasattr(df, "attrs"):
            df.attrs["pim_delimiter_detected"] = delim
            df.attrs["pim_delimiter_uncertain"] = uncertain
        yield None, df
        return
    try:
        import openpyxl  # noqa: F401 — pandas read_excel(engine="openpyxl") needs this package at runtime
    except ModuleNotFoundError as e:
        raise HTTPException(
            status_code=503,
            detail=(
                "Excel (.xlsx/.xlsm) requires the 'openpyxl' package on the ETL server. "
                "Install: pip install openpyxl   "
                "or from repo root: pip install -r backend-python/requirements.txt"
            ),
        ) from e
    bio = io.BytesIO(content)
    try:
        xl = pd.ExcelFile(bio, engine="openpyxl")
    except ImportError as e:
        raise HTTPException(
            status_code=503,
            detail=(
                "Could not load Excel engine. Install openpyxl: pip install openpyxl "
                "(see backend-python/requirements.txt)."
            ),
        ) from e
    saw_master = False
    for sheet in xl.sheet_names:
        if sheet != PRODUCT_MASTER_SHEET_NAME:
            continue
        df = pd.read_excel(
            xl,
            sheet_name=sheet,
            header=0,
            dtype=object,
            keep_default_na=False,
        )
        headers = [str(c).strip() for c in df.columns if str(c).strip() != ""]
        if df.shape[0] == 0 or not headers:
            continue
        saw_master = True
        yield sheet, df
    if not saw_master:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Excel workbook must contain a sheet named {PRODUCT_MASTER_SHEET_NAME!r} "
                "with a header row and at least one data row. Other tabs are ignored."
            ),
        )


def _sniff_csv_delimiter(first_line: str) -> str:
    if "\t" in first_line and first_line.count("\t") >= max(1, first_line.count(",")):
        return "\t"
    return ","


def _csv_headers_and_row_count_fast(content: bytes) -> tuple[list[str], int]:
    """
    One streaming pass over bytes — header cells + data row count.
    Used by upload-raw only; avoids pandas (huge RAM + double parse) on large files.
    """
    if not content:
        return [], 0
    sample = content[: min(len(content), 262144)].decode("utf-8-sig", errors="replace")
    lines = sample.splitlines()
    if not lines:
        return [], 0
    delim = _sniff_csv_delimiter(lines[0])
    bio = io.BytesIO(content)
    text_io = io.TextIOWrapper(bio, encoding="utf-8-sig", errors="replace", newline="")
    try:
        reader = csv.reader(text_io, delimiter=delim)
        rows_iter = iter(reader)
        header_row = next(rows_iter, [])
        headers = [str(h).strip() for h in header_row if str(h).strip() != ""]
        if not headers:
            return [], 0
        row_count = sum(1 for _ in rows_iter)
        return headers, row_count
    finally:
        text_io.close()


def _get_from_row(row: dict[str, Any], *candidates: str) -> str | None:
    norm_map = {_normalize_header(k): k for k in row.keys()}
    for cand in candidates:
        key = _normalize_header(cand)
        if key in norm_map:
            v = row.get(norm_map[key])
            s = _cell_to_json_safe(v)
            if s is not None and str(s).strip() != "":
                return str(s).strip()
    return None

def _parse_quantity(row: dict[str, Any]) -> float:
    raw = _get_from_row(
        row,
        "quantity",
        "shipped_quantity",
        "requested_quantity",
        "disposed_quantity",
        "cancelled_quantity",
        "Quantity",
    )
    if raw is None:
        return 0.0
    try:
        return float(str(raw).replace(",", ""))
    except ValueError:
        return 0.0

def _group_key_for_row(row: dict[str, Any]) -> str | None:
    track = _get_from_row(
        row,
        "tracking_number",
        "tracking-number",
        "tracking_id",
        "tracking-id",
    )
    if track:
        return f"tn:{track}"
    oid = _get_from_row(
        row,
        "order_id",
        "order-id",
        "removal_order_id",
        "removal-order-id",
        "amazon_order_id",
    )
    if oid:
        return f"oid:{oid}"
    return None

# --- Removal ETL: field mapping and extraction ---
REMOVAL_FIELD_MAP: dict[str, list[str]] = {
    "order_id":           [
        "order-id",
        "order_id",
        "removal_order_id",
        "removal-order-id",
        "amazon_order_id",
        "amazon-order-id",
    ],
    "order_source":       ["order-source", "order_source", "order source"],
    "order_type":         ["order-type", "order_type", "order type"],
    "order_status":       ["order-status", "order_status"],
    "sku":                ["sku", "merchant_sku", "merchant-sku"],
    "fnsku":              ["fnsku"],
    "disposition":        ["disposition", "fc-disposition", "detailed-disposition"],
    "shipped_quantity":   ["shipped-quantity", "shipped_quantity"],
    "requested_quantity": ["requested-quantity", "requested_quantity", "quantity"],
    "cancelled_quantity": ["cancelled-quantity", "cancelled_quantity"],
    "disposed_quantity":  ["disposed-quantity", "disposed_quantity"],
    "in_process_quantity": ["in-process-quantity", "in_process_quantity", "in process quantity"],
    "tracking_number":    ["tracking-number", "tracking_number"],
    "carrier":            ["carrier", "carrier-name", "carrier_name"],
    "shipment_date":      ["carrier-shipment-date", "shipment-date", "ship-date", "shipped-date"],
    "order_date":         ["request-date", "order-date", "order_date", "request_date"],
    "last_updated_date":  ["last-updated-date", "last_updated_date", "last updated date"],
    "removal_fee":        ["removal-fee", "removal_fee", "removal fee"],
    "currency":           ["currency"],
}

_INT_REMOVAL_COLS = {
    "shipped_quantity",
    "requested_quantity",
    "cancelled_quantity",
    "disposed_quantity",
    "in_process_quantity",
}

def _resolve_import_store_id_from_upload(db: Any, organization_id: str, upload_id: str | None) -> str | None:
    if not upload_id:
        return None
    try:
        res = (
            db.table("raw_report_uploads")
            .select("metadata")
            .eq("id", upload_id)
            .eq("organization_id", organization_id)
            .limit(1)
            .execute()
        )
        row = (res.data or [None])[0] if isinstance(res.data, list) else res.data
        meta = row.get("metadata") if isinstance(row, dict) else None
        if not isinstance(meta, dict):
            return None
        raw = meta.get("import_store_id") or meta.get("ledger_store_id")
        if not raw:
            return None
        sid = str(raw).strip()
        uuid.UUID(sid)
        return sid
    except Exception:
        return None

def _extract_removal_row(
    raw: dict[str, Any],
    organization_id: str,
    upload_id: str | None,
) -> dict[str, Any]:
    result: dict[str, Any] = {"organization_id": organization_id}
    if upload_id:
        result["upload_id"] = upload_id

    for db_col, candidates in REMOVAL_FIELD_MAP.items():
        val = _get_from_row(raw, *candidates)
        if db_col in _INT_REMOVAL_COLS:
            result[db_col] = _safe_int(val) if val is not None else None
        elif db_col in ("order_date", "last_updated_date"):
            if val is not None:
                try:
                    d = pd.to_datetime(val, errors="coerce")
                    result[db_col] = d.date().isoformat() if pd.notna(d) else None
                except Exception:
                    result[db_col] = None
            else:
                result[db_col] = None
        elif db_col == "removal_fee":
            result[db_col] = _safe_float(val) if val is not None else None
        else:
            result[db_col] = val

    if not _pg_text_unique_field(result.get("sku")):
        fn = _get_from_row(raw, "fnsku", "asin")
        if fn:
            result["sku"] = str(fn).strip()

    for nullable in ("sku", "fnsku", "disposition", "tracking_number"):
        t = _pg_text_unique_field(result.get(nullable))
        result[nullable] = t

    return result

_REMOVALS_CONFLICT = "organization_id,store_id,order_id,sku,fnsku,disposition,requested_quantity,shipped_quantity,disposed_quantity,cancelled_quantity,order_date,order_type"
_REMOVAL_SHIPMENT_CONFLICT = (
    "organization_id,store_id,order_id,tracking_number,sku,fnsku,disposition,"
    "requested_quantity,shipped_quantity,disposed_quantity,cancelled_quantity,order_date,order_type"
)
_REMOVAL_WRITE_COLS = frozenset(REMOVAL_FIELD_MAP.keys()) | {
    "organization_id",
    "upload_id",
    "raw_data",
    "source_staging_id",
    "store_id",
}

_DB_SYSTEM_COLS = frozenset({"id", "created_at", "updated_at"})

def _removal_order_date_key(val: Any) -> str | None:
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except Exception:
        pass
    if hasattr(val, "isoformat"):
        try:
            return str(val.isoformat())[:10]
        except Exception:
            pass
    s = str(val).strip()
    if not s:
        return None
    return s[:10]

def _removal_null_slot_match_key(r: dict[str, Any]) -> tuple[Any, ...]:
    sid = r.get("store_id")
    store_part = str(sid).strip() if sid is not None and str(sid).strip() != "" else ""
    return (
        str(r.get("organization_id") or "").strip(),
        store_part,
        _pg_text_unique_field(r.get("order_id")),
        _pg_text_unique_field(r.get("order_type")),
        _pg_text_unique_field(r.get("sku")),
        _pg_text_unique_field(r.get("fnsku")),
        _pg_text_unique_field(r.get("disposition")),
    )

def _removal_row_for_write(row: dict[str, Any]) -> dict[str, Any]:
    return {k: row[k] for k in _REMOVAL_WRITE_COLS if k in row}

def _merge_shipment_into_null_slot(existing: dict[str, Any], shipment: dict[str, Any]) -> dict[str, Any]:
    out = dict(existing)
    tn = _pg_text_unique_field(shipment.get("tracking_number"))
    if tn:
        out["tracking_number"] = tn
    inc_c = shipment.get("carrier")
    if inc_c is not None and str(inc_c).strip() != "":
        if not _pg_text_unique_field(existing.get("carrier")):
            out["carrier"] = inc_c
    inc_sd = shipment.get("shipment_date")
    if inc_sd is not None and str(inc_sd).strip() != "":
        if existing.get("shipment_date") is None or str(existing.get("shipment_date")).strip() == "":
            out["shipment_date"] = inc_sd
    return out

EXPECTED_PKG_AMAZON_COLS = frozenset({
    "organization_id",
    "store_id",
    "upload_id",
    "source_staging_id",
    "order_type",
    "order_id",
    "sku",
    "fnsku",
    "tracking_number",
    "shipped_quantity",
    "requested_quantity",
    "disposed_quantity",
    "cancelled_quantity",
    "order_status",
    "disposition",
    "order_date",
    "carrier",
    "shipment_date",
})

_EP_WORKLIST_AMAZON_FILL_NULL = frozenset({
    "tracking_number",
    "carrier",
    "shipment_date",
    "order_date",
    "fnsku",
    "store_id",
    "order_type",
})

_EXPECTED_PKG_CONFLICT = (
    "organization_id,store_id,order_id,order_type,sku,fnsku,disposition"
)

def _ep_value_absent(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and str(v).strip() == "":
        return True
    return False

def _expected_pkg_identity_key(row: dict[str, Any]) -> tuple[Any, ...]:
    st = row.get("store_id")
    sid = str(st).strip() if st is not None and str(st).strip() != "" else ""
    return (
        str(row.get("organization_id", "")),
        sid,
        _pg_text_unique_field(row.get("order_id")),
        _pg_text_unique_field(row.get("order_type")),
        _pg_text_unique_field(row.get("sku")),
        _pg_text_unique_field(row.get("fnsku")),
        _pg_text_unique_field(row.get("disposition")),
    )

_EP_UPSERT_COLLAPSE_PREFER_FIELDS = frozenset({
    "tracking_number",
    "carrier",
    "shipment_date",
    "order_type",
    "fnsku",
})

def _dedupe_merged_rows_for_expected_packages_upsert(
    merged_rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], int]:
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    key_order: list[tuple[Any, ...]] = []
    for r in merged_rows:
        k = _expected_pkg_identity_key(r)
        if k not in groups:
            key_order.append(k)
            groups[k] = []
        groups[k].append(r)

    out: list[dict[str, Any]] = []
    collapsed = 0
    for k in key_order:
        group = groups[k]
        if len(group) == 1:
            out.append(group[0])
            continue
        collapsed += len(group) - 1
        merged = dict(group[0])
        for other in group[1:]:
            for fk in _EP_UPSERT_COLLAPSE_PREFER_FIELDS:
                if fk not in merged or _ep_value_absent(merged.get(fk)):
                    v = other.get(fk)
                    if not _ep_value_absent(v):
                        merged[fk] = v
        out.append(merged)

    return out, collapsed

def _merge_expected_pkg_amazon_only(existing: dict[str, Any], incoming_amazon: dict[str, Any]) -> dict[str, Any]:
    out = dict(existing)
    for k, v in incoming_amazon.items():
        if k not in EXPECTED_PKG_AMAZON_COLS:
            continue
        if k in _EP_WORKLIST_AMAZON_FILL_NULL:
            if _ep_value_absent(v):
                continue
            out[k] = v
        else:
            out[k] = v
    return out

def _is_return_order_type_for_worklist(row: dict[str, Any]) -> bool:
    t = _pg_text_unique_field(row.get("order_type"))
    if not t:
        return True
    return t.strip().lower() == "return"

class AmazonOrderSync(BaseModel):
    amazon_order_id: str
    org_id: str
    store_id: str
    raw_data: dict

class SyncRemovalsRequest(BaseModel):
    organization_id: str
    upload_id: str | None = None

class GenerateWorklistRequest(BaseModel):
    organization_id: str
    upload_id: str | None = None

@app.get("/")
def read_root():
    return {"status": "Agent Backend is Live!", "service": "AI Logistics"}

@app.get("/agent/pending-claims")
async def get_pending_claims(
    organization_id: str = Query(..., description="Tenant UUID (required)"),
    store_id: str | None = Query(None, description="Optional store UUID — only submissions for this store"),
):
    db = _require_supabase()
    try:
        try:
            org_uuid = str(uuid.UUID(organization_id.strip()))
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="Invalid organization_id") from exc
        store_uuid: str | None = None
        if store_id and str(store_id).strip():
            try:
                store_uuid = str(uuid.UUID(str(store_id).strip()))
            except ValueError as exc:
                raise HTTPException(status_code=400, detail="Invalid store_id") from exc

        q = (
            db.table("claim_submissions")
            .select("*")
            .eq("status", "ready_to_send")
            .eq("organization_id", org_uuid)
        )
        if store_uuid:
            q = q.eq("store_id", store_uuid)
        response = q.execute()
        return {"count": len(response.data), "claims": response.data}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/sync/order")
async def save_raw_amazon_order(order: AmazonOrderSync):
    db = _require_supabase()
    try:
        data = {
            "organization_id": order.org_id,
            "store_id": order.store_id,
            "amazon_order_id": order.amazon_order_id,
            "raw_data": order.raw_data,
            "status": "synced",
        }
        db.table("marketplace_orders").upsert(data).execute()
        return {"status": "success", "message": "Order synced to Landing Zone"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

STAGING_INSERT_BATCH = 500
WORKLIST_FETCH_PAGE = 1000

def _fetch_org_table_all(
    db: Any,
    table: str,
    organization_id: str,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    offset = 0
    while True:
        res = (
            db.table(table)
            .select("*")
            .eq("organization_id", organization_id)
            .order("id")
            .range(offset, offset + WORKLIST_FETCH_PAGE - 1)
            .execute()
        )
        chunk = res.data or []
        out.extend(chunk)
        if len(chunk) < WORKLIST_FETCH_PAGE:
            break
        offset += WORKLIST_FETCH_PAGE
    return out

def _fetch_amazon_removal_shipments_by_upload(
    db: Any,
    organization_id: str,
    upload_id: str,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    offset = 0
    while True:
        res = (
            db.table("amazon_removal_shipments")
            .select("*")
            .eq("organization_id", organization_id)
            .eq("upload_id", upload_id)
            .order("id")
            .range(offset, offset + WORKLIST_FETCH_PAGE - 1)
            .execute()
        )
        chunk = res.data or []
        out.extend(chunk)
        if len(chunk) < WORKLIST_FETCH_PAGE:
            break
        offset += WORKLIST_FETCH_PAGE
    return out

def _removal_like_from_shipment_archive(sh: dict[str, Any]) -> dict[str, Any]:
    row = dict(sh)
    sid = sh.get("amazon_staging_id")
    if sid is not None and str(sid).strip() != "":
        row["source_staging_id"] = sid
    return row

@app.post("/etl/upload-removal")
async def etl_upload_removal(
    file: UploadFile = File(...),
    organization_id: str = Form(...),
    batch_id: str | None = Form(None),
):
    db = _require_supabase()
    try:
        uuid.UUID(organization_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="organization_id must be a valid UUID.")

    try:
        raw_bytes = await file.read()
        if not raw_bytes:
            raise HTTPException(status_code=400, detail="Empty file.")

        batch = batch_id.strip() if batch_id and batch_id.strip() else str(uuid.uuid4())
        try:
            uuid.UUID(batch)
        except ValueError:
            raise HTTPException(
                status_code=400,
                detail="batch_id must be a valid UUID when provided.",
            )

        df = _read_tabular_file(raw_bytes)
        df.columns = [_normalize_header(str(c)) for c in df.columns]
        df = df.dropna(axis=1, how="all")

        if df.empty:
            raise HTTPException(status_code=400, detail="No data rows in file.")

        rows_to_insert: list[dict[str, Any]] = []
        for _, series in df.iterrows():
            row_dict: dict[str, Any] = {}
            for col in df.columns:
                row_dict[col] = _cell_to_json_safe(series[col])

            payload = {"batch_id": batch, **row_dict}

            snapshot_date = None
            for date_key in (
                "date",
                "date_time",
                "posted_date",
                "order_date",
                "request_date",
                "snapshot_date",
            ):
                if date_key in payload and payload[date_key]:
                    try:
                        d = pd.to_datetime(payload[date_key], errors="coerce")
                        if pd.notna(d):
                            snapshot_date = d.date().isoformat()
                            break
                    except Exception:
                        continue

            row_insert: dict[str, Any] = {
                "organization_id": organization_id,
                "raw_row": payload,
                "batch_id": batch,
            }
            if snapshot_date:
                row_insert["snapshot_date"] = snapshot_date

            rows_to_insert.append(row_insert)

        inserted = 0
        for i in range(0, len(rows_to_insert), STAGING_INSERT_BATCH):
            chunk = rows_to_insert[i : i + STAGING_INSERT_BATCH]
            _execute_with_import_retries(
                f"amazon_staging insert offset={i} size={len(chunk)}",
                lambda c=chunk: db.table("amazon_staging").insert(c).execute(),
            )
            inserted += len(chunk)

        return {
            "status": "success",
            "batch_id": batch,
            "rows_inserted": inserted,
            "message": "Staging load completed.",
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"ETL upload failed: {e!s}",
        )

@app.post("/etl/process-staging")
async def etl_process_staging(organization_id: str = Form(...)):
    db = _require_supabase()
    try:
        uuid.UUID(organization_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="organization_id must be a valid UUID.")

    try:
        res = (
            db.table("amazon_staging")
            .select("id,organization_id,raw_row")
            .eq("organization_id", organization_id)
            .execute()
        )
        staging_rows = res.data or []
        if not staging_rows:
            return {
                "status": "success",
                "message": "No staging rows to process.",
                "pallets_created": 0,
                "items_created": 0,
                "staging_cleared": False,
            }

        groups: dict[str, list[dict[str, Any]]] = {}
        skipped = 0
        for row in staging_rows:
            raw = row.get("raw_row")
            if isinstance(raw, str):
                try:
                    raw = json.loads(raw)
                except json.JSONDecodeError:
                    skipped += 1
                    continue
            if not isinstance(raw, dict):
                skipped += 1
                continue

            gkey = _group_key_for_row(raw)
            if not gkey:
                skipped += 1
                continue
            groups.setdefault(gkey, []).append(raw)

        if not groups:
            raise HTTPException(
                status_code=400,
                detail="Could not derive tracking_number or order_id for any row; nothing to process.",
            )

        pallet_payloads: list[dict[str, Any]] = []
        group_order: list[str] = []

        for gkey, raws in groups.items():
            group_order.append(gkey)
            first = raws[0]
            track = _get_from_row(
                first,
                "tracking_number",
                "tracking-number",
                "tracking_id",
                "tracking-id",
            )
            oid = _get_from_row(
                first,
                "order_id",
                "order-id",
                "removal_order_id",
                "removal-order-id",
                "amazon_order_id",
            )
            batch = _get_from_row(first, "batch_id")

            pallet_payloads.append(
                {
                    "organization_id": organization_id,
                    "status": "Pending",
                    "tracking_number": track,
                    "order_id": oid,
                    "batch_id": batch if batch else None,
                }
            )

        pal_res = _execute_with_import_retries(
            "expected_pallets bulk insert",
            lambda: db.table("expected_pallets").insert(pallet_payloads).select("id").execute(),
        )
        pallet_ids = [r["id"] for r in (pal_res.data or [])]

        if len(pallet_ids) != len(group_order):
            raise HTTPException(
                status_code=500,
                detail="Insert into expected_pallets did not return all IDs.",
            )

        gkey_to_pallet = dict(zip(group_order, pallet_ids, strict=True))

        item_rows: list[dict[str, Any]] = []
        for gkey, raws in groups.items():
            pid = gkey_to_pallet[gkey]
            sku_qty: dict[str, float] = {}
            for raw in raws:
                sku = _get_from_row(raw, "sku", "fnsku", "asin")
                if not sku:
                    continue
                qty = _parse_quantity(raw)
                sku_qty[sku] = sku_qty.get(sku, 0.0) + qty

            for sku, qty in sku_qty.items():
                item_rows.append(
                    {
                        "expected_pallet_id": pid,
                        "sku": sku,
                        "quantity": qty,
                    }
                )

        items_created = 0
        if item_rows:
            for i in range(0, len(item_rows), STAGING_INSERT_BATCH):
                chunk = item_rows[i : i + STAGING_INSERT_BATCH]
                _execute_with_import_retries(
                    f"expected_items insert offset={i} size={len(chunk)}",
                    lambda c=chunk: db.table("expected_items").insert(c).execute(),
                )
                items_created += len(chunk)

        _execute_with_import_retries(
            "amazon_staging delete by organization_id",
            lambda: db.table("amazon_staging").delete().eq("organization_id", organization_id).execute(),
        )

        return {
            "status": "success",
            "pallets_created": len(pallet_ids),
            "items_created": items_created,
            "rows_skipped_no_key": skipped,
            "staging_cleared": True,
            "message": "Staging processed and cleared.",
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"ETL process failed: {e!s}",
        )

def _run_sync_removals(
    task_id: str,
    organization_id: str,
    upload_id: str | None,
) -> None:
    print(f"[sync-removals] START task_id={task_id} org={organization_id} upload_id={upload_id}")
    log.info(
        "[sync-removals] START task_id=%s org=%s upload_id=%s",
        task_id, organization_id, upload_id,
    )
    _update_task(task_id, 5, "Task is alive. Connecting to database...")

    _url = os.getenv("SUPABASE_URL")
    _key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    if not _url or not _key:
        log.error("[sync-removals] ERROR: Missing Supabase credentials")
        _update_task(task_id, 0, "Missing Supabase credentials in environment.", status="failed")
        return

    try:
        db = create_client(_url, _key)
        print("[sync-removals] Supabase client created successfully")
    except Exception as conn_err:
        log.error("[sync-removals] CONNECTION ERROR: %s", conn_err)
        _update_task(task_id, 0, f"Failed to connect to database: {conn_err!s}", status="failed")
        return

    store_id = _resolve_import_store_id_from_upload(db, organization_id, upload_id)
    if not store_id:
        log.error("[sync-removals] Missing Imports Target Store on raw_report_uploads metadata")
        _update_task(
            task_id,
            0,
            "Imports Target Store missing — set import_store_id or ledger_store_id on the upload metadata.",
            status="failed",
        )
        return
    log.info("[sync-removals] Using store_id=%s", store_id)

    enriched_count = 0
    unmatched_shipment_rows: list[dict[str, Any]] = []
    no_tracking_rows: list[dict[str, Any]] = []

    try:
        _update_task(task_id, 10, "Fetching rows from amazon_staging...")
        query = db.table("amazon_staging").select("*").eq("organization_id", organization_id)
        if upload_id:
            query = query.eq("upload_id", upload_id)
        staging_rows: list[dict[str, Any]] = (query.execute().data or [])
        log.info("[sync-removals] Fetched %d rows from amazon_staging", len(staging_rows))

        if not staging_rows:
            _update_task(task_id, 100, "No staging rows found. Nothing to sync.", status="completed")
            return

        _update_task(task_id, 18, f"Extracting {len(staging_rows)} staging rows...")

        packed: list[tuple[str, dict[str, Any], dict[str, Any]]] = []
        staging_ids: list[str] = []
        skipped = 0
        skipped_no_order_id = 0

        for row in staging_rows:
            raw = row.get("raw_row", {})
            if isinstance(raw, str):
                try:
                    raw = json.loads(raw)
                except json.JSONDecodeError:
                    skipped += 1
                    continue
            if not isinstance(raw, dict):
                skipped += 1
                continue

            row_upload_id: str | None = row.get("upload_id") or upload_id
            extracted = _extract_removal_row(raw, organization_id, row_upload_id)

            if not _pg_text_unique_field(extracted.get("order_id")):
                skipped += 1
                skipped_no_order_id += 1
                continue

            staging_uuid = str(row["id"])
            extracted["source_staging_id"] = staging_uuid
            extracted["store_id"] = store_id

            packed.append((staging_uuid, raw, extracted))
            staging_ids.append(str(row["id"]))

        log.info("[sync-removals] Extracted %d valid rows; skipped %d", len(packed), skipped)
        if skipped_no_order_id > 0:
            log.warning(
                "[sync-removals] SKIPPED %d staging rows with no valid order_id — "
                "these rows remain in amazon_staging and are NOT written to amazon_removals.",
                skipped_no_order_id,
            )

        if not packed:
            _update_task(task_id, 100, "No valid removal rows extracted (missing order_id).", status="completed")
            return

        shipment_history_rows: list[dict[str, Any]] = []
        for staging_id, raw, ext in packed:
            row_ship: dict[str, Any] = {
                "organization_id": organization_id,
                "upload_id": ext.get("upload_id"),
                "amazon_staging_id": staging_id,
                "store_id": store_id,
                "raw_row": raw,
                "order_id": ext.get("order_id"),
                "sku": ext.get("sku"),
                "fnsku": ext.get("fnsku"),
                "disposition": ext.get("disposition"),
                "tracking_number": ext.get("tracking_number"),
                "carrier": ext.get("carrier"),
                "shipment_date": ext.get("shipment_date"),
                "order_date": ext.get("order_date"),
                "order_type": ext.get("order_type"),
                "requested_quantity": ext.get("requested_quantity"),
                "shipped_quantity": ext.get("shipped_quantity"),
                "disposed_quantity": ext.get("disposed_quantity"),
                "cancelled_quantity": ext.get("cancelled_quantity"),
            }
            shipment_history_rows.append(row_ship)
        if shipment_history_rows:
            _update_task(task_id, 20, f"Archiving {len(shipment_history_rows)} raw shipment rows...")
            for i in range(0, len(shipment_history_rows), STAGING_INSERT_BATCH):
                chunk = shipment_history_rows[i : i + STAGING_INSERT_BATCH]
                db.table("amazon_removal_shipments").upsert(
                    chunk,
                    on_conflict=_REMOVAL_SHIPMENT_CONFLICT,
                ).execute()
            log.info(
                "[sync-removals] Archived %d rows to amazon_removal_shipments",
                len(shipment_history_rows),
            )

        extracted_rows = [p[2] for p in packed]
        _update_task(task_id, 24, f"Prepared {len(extracted_rows)} removal line(s) (one per staging row)...")

        has_tracking_rows: list[dict[str, Any]] = []
        for r in extracted_rows:
            if _pg_text_unique_field(r.get("tracking_number")):
                has_tracking_rows.append(r)
            else:
                no_tracking_rows.append(r)

        log.info(
            "[sync-removals] Split: %d order rows (no tracking), %d rows with tracking",
            len(no_tracking_rows), len(has_tracking_rows),
        )

        upserted_total = 0

        if no_tracking_rows:
            n_order = len(no_tracking_rows)
            _update_task(task_id, 28, f"Upserting {n_order} order rows (no tracking)...")
            for i in range(0, n_order, STAGING_INSERT_BATCH):
                chunk = [_removal_row_for_write(r) for r in no_tracking_rows[i : i + STAGING_INSERT_BATCH]]
                db.table("amazon_removals").upsert(
                    chunk, on_conflict=_REMOVALS_CONFLICT,
                ).execute()
                upserted_total += len(chunk)
                progress = 28 + int((min(i + STAGING_INSERT_BATCH, n_order) / max(n_order, 1)) * 18)
                _update_task(task_id, progress, f"Upserted order rows… {min(i + STAGING_INSERT_BATCH, n_order)}/{n_order}")
            log.info("[sync-removals] Upserted %d order rows into amazon_removals", n_order)

        if has_tracking_rows:
            _update_task(
                task_id, 48,
                f"Matching {len(has_tracking_rows)} shipment rows to NULL-tracking removals...",
            )

            try:
                existing_null_rows: list[dict[str, Any]] = (
                    db.table("amazon_removals")
                    .select("*")
                    .eq("organization_id", organization_id)
                    .eq("store_id", store_id)
                    .is_("tracking_number", "null")
                    .execute()
                    .data or []
                )
            except Exception as fetch_err:
                log.warning(
                    "[sync-removals] Could not fetch NULL-tracking rows (%s) — "
                    "all shipment rows will be upserted as new entries.", fetch_err,
                )
                existing_null_rows = []

            log.info(
                "[sync-removals] Found %d NULL-tracking rows for enrichment",
                len(existing_null_rows),
            )

            null_match: dict[tuple[Any, ...], deque] = defaultdict(deque)
            null_rows_by_id: dict[str, dict[str, Any]] = {}

            for db_row in existing_null_rows:
                mk = _removal_null_slot_match_key(db_row)
                rid = str(db_row["id"])
                null_match[mk].append(rid)
                null_rows_by_id[rid] = db_row

            for s_row in has_tracking_rows:
                match_key = _removal_null_slot_match_key(s_row)

                if null_match[match_key]:
                    match_id = null_match[match_key].popleft()
                    existing = null_rows_by_id[match_id]
                    merged = _merge_shipment_into_null_slot(existing, s_row)
                    payload = _removal_row_for_write(merged)
                    db.table("amazon_removals").update(payload).eq("id", match_id).execute()
                    enriched_count += 1
                else:
                    unmatched_shipment_rows.append(s_row)

            log.info(
                "[sync-removals] Shipment handling: %d rows updated in place, %d upserted as new",
                enriched_count, len(unmatched_shipment_rows),
            )

            if unmatched_shipment_rows:
                _update_task(task_id, 68, f"Upserting {len(unmatched_shipment_rows)} new shipment rows...")
                for i in range(0, len(unmatched_shipment_rows), STAGING_INSERT_BATCH):
                    chunk = [
                        _removal_row_for_write(r)
                        for r in unmatched_shipment_rows[i : i + STAGING_INSERT_BATCH]
                    ]
                    db.table("amazon_removals").upsert(
                        chunk, on_conflict=_REMOVALS_CONFLICT,
                    ).execute()
                    upserted_total += len(chunk)

        staged_line_count = len(extracted_rows)
        accounted_for = upserted_total + enriched_count
        if accounted_for < staged_line_count:
            log.warning(
                "[sync-removals] ROW COUNT MISMATCH: %d staging lines but only %d accounted for "
                "(upserted=%d + enriched_in_place=%d). %d rows may have been silently dropped. "
                "Check for upstream errors in the log above.",
                staged_line_count,
                accounted_for,
                upserted_total,
                enriched_count,
                staged_line_count - accounted_for,
            )
        else:
            log.info(
                "[sync-removals] Row count OK: all %d staging lines accounted for "
                "(upserted=%d, enriched_in_place=%d).",
                staged_line_count,
                upserted_total,
                enriched_count,
            )
        if skipped > 0:
            log.info(
                "[sync-removals] Staging rows skipped (total=%d, no_order_id=%d): "
                "these remain in amazon_staging and are excluded from the counts above.",
                skipped,
                skipped_no_order_id,
            )

        _update_task(task_id, 82, f"Cleaning up {len(staging_ids)} rows from amazon_staging...")
        log.info("[sync-removals] Deleting %d staging rows by id", len(staging_ids))
        for i in range(0, len(staging_ids), STAGING_INSERT_BATCH):
            db.table("amazon_staging").delete().in_(
                "id", staging_ids[i : i + STAGING_INSERT_BATCH]
            ).execute()

        log.info(
            "[sync-removals] wave1_reconciliation %s",
            json.dumps(
                {
                    "store_id": store_id,
                    "staging_lines": staged_line_count,
                    "upsert_chunks_rows": upserted_total,
                    "enriched_in_place": enriched_count,
                    "skipped_staging": skipped,
                },
                default=str,
            ),
        )
        log.info(
            "[sync-removals] DONE — upserted=%d updated_in_place=%d staging_skipped=%d",
            upserted_total,
            enriched_count,
            skipped,
        )
        _update_task(
            task_id, 100,
            (
                f"Sync complete: {upserted_total} upserts + {enriched_count} in-place enrichments "
                f"into amazon_removals ({len(shipment_history_rows)} raw shipment rows archived). "
                "Run Generate Worklist to update expected_packages."
            ),
            status="completed",
        )

    except Exception as e:
        log.error("[sync-removals] FATAL ERROR:\n%s", traceback.format_exc())
        _update_task(task_id, 0, f"Sync failed: {e!s}", status="failed")

def _generate_worklist_core(
    db: Any,
    task_id: str,
    organization_id: str,
    *,
    progress_start: int = 0,
    upload_id: str | None = None,
) -> None:

    def _prog(fraction: float, msg: str) -> None:
        actual = min(int(progress_start + fraction * (100 - progress_start)), 99)
        _update_task(task_id, actual, msg)

    try:
        session_upload = str(upload_id).strip() if upload_id else None

        _prog(0.02, "Fetching amazon_removals rows for worklist generation...")
        log.info("[worklist-core] Fetching amazon_removals for org=%s", organization_id)

        try:
            removal_rows = _fetch_org_table_all(db, "amazon_removals", organization_id)
        except Exception as fetch_err:
            log.exception("[worklist-core] FETCH amazon_removals FAILED")
            _update_task(task_id, 0, f"Worklist fetch failed: {fetch_err!s}", status="failed")
            return

        n_removals_loaded = len(removal_rows)
        log.info("[worklist-core] amazon_removals rows loaded for generation: %d", n_removals_loaded)

        if session_upload:
            scoped = [r for r in removal_rows if str(r.get("upload_id") or "").strip() == session_upload]
            if scoped:
                removal_rows = scoped
                log.info(
                    "[worklist-core] scoped to upload_id=%s: %d amazon_removals row(s)",
                    session_upload,
                    len(removal_rows),
                )
            else:
                ship_rows = _fetch_amazon_removal_shipments_by_upload(db, organization_id, session_upload)
                removal_rows = [_removal_like_from_shipment_archive(x) for x in ship_rows]
                log.info(
                    "[worklist-core] no amazon_removals for upload_id=%s: using %d amazon_removal_shipments row(s)",
                    session_upload,
                    len(removal_rows),
                )

        if not removal_rows:
            log.warning("[worklist-core] no input rows for this worklist run — nothing to write")
            _update_task(task_id, 100, "No removal or shipment rows for this upload; worklist unchanged.", status="completed")
            return

        _prog(0.08, "Loading existing expected_packages for merge...")
        try:
            existing_ep = _fetch_org_table_all(db, "expected_packages", organization_id)
        except Exception as ep_err:
            log.exception("[worklist-core] FETCH expected_packages FAILED")
            _update_task(task_id, 0, f"Worklist existing fetch failed: {ep_err!s}", status="failed")
            return

        existing_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}
        for er in existing_ep:
            existing_by_key[_expected_pkg_identity_key(er)] = er

        _prog(0.12, f"Mapping {len(removal_rows)} rows to expected_packages format...")

        merged_rows: list[dict[str, Any]] = []
        skipped_no_order = 0
        skipped_non_return = 0

        for row in removal_rows:
            order_id = _pg_text_unique_field(row.get("order_id"))
            if not order_id:
                skipped_no_order += 1
                continue
            if not _is_return_order_type_for_worklist(row):
                skipped_non_return += 1
                continue

            incoming: dict[str, Any] = {"organization_id": organization_id, "order_id": order_id}
            for k in EXPECTED_PKG_AMAZON_COLS:
                if k in ("organization_id", "order_id"):
                    continue
                if k not in row:
                    continue
                val = row[k]
                if k in _EP_WORKLIST_AMAZON_FILL_NULL and _ep_value_absent(val):
                    continue
                incoming[k] = val

            if not incoming.get("sku"):
                incoming["sku"] = ""
            if not incoming.get("order_status"):
                incoming["order_status"] = "Pending"

            if session_upload:
                incoming["upload_id"] = session_upload

            k = _expected_pkg_identity_key(incoming)
            if k in existing_by_key:
                merged_rows.append(_merge_expected_pkg_amazon_only(existing_by_key[k], incoming))
            else:
                merged_rows.append(incoming)

        log.info(
            "[worklist-core] Prepared %d upsert rows (skipped %d no order_id, %d non-Return order_type)",
            len(merged_rows), skipped_no_order, skipped_non_return,
        )

        if skipped_no_order > 0:
            log.warning(
                "[worklist-core] SKIPPED %d removal rows with no order_id — "
                "these will NOT appear in expected_packages.",
                skipped_no_order,
            )
        if skipped_non_return > 0:
            log.info(
                "[worklist-core] SKIPPED %d removal rows (order_type is not Return) — "
                "expected_packages only includes Return lines.",
                skipped_non_return,
            )

        if not merged_rows:
            log.info(
                "[worklist-core] rows skipped before insert: no_order=%d non_return=%d — no merged rows",
                skipped_no_order,
                skipped_non_return,
            )
            _update_task(task_id, 100, "No worklist rows produced from amazon_removals.", status="completed")
            return

        merged_before_dedupe = len(merged_rows)
        merged_rows, merged_collapsed = _dedupe_merged_rows_for_expected_packages_upsert(merged_rows)
        merged_after_dedupe = len(merged_rows)
        log.info("[worklist-core] merged_rows_before_dedupe: %d", merged_before_dedupe)
        log.info("[worklist-core] merged_rows_after_dedupe: %d", merged_after_dedupe)
        log.info("[worklist-core] merged_rows_collapsed_by_conflict_key: %d", merged_collapsed)

        _prog(0.28, f"Upserting {len(merged_rows)} rows into expected_packages...")

        upserted = 0
        n_pkg = len(merged_rows)
        for i in range(0, n_pkg, STAGING_INSERT_BATCH):
            chunk = [
                {k: v for k, v in r.items() if k not in _DB_SYSTEM_COLS}
                for r in merged_rows[i : i + STAGING_INSERT_BATCH]
            ]
            try:
                db.table("expected_packages").upsert(
                    chunk, on_conflict=_EXPECTED_PKG_CONFLICT,
                ).execute()
                upserted += len(chunk)
                frac = 0.28 + (upserted / max(n_pkg, 1)) * 0.70
                _prog(frac, f"Upserted {upserted}/{n_pkg} expected_packages rows...")
            except Exception as ins_err:
                log.exception("[worklist-core] UPSERT chunk at offset %d FAILED", i)
                _update_task(
                    task_id, 0,
                    f"expected_packages upsert failed at offset {i}: {ins_err!s}",
                    status="failed",
                )
                return

        ep_rows_generated = len(merged_rows)
        ep_rows_populated_shipment_derived = sum(
            1
            for r in merged_rows
            if any(not _ep_value_absent(r.get(fk)) for fk in _EP_WORKLIST_AMAZON_FILL_NULL)
        )
        ep_rows_missing_tracking_carrier_or_shipment_date = sum(
            1
            for r in merged_rows
            if _ep_value_absent(r.get("tracking_number"))
            or _ep_value_absent(r.get("carrier"))
            or _ep_value_absent(r.get("shipment_date"))
        )

        log.info(
            "[worklist-core] REMOVAL_PIPELINE %s",
            json.dumps(
                {
                    "checkpoint": "REMOVAL_PIPELINE",
                    "stage": "C_expected_packages",
                    "organization_id": organization_id,
                    "expected_rows_generated": ep_rows_generated,
                    "expected_rows_populated_shipment_derived": ep_rows_populated_shipment_derived,
                    "expected_rows_missing_tracking_carrier_or_shipment_date": (
                        ep_rows_missing_tracking_carrier_or_shipment_date
                    ),
                    "expected_packages_upserted": upserted,
                    "skipped_no_order": skipped_no_order,
                    "skipped_non_return": skipped_non_return,
                },
                default=str,
            ),
        )

        log.info("[worklist-core] expected_packages rows inserted: %d", upserted)

        if upserted < len(merged_rows):
            log.warning(
                "[worklist-core] ROW COUNT MISMATCH: prepared %d rows but only %d were upserted.",
                len(merged_rows), upserted,
            )
        explained = skipped_no_order + skipped_non_return + len(merged_rows)
        if explained != len(removal_rows):
            log.warning(
                "[worklist-core] ROW COUNT WARNING: input_rows=%d but "
                "skipped_no_order=%d + skipped_non_return=%d + worklist=%d = %d.",
                len(removal_rows),
                skipped_no_order,
                skipped_non_return,
                len(merged_rows),
                explained,
            )
        else:
            log.info(
                "[worklist-core] Row count OK: %d input = %d skipped (no order) + "
                "%d skipped (non-Return) + %d worklist row(s).",
                len(removal_rows), skipped_no_order, skipped_non_return, len(merged_rows),
            )

        log.info(
            "[worklist-core] DONE — upserted %d rows (input_rows=%d, skipped_no_order=%d, skipped_non_return=%d)",
            upserted, len(removal_rows), skipped_no_order, skipped_non_return,
        )
        log.info(
            "[worklist-core] wave1_reconciliation %s",
            json.dumps(
                {
                    "phase": "generate_worklist",
                    "upload_id": session_upload,
                    "input_rows": len(removal_rows),
                    "expected_packages_upserted": upserted,
                    "expected_packages_rows_generated": ep_rows_generated,
                    "expected_packages_rows_populated_shipment_derived": ep_rows_populated_shipment_derived,
                    "expected_packages_rows_missing_tracking_carrier_or_shipment_date": (
                        ep_rows_missing_tracking_carrier_or_shipment_date
                    ),
                    "skipped_no_order": skipped_no_order,
                    "skipped_non_return": skipped_non_return,
                },
                default=str,
            ),
        )
        _update_task(
            task_id, 100,
            f"Worklist complete: {upserted} rows upserted into expected_packages "
            f"(Amazon fields updated; warehouse scan progress preserved).",
            status="completed",
        )

    except Exception as e:
        log.exception("[worklist-core] Phase 4 (generate worklist) failed")
        _update_task(task_id, 0, f"Worklist failed: {e!s}", status="failed")

def _run_generate_worklist(
    task_id: str,
    organization_id: str,
    upload_id: str | None,
) -> None:
    print(f"[generate-worklist] START task_id={task_id} org={organization_id}")
    _update_task(task_id, 5, "Task is alive. Connecting to database...")

    _url = os.getenv("SUPABASE_URL")
    _key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    if not _url or not _key:
        _update_task(task_id, 0, "Missing Supabase credentials in environment.", status="failed")
        return

    try:
        db = create_client(_url, _key)
        print("[generate-worklist] Supabase client created successfully")
    except Exception as conn_err:
        _update_task(task_id, 0, f"Failed to connect to database: {conn_err!s}", status="failed")
        return

    try:
        _generate_worklist_core(db, task_id, organization_id, progress_start=0, upload_id=upload_id)
    except Exception:
        log.exception("[generate-worklist] background task crashed")
        _update_task(task_id, 0, "Worklist task crashed — see server logs.", status="failed")

@app.post("/etl/sync-removals")
async def etl_sync_removals(
    request: SyncRemovalsRequest,
    background_tasks: BackgroundTasks,
):
    try:
        uuid.UUID(request.organization_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="organization_id must be a valid UUID.")

    if request.upload_id:
        try:
            uuid.UUID(request.upload_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="upload_id must be a valid UUID.")

    task_id = str(uuid.uuid4())
    _update_task(task_id, 0, "Task queued.", status="queued")
    background_tasks.add_task(_run_sync_removals, task_id, request.organization_id, request.upload_id)
    return {"task_id": task_id, "status": "queued", "poll_url": f"/etl/task/{task_id}"}

@app.post("/etl/generate-worklist")
async def etl_generate_worklist(
    request: GenerateWorklistRequest,
    background_tasks: BackgroundTasks,
):
    try:
        uuid.UUID(request.organization_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="organization_id must be a valid UUID.")

    if request.upload_id:
        try:
            uuid.UUID(request.upload_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="upload_id must be a valid UUID.")

    task_id = str(uuid.uuid4())
    _update_task(task_id, 0, "Task queued.", status="queued")
    background_tasks.add_task(_run_generate_worklist, task_id, request.organization_id, request.upload_id)
    return {"task_id": task_id, "status": "queued", "poll_url": f"/etl/task/{task_id}"}

@app.get("/etl/task/{task_id}")
async def get_task_status(task_id: str):
    if task_id not in task_store:
        raise HTTPException(status_code=404, detail=f"Task '{task_id}' not found.")
    return task_store[task_id]

# -------------------------------------------------------------------------
# UI ARCHITECTURE ENDPOINTS & DATA ROUTING
# -------------------------------------------------------------------------

def _resolve_openai_api_key_from_db(db: Any, org_id: str) -> str | None:
    """
    Match Next.js getOrganizationOpenAIApiKey / Settings upsertProviderApiKey:
    plaintext in api_key, role llm_provider, name often 'OpenAI' (not 'openai_api_key').
    """
    try:
        res = (
            db.table("organization_api_keys")
            .select("api_key, name")
            .eq("organization_id", org_id)
            .eq("role", "llm_provider")
            .order("created_at", desc=True)
            .limit(24)
            .execute()
        )
        rows = list(res.data or [])
        for row in rows:
            if str(row.get("name") or "").strip() == "OpenAI":
                k = str(row.get("api_key") or "").strip()
                if k:
                    return k
        for row in rows:
            n = str(row.get("name") or "").strip().lower()
            k = str(row.get("api_key") or "").strip()
            if not k:
                continue
            if "openai" in n or "gpt" in n or n in ("chatgpt", "llm"):
                return k
    except Exception as e:
        log.warning("OpenAI llm_provider key lookup failed org=%s: %s", org_id, e)

    # Legacy: row name exactly openai_api_key (any role)
    try:
        res = (
            db.table("organization_api_keys")
            .select("api_key")
            .eq("organization_id", org_id)
            .eq("name", "openai_api_key")
            .order("created_at", desc=True)
            .limit(1)
            .execute()
        )
        if res.data and res.data[0].get("api_key"):
            k = str(res.data[0]["api_key"]).strip()
            if k:
                return k
    except Exception as e:
        log.warning("OpenAI legacy name=openai_api_key lookup failed org=%s: %s", org_id, e)

    # Last resort: any row for this org whose name suggests OpenAI (mis-set role)
    try:
        res = db.table("organization_api_keys").select("api_key, name").eq("organization_id", org_id).execute()
        for row in res.data or []:
            n = str(row.get("name") or "").strip().lower()
            if "openai" not in n and "gpt" not in n:
                continue
            k = str(row.get("api_key") or "").strip()
            if k:
                return k
    except Exception as e:
        log.warning("OpenAI broad name scan failed org=%s: %s", org_id, e)

    return None


def _get_openai_key(db: Any, org_id: str) -> str | None:
    """Organization DB first (llm_provider / OpenAI), then OPENAI_API_KEY env."""
    k = _resolve_openai_api_key_from_db(db, org_id)
    if k:
        return k
    env_key = os.getenv("OPENAI_API_KEY")
    return env_key.strip() if env_key and str(env_key).strip() else None

def _normalize_to_ui_report_slug(raw: str | None) -> str:
    """
    Maps GPT / user / legacy tokens to a small set of UI routing slugs.
    DB CHECK constraints use different spellings — see _ui_slug_to_db_report_type().
    """
    if not raw:
        return "unknown"
    s = str(raw).strip().lower()
    for junk in ('"', "'", "`", ".", ",", ";", ":"):
        s = s.replace(junk, "")
    s = s.split()[0] if s else ""

    synonyms: dict[str, str] = {
        "inventory_ledger": "inventory_ledger",
        "inventory": "inventory_ledger",
        "fba_inventory_ledger": "inventory_ledger",
        "ledger": "inventory_ledger",
        "inventory-ledger": "inventory_ledger",
        "reimbursements": "reimbursements",
        "reimbursement": "reimbursements",
        "fba_reimbursements": "reimbursements",
        "removals": "removals",
        "removal": "removals",
        "removal_shipment": "removals",
        "removal_shipments": "removals",
        "removal_shipment_detail": "removals",
        "removal_order": "removals",
        "removal_orders": "removals",
        "removal-order-id": "removals",
        "amazon_all_orders": "amazon_all_orders",
        "all_orders": "amazon_all_orders",
        # GPT might echo Postgres-safe labels — fold back to UI slug for routing
        "removal_shipment": "removals",
        "unknown": "unknown",
    }
    return synonyms.get(s, "unknown")

def _normalized_header_set(headers: list[str]) -> set[str]:
    return {_normalize_header(str(h)) for h in headers if str(h).strip()}

def _detect_report_type_by_rules(headers: list[str]) -> tuple[str, float, str]:
    """
    Deterministic first pass. This avoids unnecessary GPT calls for the common
    Amazon exports and makes the UI feel instant/reliable.
    """
    h = _normalized_header_set(headers)
    if not h:
        return "unknown", 0.0, "rules"

    if {"date", "fnsku", "asin", "msku", "event_type", "quantity"}.issubset(h):
        return "inventory_ledger", 0.98, "rules"
    if {"date_and_time", "fnsku", "asin", "msku", "event_type", "quantity"}.issubset(h):
        return "inventory_ledger", 0.98, "rules"
    if {"reimbursement_id", "reason", "sku"}.issubset(h) or {"approval_date", "reimbursement_id"}.issubset(h):
        return "reimbursements", 0.98, "rules"
    if {"order_id", "sku", "fnsku", "disposition"}.issubset(h) and (
        "requested_quantity" in h or "shipped_quantity" in h or "tracking_number" in h
    ):
        return "removals", 0.96, "rules"
    if {"amazon_order_id", "purchase_date", "order_status"}.issubset(h):
        return "amazon_all_orders", 0.98, "rules"
    if {"order_id", "purchase_date", "order_status"}.issubset(h):
        return "amazon_all_orders", 0.9, "rules"

    return "unknown", 0.0, "rules"


def _ui_slug_to_db_report_type(ui_slug: str) -> str:
    """Maps UI slug to a value that satisfies raw_report_uploads.report_type CHECK."""
    mapping = {
        "inventory_ledger": "inventory_ledger",
        "reimbursements": "reimbursements",
        "removals": "REMOVAL_SHIPMENT",
        "amazon_all_orders": "ALL_ORDERS",
        "unknown": "UNKNOWN",
    }
    return mapping.get(ui_slug, "UNKNOWN")


ALLOWED_UI_REPORT_SLUGS = frozenset(
    {"inventory_ledger", "reimbursements", "removals", "amazon_all_orders"}
)


def _detect_report_type_with_gpt(headers: list[str], api_key: str) -> tuple[str, float, str]:
    """Uses GPT to classify the report; returns a normalized UI slug."""
    try:
        from openai import OpenAI

        client = OpenAI(api_key=api_key)
        header_sample = ", ".join(headers[:80])
        prompt = (
            "You classify Amazon Seller Central export CSVs using COLUMN HEADERS ONLY.\n\n"
            "Reply with exactly ONE lowercase slug from this list (no punctuation, no explanation):\n"
            "- inventory_ledger — FBA Inventory Ledger / inventory events (SKU, fulfillment center, quantities, event types)\n"
            "- reimbursements — FBA reimbursements / repayment (reimbursement id, fee/reason, currency amounts)\n"
            "- removals — Removal orders / removal shipments / disposal (removal-order-id, disposition, shipped quantity)\n"
            "- amazon_all_orders — All Orders style report (amazon-order-id, order status, purchase date)\n"
            "- unknown — if none of the above fit\n\n"
            f"Headers:\n{header_sample}"
        )
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.0,
            max_tokens=24,
        )
        raw_out = response.choices[0].message.content
        detected = _normalize_to_ui_report_slug(raw_out)
        return detected, 0.8 if detected != "unknown" else 0.0, "gpt"
    except Exception as e:
        log.warning(f"GPT detection failed: {e}")
        return "unknown", 0.0, "gpt"

class DetectHeadersRequest(BaseModel):
    headers: list[str]
    organization_id: str


class PimImportPreviewStepBody(BaseModel):
    organization_id: str
    store_id: str
    upload_id: str
    row_chunk: int | None = None
    """Optional client hint for diagnostics only; server uses persisted scan_cursor."""
    scan_data_row_hint: int | None = None


class PimImportApplyStepBody(BaseModel):
    organization_id: str
    store_id: str
    upload_id: str
    confirm: str = "false"
    row_chunk: int | None = None
    skip_conflicts: bool = False
    import_safe_rows_only: bool = False

    @property
    def safe_rows_only(self) -> bool:
        """Canonical check — True when either legacy or new flag is set."""
        return self.skip_conflicts or self.import_safe_rows_only


class PimImportPriceBackfillStepBody(BaseModel):
    organization_id: str
    upload_id: str
    row_chunk: int | None = None
    restart: bool = False
    cancel: bool = False


@app.post("/etl/detect-headers")
async def etl_detect_headers(request: DetectHeadersRequest):
    """Architectural Route for Client-Side Slicing auto-detection."""
    db = _require_supabase()
    rules_type, confidence, method = _detect_report_type_by_rules(request.headers)
    if rules_type != "unknown":
        return {
            "detected_type": rules_type,
            "confidence": confidence,
            "method": method,
        }

    api_key = _get_openai_key(db, request.organization_id)
    if not api_key:
        return {"detected_type": "unknown", "confidence": 0.0, "method": "none"}

    detected, confidence, method = _detect_report_type_with_gpt(request.headers, api_key)
    return {
        "detected_type": detected,
        "confidence": confidence,
        "method": method,
    }

@app.get("/etl/upload-history/{organization_id}")
async def etl_upload_history(organization_id: str):
    """Fetches the upload history for the frontend table + per-upload pipeline progress."""
    db = _require_supabase()
    try:
        uuid.UUID(organization_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid organization ID.")

    try:
        res = (
            db.table("raw_report_uploads")
            .select("*")
            .eq("organization_id", organization_id)
            .order("created_at", desc=True)
            .limit(50)
            .execute()
        )
        rows = list(res.data or [])
        ids = [r.get("id") for r in rows if r.get("id")]
        fps_by_upload: dict[str, dict[str, Any]] = {}
        if ids:
            try:
                fps_res = (
                    db.table("file_processing_status")
                    .select("*")
                    .in_("upload_id", ids)
                    .execute()
                )
                for fps in fps_res.data or []:
                    uid = fps.get("upload_id")
                    if isinstance(uid, str):
                        fps_by_upload[uid] = fps
            except Exception as fe:
                log.warning(f"file_processing_status join skipped: {fe}")
        for r in rows:
            uid = r.get("id")
            if isinstance(uid, str) and uid in fps_by_upload:
                r["pipeline"] = fps_by_upload[uid]
        return {"status": "success", "history": rows}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database fetch error: {e}")

@app.delete("/etl/upload/{upload_id}")
async def etl_delete_upload(upload_id: str):
    """Deletes an upload record and its physical file from storage."""
    db = _require_supabase()
    try:
        res = db.table("raw_report_uploads").select("metadata").eq("id", upload_id).execute()
        if res.data and res.data[0].get("metadata"):
            storage_path = res.data[0]["metadata"].get("storage_path")
            if storage_path:
                try:
                    db.storage.from_("raw-reports").remove([storage_path])
                except Exception as e:
                    log.warning(f"Failed to delete physical file: {e}")
                    
        db.table("raw_report_uploads").delete().eq("id", upload_id).execute()
        return {"status": "success", "message": "Record deleted."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Deletion failed: {e}")

@app.post("/etl/upload-raw")
async def etl_upload_raw(
    file: UploadFile = File(...),
    report_type: str = Form(...),
    organization_id: str = Form(...),
    store_id: str = Form(...)
):
    db = _require_supabase()
    
    try:
        uuid.UUID(organization_id)
        uuid.UUID(store_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid UUID format for organization_id or store_id.")

    raw_bytes = await file.read()
    if not raw_bytes:
        raise HTTPException(status_code=400, detail="Empty file.")
    file_hash = hashlib.sha256(raw_bytes).hexdigest()

    # Idempotency Check. Only completed/synced uploads are considered duplicates.
    # Previous failed or stuck "processing" rows must not block a retry because
    # they may have been created before the raw file reached storage.
    try:
        existing_res = (
            db.table("raw_report_uploads")
            .select("id, report_type, status, metadata")
            .eq("organization_id", organization_id)
            .execute()
        )
        for row in existing_res.data:
            meta = row.get("metadata") or {}
            status = str(row.get("status") or "").lower()
            reusable_status = status in {"complete", "synced", "mapped", "ready", "uploaded"}
            if meta.get("file_hash") == file_hash and reusable_status:
                return {
                    "status": "success",
                    "message": "Duplicate file prevented. File already exists in the system.",
                    "upload_id": row["id"],
                    "detected_type": (row.get("metadata") or {}).get("ui_report_slug") or row.get("report_type"),
                }
    except Exception as e:
        log.warning(f"Duplicate check bypassed: {e}")

    headers, row_count = _csv_headers_and_row_count_fast(raw_bytes)
    if not headers:
        raise HTTPException(
            status_code=400,
            detail="Could not read a header row. Ensure the file is UTF-8 CSV or TSV.",
        )
    if row_count < 1:
        raise HTTPException(status_code=400, detail="No data rows in file.")

    raw_rt = (report_type or "").strip()
    ui_slug = _normalize_to_ui_report_slug(raw_rt)

    if ui_slug == "unknown" or raw_rt.lower() in (
        "auto",
        "auto-detect",
        "",
        "null",
    ):
        api_key = _get_openai_key(db, organization_id)
        if api_key:
            ui_slug, _, _ = _detect_report_type_with_gpt(headers, api_key)
        else:
            ui_slug = "unknown"

    if ui_slug not in ALLOWED_UI_REPORT_SLUGS:
        raise HTTPException(
            status_code=400,
            detail=(
                "Report type could not be determined. "
                "Pick inventory_ledger, reimbursements, removals, or amazon_all_orders from the dropdown."
            ),
        )

    db_report_type = _ui_slug_to_db_report_type(ui_slug)
    target_table_by_slug = {
        "inventory_ledger": "amazon_inventory_ledger",
        "reimbursements": "amazon_reimbursements",
        "removals": "amazon_staging",
        "amazon_all_orders": "amazon_all_orders",
    }
    target_table = target_table_by_slug.get(ui_slug)

    safe_filename = (file.filename or "upload.csv").replace(" ", "_")
    storage_path = f"{organization_id}/{store_id}/{file_hash}_{safe_filename}"

    try:
        db.storage.from_("raw-reports").upload(
            path=storage_path,
            file=raw_bytes,
            file_options={
                "content-type": file.content_type or "text/csv",
                "upsert": "true",
            },
        )
    except Exception as e:
        log.exception(f"Storage upload failed for {storage_path}: {e}")
        raise HTTPException(
            status_code=502,
            detail=f"Could not save file to storage. Check bucket 'raw-reports' and service role permissions. ({e})",
        ) from e

    upload_id: str | None = None
    try:
        upload_record = {
            "organization_id": organization_id,
            "report_type": db_report_type,
            "file_name": file.filename or "upload.csv",
            "metadata": {
                "file_hash": file_hash,
                "content_sha256": file_hash,
                "row_count": row_count,
                "total_rows": row_count,
                "import_store_id": store_id,
                "ledger_store_id": store_id,
                "storage_path": storage_path,
                "raw_file_path": storage_path,
                "upload_chunks_count": 1,
                "total_parts": 1,
                "upload_progress": 100,
                "total_bytes": len(raw_bytes),
                "uploaded_bytes": len(raw_bytes),
                "csv_headers": headers,
                "headers": headers[:15],
                "ui_report_slug": ui_slug,
                "target_table": target_table,
                "etl_source": "amazon_etl_quick_upload",
            },
            "status": "mapped",
        }

        upload_res = db.table("raw_report_uploads").insert(upload_record).execute()
        if not upload_res.data:
            raise RuntimeError("Insert returned no row — check raw_report_uploads schema and RLS.")
        upload_id = str(upload_res.data[0]["id"])

        return {
            "status": "success",
            "message": (
                f"File saved to storage ({row_count:,} data rows). "
                "Open Imports → find this upload → Process (staging), then Sync (Amazon tables)."
            ),
            "upload_id": upload_id,
            "detected_type": ui_slug,
            "report_type_db": db_report_type,
            "next_step": "imports_process",
        }

    except HTTPException:
        raise
    except Exception as e:
        log.error(f"Error finalizing upload: {traceback.format_exc()}")
        if upload_id:
            try:
                db.table("raw_report_uploads").update({"status": "failed"}).eq("id", upload_id).execute()
            except Exception as ue:
                log.warning(f"Could not mark upload failed: {ue}")
        raise HTTPException(status_code=500, detail=f"System failed to finalize the upload: {e}") from e






# API GOOGLE SHEETS
# -------------------------------------------------------------------------
# PHASE 1: ENTERPRISE CATALOG SEEDING (Live Amazon, GPT Mapping, & Vendors)
# -------------------------------------------------------------------------
import requests
import time
from datetime import datetime, timezone

from dataclasses import dataclass, replace

from pim_product_master import (
    PRODUCT_MASTER_SHEET_NAME,
    canonical_category_display_name,
    collect_product_attributes_from_row,
    merge_product_attributes_into_metadata,
    normalize_category_label_for_key,
    normalize_pim_status,
)
from pim_seed_cleaning import (
    ParsedIdentifierRow,
    build_identifier_map_variants,
    finalize_ambiguity_with_db,
    ordered_unique,
    parse_identifier_row,
    trim_cell_value,
)

def _get_api_credentials(db: Any, org_id: str, api_name: str) -> Any:
    if api_name == "openai_api_key":
        k = _resolve_openai_api_key_from_db(db, org_id)
        if k:
            return k
        env_key = os.getenv("OPENAI_API_KEY")
        return env_key.strip() if env_key and str(env_key).strip() else None

    try:
        res = db.table("organization_api_keys").select("api_key").eq("organization_id", org_id).eq("name", api_name).execute()
        if res.data and res.data[0].get("api_key"):
            key_val = res.data[0]["api_key"]
            if api_name == "amazon_sp_api":
                return json.loads(key_val) if isinstance(key_val, str) else key_val
            return key_val
    except Exception as e:
        log.warning(f"Failed to fetch {api_name} key from DB: {e}")
    return None

def _fetch_amazon_catalog_data(asin: str, creds: dict) -> dict:
    try:
        token_res = requests.post(
            "https://api.amazon.com/auth/o2/token",
            data={
                "grant_type": "refresh_token",
                "refresh_token": creds.get("refresh_token"),
                "client_id": creds.get("client_id"),
                "client_secret": creds.get("client_secret"),
            },
        )
        token_res.raise_for_status()
        access_token = token_res.json().get("access_token")

        headers = {"x-amz-access-token": access_token, "Content-Type": "application/json"}
        params = {"marketplaceIds": "ATVPDKIKX0DER", "includedData": "summaries,images"}
        cat_res = requests.get(
            f"https://sellingpartnerapi-na.amazon.com/catalog/2022-04-01/items/{asin}",
            headers=headers,
            params=params,
        )
        cat_res.raise_for_status()
        item = cat_res.json()

        title = item.get("summaries", [{}])[0].get("itemName")
        brand = item.get("summaries", [{}])[0].get("brand")
        images = item.get("images", [{}])[0].get("images", [])
        image_url = images[0].get("link") if images else None

        return {"product_name": title, "brand": brand, "main_image_url": image_url, "amazon_raw": item}
    except Exception as e:
        log.warning(
            "Amazon SP-API catalog fetch failed asin=%s error=%s",
            asin,
            e,
            exc_info=True,
        )
        return {}


_PIM_MAP_STANDARD_KEYS = (
    "vendor",
    "category",
    "mfg_part",
    "product_name",
    "seller_sku",
    "asin",
    "fnsku",
    "upc",
    "cost",
    "status",
)


def _pim_public_column_map(column_map: dict[str, Any]) -> dict[str, str]:
    out: dict[str, str] = {}
    for k in _PIM_MAP_STANDARD_KEYS:
        v = column_map.get(k)
        if v is not None and str(v).strip():
            out[str(k)] = str(v).strip()
    return out


def _pim_build_header_norm_index(headers: list[str]) -> dict[str, str]:
    """normalized_header -> first exact header string from the file."""
    out: dict[str, str] = {}
    for h in headers:
        t = str(h).strip()
        if not t:
            continue
        nk = _normalize_header(t)
        if nk not in out:
            out[nk] = t
    return out


def _pim_try_deterministic_column_map(headers: list[str]) -> dict[str, Any] | None:
    """
    Rule-based map for common PIM exports (e.g. UPC, Vendor, Seller SKU, Mfg #, FNSKU, ASIN, Product Name).
    Skips GPT when we have product_name plus at least one identifier column.
    """
    idx = _pim_build_header_norm_index(headers)

    def pick(*norm_keys: str) -> str | None:
        for k in norm_keys:
            v = idx.get(k)
            if v:
                return v
        return None

    m: dict[str, Any] = {}
    maybe_vendor = pick("vendor", "supplier", "distributor", "manufacturer", "brand_owner")
    maybe_category = pick(
        "category",
        "product_category",
        "product_type",
        "producttype",
        "department",
        "amazon_category",
    )
    maybe_product = pick(
        "product_name",
        "productname",
        "title",
        "item_name",
        "itemname",
        "product_title",
        "producttitle",
        "name",
        "listing_title",
    )
    maybe_sku = pick("seller_sku", "sellersku", "merchant_sku", "merchantsku", "msku", "sku")
    maybe_mfg = pick(
        "mfg_#",
        "mfg_hash",
        "mfg_no",
        "mfg_num",
        "mfg_number",
        "mfg_part",
        "mfg_part_number",
        "manufacturer_part_number",
        "manufacturer_part",
        "mpn",
        "part_number",
    )
    maybe_asin = pick("asin")
    maybe_fnsku = pick("fnsku", "fulfillment_channel_sku", "fulfillmentchannelsku")
    maybe_upc = pick("upc", "upc_ean", "upc_code", "ean", "gtin", "barcode")
    maybe_cost = pick("cost", "unit_cost", "unitcost", "list_price", "listprice", "price", "last_cost")
    maybe_status = pick("status", "product_status", "lifecycle", "state", "active", "listing_status")

    if maybe_vendor:
        m["vendor"] = maybe_vendor
    if maybe_category:
        m["category"] = maybe_category
    if maybe_product:
        m["product_name"] = maybe_product
    if maybe_sku:
        m["seller_sku"] = maybe_sku
    if maybe_mfg:
        m["mfg_part"] = maybe_mfg
    if maybe_asin:
        m["asin"] = maybe_asin
    if maybe_fnsku:
        m["fnsku"] = maybe_fnsku
    if maybe_upc:
        m["upc"] = maybe_upc
    if maybe_cost:
        m["cost"] = maybe_cost
    if maybe_status:
        m["status"] = maybe_status

    ident = bool(m.get("seller_sku") or m.get("asin") or m.get("fnsku") or m.get("upc"))
    if ident and m.get("product_name"):
        return m
    return None


def _pim_resolve_column_map(headers: list[str], openai_key: str | None) -> tuple[dict[str, Any], str]:
    det = _pim_try_deterministic_column_map(headers)
    if det is not None:
        return det, "deterministic"
    key = (openai_key or "").strip()
    if not key:
        raise HTTPException(
            status_code=400,
            detail={
                "error": "column_mapping_requires_openai",
                "message": (
                    "Could not map columns from headers without GPT. Save an OpenAI LLM key in Settings "
                    "(organization_api_keys) or set OPENAI_API_KEY on the ETL server — or use a file whose headers "
                    "include Product Name plus at least one of: Seller SKU, ASIN, FNSKU, UPC."
                ),
                "headers": headers,
            },
        )
    return _gpt_map_catalog_columns_or_raise(headers, key), "gpt"


def _pim_coerce_mapped_columns_string(df: pd.DataFrame, handles: dict[str, str | None]) -> pd.DataFrame:
    """Treat mapped identity / text columns as plain strings (leading zeros, avoid float UPC drift)."""
    out = df.copy()
    sci_re = re.compile(r"^-?\d+(\.\d+)?[eE][+-]?\d+$")

    def to_text(v: Any) -> str:
        if v is None:
            return ""
        try:
            if isinstance(v, float) and pd.isna(v):
                return ""
        except Exception:
            pass
        if isinstance(v, bool):
            return "TRUE" if v else "FALSE"
        if isinstance(v, float):
            if math.isfinite(v) and v == int(v):
                return str(int(v))
            return str(v).strip()
        if isinstance(v, int) and not isinstance(v, bool):
            return str(v)
        s = str(v).strip()
        if sci_re.match(s):
            try:
                f = float(s)
                if math.isfinite(f) and f == int(f):
                    return str(int(f))
            except (TypeError, ValueError):
                pass
        return s

    for std in _PIM_MAP_STANDARD_KEYS:
        col = handles.get(std)
        if not col or col not in out.columns:
            continue
        out[col] = out[col].map(to_text)
    return out


def _pim_seed_history_insert_preview(
    db: Any,
    organization_id: str,
    store_id: str,
    file_name: str | None,
    column_map: dict[str, str],
    mapping_source: str,
    quality: dict[str, Any],
) -> str | None:
    """Insert raw_report_uploads row for PIM seed preview (best-effort; returns id or None)."""
    try:
        now = datetime.now(timezone.utc).isoformat()
        meta: dict[str, Any] = {
            "pim_catalog_seed": True,
            "module": "pim",
            "import_area": "product_master",
            "source_kind": "file",
            "store_id": store_id,
            "phase": "previewed",
            "mapping_source": mapping_source,
            "column_mapping": dict(column_map),
            "quality_snapshot": quality,
            "started_at": now,
            "finished_at": now,
            "preview_status": "preview_ready",
            "confirm_required": True,
        }
        ins = (
            db.table("raw_report_uploads")
            .insert(
                {
                    "organization_id": organization_id,
                    "file_name": (file_name or "catalog_upload").strip() or "catalog_upload",
                    "report_type": "pim_product_master",
                    "status": "mapped",
                    "metadata": meta,
                    "row_count": int(quality.get("rows_total") or 0),
                }
            )
            .execute()
        )
        if ins.data:
            return str(ins.data[0]["id"])
    except Exception as e:
        log.warning("PIM seed history insert (preview) skipped: %s", e)
    return None


def _pim_seed_history_finalize_apply(
    db: Any,
    seed_session_id: str | None,
    organization_id: str,
    store_id: str,
    *,
    success: bool,
    metrics: dict[str, Any] | None,
    quality: dict[str, Any] | None,
    error_message: str | None,
) -> None:
    if not seed_session_id:
        return
    try:
        uuid.UUID(str(seed_session_id).strip())
    except ValueError:
        return
    sid = str(seed_session_id).strip()
    try:
        res = (
            db.table("raw_report_uploads")
            .select("id,organization_id,metadata")
            .eq("id", sid)
            .eq("organization_id", organization_id)
            .limit(1)
            .execute()
        )
        rows = res.data or []
        if not rows:
            return
        prev_meta = rows[0].get("metadata") if isinstance(rows[0].get("metadata"), dict) else {}
        merged: dict[str, Any] = dict(prev_meta) if isinstance(prev_meta, dict) else {}
        merged["pim_catalog_seed"] = True
        merged.setdefault("module", "pim")
        merged.setdefault("import_area", "product_master")
        merged["store_id"] = store_id
        if error_message == "apply_blocked_by_dirty_rate":
            merged["phase"] = "blocked"
        else:
            merged["phase"] = "completed" if success else "failed"
        merged["apply_finished_at"] = datetime.now(timezone.utc).isoformat()
        if metrics is not None:
            merged["apply_metrics"] = metrics
        if quality is not None:
            merged["apply_quality"] = quality
        row_status = "complete" if success else "failed"
        upd: dict[str, Any] = {
            "metadata": merged,
            "status": row_status,
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }
        if error_message:
            upd["error_message"] = error_message[:8000]
        elif success:
            upd["error_message"] = None
        db.table("raw_report_uploads").update(upd).eq("id", sid).eq("organization_id", organization_id).execute()
    except Exception as e:
        log.warning("PIM seed history finalize (apply) skipped: %s", e)


# Catalog seed / Google Sheets: reject apply when too many rows fail quality classification (preview still returns 200).
PIM_SEED_MAX_DIRTY_RATE = float(os.environ.get("PIM_SEED_MAX_DIRTY_RATE", "0.2"))


def _pim_catalog_seed_confirm_true(raw: str | None) -> bool:
    v = (raw or "").strip().lower()
    return v in ("1", "true", "yes", "on", "confirm")


def _validate_pim_org_store(organization_id: str, store_id: str) -> tuple[str, str]:
    """Accept any 128-bit UUID string (matches TS `isUuidString` / `uuid.UUID`), not RFC version/variant only."""
    oid = (organization_id or "").strip()
    sid = (store_id or "").strip()
    if not oid:
        raise HTTPException(status_code=400, detail="organization_id is required (UUID).")
    if not sid:
        raise HTTPException(status_code=400, detail="store_id is required (imports target store UUID).")
    try:
        uuid.UUID(oid)
    except ValueError:
        raise HTTPException(
            status_code=400,
            detail=f"organization_id must be a valid UUID (got {oid[:48]!r}).",
        ) from None
    try:
        uuid.UUID(sid)
    except ValueError:
        raise HTTPException(
            status_code=400,
            detail=f"store_id must be a valid UUID (got {sid[:48]!r}).",
        ) from None
    return oid, sid


def _pim_is_valid_store_uuid(value: Any) -> bool:
    """True iff `value` is a non-empty UUID string (matches `_validate_pim_org_store`).

    Used as the entry-guard for `_process_pim_seed_row` so PIM never silently
    INSERTs a duplicate `products` row under an unscoped (org, store) pair when
    `metadata.import_store_id` is missing or malformed. See PATCH-01.
    """
    if value is None:
        return False
    s = str(value).strip()
    if not s:
        return False
    try:
        uuid.UUID(s)
    except (ValueError, AttributeError, TypeError):
        return False
    return True


def _empty_pim_seed_metrics() -> dict[str, Any]:
    return {
        "rows_processed": 0,
        "sheets_processed": 0,
        "rows_per_sheet": {},
        "vendors_created": 0,
        "vendors_reused": 0,
        "categories_created": 0,
        "categories_reused": 0,
        "products_created": 0,
        "products_updated": 0,
        "identifiers_created": 0,
        "identifiers_updated": 0,
        "prices_inserted": 0,
        "products_enriched_by_amazon": 0,
        "skipped_no_identity": 0,
        "skipped_ambiguous": 0,
        "skipped_invalid_store_scope": 0,
        "blocked_new_without_seller_sku": 0,
        "errors": [],
        "fields_trimmed": 0,
        "multi_identifier_cells_split": 0,
        "identifier_tokens_accepted": 0,
        "identifier_tokens_rejected": 0,
        "duplicate_identifier_tokens_collapsed": 0,
        "ambiguous_multi_identifier_rows": 0,
        "multi_identifier_rows_allowed": 0,
        "multi_identifier_rows_conflicting": 0,
        "canonical_products_from_multi_id_rows": 0,
        "identifier_tokens_attached": 0,
        "conflict_rows_blocked": 0,
        "prices_skipped_duplicate": 0,
        "prices_backfill_skipped_no_product": 0,
        "prices_skipped_no_valid_price": 0,
    }


def _pim_seed_quality_acc() -> dict[str, Any]:
    return {
        "rows_total": 0,
        "dirty_rows": 0,
        "blocked_new_without_seller_sku": 0,
        "products_would_update": 0,
        "products_would_create": 0,
        "vendors_would_create": 0,
        "vendors_reused": 0,
        "categories_would_create": 0,
        "categories_reused": 0,
        "prices_would_insert": 0,
        "skipped_no_identity": 0,
        "skipped_ambiguous": 0,
        "invalid_price_rows": 0,
        "preview_errors": [],
        "fields_trimmed": 0,
        "multi_identifier_cells_split": 0,
        "identifier_tokens_accepted": 0,
        "identifier_tokens_rejected": 0,
        "duplicate_identifier_tokens_collapsed": 0,
        "ambiguous_multi_identifier_rows": 0,
        "identifier_map_rows_would_insert": 0,
        "identifier_map_rows_would_update": 0,
        "identifier_rows_would_insert": 0,
        "identifier_rows_would_update": 0,
        "duplicates_reused": 0,
        "existing_products_matched": 0,
        "already_complete": 0,
        "metadata_attributes_detected": 0,
        "multi_identifier_rows_allowed": 0,
        "multi_identifier_rows_conflicting": 0,
        "canonical_products_from_multi_id_rows": 0,
        "identifier_tokens_attached": 0,
        "conflict_rows_blocked": 0,
        "pim_category_debug_samples": [],
        "preview_rows_multi_identifier": 0,
        "preview_category_nonempty_rows": 0,
        "preview_category_simulated_rows": 0,
        "clean_pm_physical_rows_split": 0,
        "clean_pm_synthetic_rows_emitted": 0,
        "conflict_detail": [],  # list[dict] — structured per-row conflict records for UI panel
    }


def _pim_seed_preview_append(acc: dict[str, Any], row_index: int | str | None, message: str, cap: int = 200) -> None:
    pe = acc.setdefault("preview_errors", [])
    if not isinstance(pe, list) or len(pe) >= cap:
        return
    prefix = f"row {row_index}: " if row_index is not None else ""
    pe.append(f"{prefix}{message}")


def _pim_bump_cleaning_counters(
    target: dict[str, Any], parsed: ParsedIdentifierRow, row_trim: int, *, include_accepted: bool
) -> None:
    target["fields_trimmed"] = int(target.get("fields_trimmed") or 0) + int(row_trim)
    target["multi_identifier_cells_split"] = int(target.get("multi_identifier_cells_split") or 0) + int(
        parsed.multi_identifier_cells_split or 0
    )
    if include_accepted:
        target["identifier_tokens_accepted"] = int(target.get("identifier_tokens_accepted") or 0) + int(
            parsed.identifier_tokens_accepted or 0
        )
    target["identifier_tokens_rejected"] = int(target.get("identifier_tokens_rejected") or 0) + int(
        parsed.identifier_tokens_rejected or 0
    )
    target["duplicate_identifier_tokens_collapsed"] = int(target.get("duplicate_identifier_tokens_collapsed") or 0) + int(
        parsed.duplicate_identifier_tokens_collapsed or 0
    )


def _pim_product_ids_for_values(db: Any, organization_id: str, store_id: str, values: list[str], col: str) -> dict[str, list[str]]:
    # NEXT-02b: defense-in-depth scope guard. Without this, store-less SELECTs
    # would silently match nothing and callers would treat the row as "insert"
    # at PIM apply time. Returns empty-list mapping shaped like a no-match run.
    if not _pim_is_valid_store_uuid(organization_id) or not _pim_is_valid_store_uuid(store_id):
        return {value: [] for value in values}

    out: dict[str, list[str]] = {}
    for v in values:
        try:
            r = (
                db.table("products")
                .select("id")
                .eq("organization_id", organization_id)
                .eq("store_id", store_id)
                .eq(col, v)
                .limit(15)
                .execute()
            )
            out[v] = [str(x["id"]) for x in (r.data or [])]
        except Exception:
            out[v] = []
    return out


def _pim_product_ids_for_values_batch(
    db: Any, organization_id: str, store_id: str, values: list[str], col: str
) -> dict[str, list[str]]:
    """One query per batch of distinct identifier values (preview performance)."""
    # NEXT-02b: defense-in-depth scope guard. Same rationale as the single-value
    # variant; preview-time callers in pim_import_async.py reach this directly.
    if not _pim_is_valid_store_uuid(organization_id) or not _pim_is_valid_store_uuid(store_id):
        return {value: [] for value in dict.fromkeys(values)}

    out: dict[str, list[str]] = {}
    uniq = [str(v).strip() for v in dict.fromkeys(values) if str(v).strip()]
    batch_n = 120
    for i in range(0, len(uniq), batch_n):
        chunk = uniq[i : i + batch_n]
        if not chunk:
            continue
        try:
            r = (
                db.table("products")
                .select(f"id,{col}")
                .eq("organization_id", organization_id)
                .eq("store_id", store_id)
                .in_(col, chunk)
                .limit(8000)
                .execute()
            )
            for row in r.data or []:
                key = str(row.get(col) or "").strip()
                pid = str(row.get("id") or "").strip()
                if not key or not pid:
                    continue
                out.setdefault(key, []).append(pid)
                if len(out[key]) > 15:
                    out[key] = out[key][:15]
        except Exception:
            for v in chunk:
                out.setdefault(v, [])
    for v in uniq:
        out.setdefault(v, [])
    return out


def _pim_identifier_preview_detail(parsed: ParsedIdentifierRow) -> str:
    chunks: list[str] = []
    for label, col in (
        ("SKU", parsed.seller_sku),
        ("ASIN", parsed.asin),
        ("FNSKU", parsed.fnsku),
        ("UPC", parsed.upc),
    ):
        if not col.original_display and not col.split_tokens:
            continue
        chunks.append(
            f"{label} cell={col.original_display!r} tokens={col.split_tokens!r} "
            f"accepted={col.accepted!r} rejected={col.rejected!r}"
        )
    return "; ".join(chunks)[:900]


def _pim_finalize_multi_sku(
    db: Any,
    organization_id: str,
    store_id: str,
    parsed: ParsedIdentifierRow,
) -> tuple[ParsedIdentifierRow, str | None]:
    """Multiple seller SKUs in one row: require a single existing product or mark ambiguous."""
    sk = parsed.seller_sku.accepted
    if len(sk) <= 1:
        return parsed, None
    # Parse / finalize_ambiguity may already mark ambiguity — never replace with SKU-only reasons.
    if parsed.ambiguous and parsed.ambiguous_reason != "multiple_seller_sku_requires_single_resolved_product":
        return parsed, None
    matched: dict[str, str] = {}
    for sku in sk:
        try:
            r = (
                db.table("products")
                .select("id")
                .eq("organization_id", organization_id)
                .eq("store_id", store_id)
                .eq("sku", sku)
                .limit(5)
                .execute()
            )
        except Exception:
            return replace(parsed, ambiguous=True, ambiguous_reason="multi_sku_db_lookup_failed"), None
        rows = r.data or []
        if len(rows) > 1:
            return replace(parsed, ambiguous=True, ambiguous_reason="sku_resolves_multiple_products"), None
        if len(rows) == 1:
            matched[sku] = str(rows[0]["id"])
    pids = {v for v in matched.values()}
    if len(pids) > 1:
        return replace(parsed, ambiguous=True, ambiguous_reason="multiple_sku_map_to_different_products"), None
    if len(pids) == 0:
        return (
            replace(parsed, ambiguous=True, ambiguous_reason="multiple_seller_sku_requires_single_resolved_product"),
            None,
        )
    pref = next((s for s in sk if s in matched), sk[0])
    return parsed, pref


def _pim_finalize_multi_sku_batched(
    parsed: ParsedIdentifierRow,
    sku_to_pids: dict[str, list[str]],
) -> tuple[ParsedIdentifierRow, str | None]:
    """Same semantics as _pim_finalize_multi_sku using a pre-fetched sku -> product_id lists map."""
    sk = parsed.seller_sku.accepted
    if len(sk) <= 1:
        return parsed, None
    if parsed.ambiguous and parsed.ambiguous_reason != "multiple_seller_sku_requires_single_resolved_product":
        return parsed, None
    matched: dict[str, str] = {}
    for sku in sk:
        rows = sku_to_pids.get(sku, [])
        if len(rows) > 1:
            return replace(parsed, ambiguous=True, ambiguous_reason="sku_resolves_multiple_products"), None
        if len(rows) == 1:
            matched[sku] = str(rows[0])
    pids = {v for v in matched.values()}
    if len(pids) > 1:
        return replace(parsed, ambiguous=True, ambiguous_reason="multiple_sku_map_to_different_products"), None
    if len(pids) == 0:
        return (
            replace(parsed, ambiguous=True, ambiguous_reason="multiple_seller_sku_requires_single_resolved_product"),
            None,
        )
    pref = next((s for s in sk if s in matched), sk[0])
    return parsed, pref


def _pim_imap_distinct_product_ids_for_tokens(
    db: Any,
    organization_id: str,
    store_id: str,
    seller_skus: list[str],
    asins: list[str],
    fnskus: list[str],
    upcs: list[str],
) -> set[str]:
    """Union of product_id values in product_identifier_map for any listed token (batched)."""
    out: set[str] = set()
    batch_n = 100
    for col, vals in (
        ("seller_sku", seller_skus),
        ("asin", asins),
        ("fnsku", fnskus),
        ("upc_code", upcs),
    ):
        uniq = ordered_unique([str(v).strip() for v in vals if str(v).strip()])
        if not uniq:
            continue
        for i in range(0, len(uniq), batch_n):
            chunk = uniq[i : i + batch_n]
            try:
                r = (
                    db.table("product_identifier_map")
                    .select("product_id")
                    .eq("organization_id", organization_id)
                    .eq("store_id", store_id)
                    .in_(col, chunk)
                    .limit(5000)
                    .execute()
                )
                for row in r.data or []:
                    pid = str(row.get("product_id") or "").strip()
                    if pid:
                        out.add(pid)
            except Exception:
                log.exception("pim_product_master imap lookup failed col=%s", col)
    return out


def _pim_pm_singleton_pids_union_for_tokens(
    tokens: list[str],
    col_map: dict[str, list[str]],
) -> tuple[set[str], str | None]:
    """Each token maps to 0 or 1 product_id in col_map; return union of singletons or ambiguity reason."""
    out: set[str] = set()
    for raw in tokens:
        t = str(raw).strip()
        if not t:
            continue
        rows = col_map.get(t, [])
        if len(rows) > 1:
            return set(), "product_master_token_resolves_multiple_products"
        if len(rows) == 1:
            out.add(str(rows[0]))
    return out, None


def _pim_product_master_resolve_multi_sku_row(
    db: Any,
    organization_id: str,
    store_id: str,
    parsed: ParsedIdentifierRow,
    sku_to_pids: dict[str, list[str]],
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
    acc: dict[str, Any] | None,
    *,
    asin_to_pids: dict[str, list[str]] | None = None,
    fnsku_to_pids: dict[str, list[str]] | None = None,
    upc_to_pids: dict[str, list[str]] | None = None,
) -> tuple[ParsedIdentifierRow, str | None]:
    """
    PIM Product Master (import_mode=product_master): multiple seller SKUs may unify only when
    every accepted SKU/ASIN/FNSKU/UPC resolves via products table (and optionally identifier_map)
    to exactly one org+store product_id. No name/vendor/category heuristics.
    """
    _ = row_cells, handles  # reserved for future diagnostics
    sk = parsed.seller_sku.accepted
    if len(sk) <= 1 or not parsed.ambiguous:
        return parsed, None
    if parsed.ambiguous_reason != "multiple_seller_sku_requires_single_resolved_product":
        return parsed, None

    asin_map = asin_to_pids or {}
    fnsku_map = fnsku_to_pids or {}
    upc_map = upc_to_pids or {}

    matched: dict[str, str] = {}
    for sku in sk:
        key = str(sku).strip()
        rows = sku_to_pids.get(key, [])
        if len(rows) > 1:
            return replace(parsed, ambiguous=True, ambiguous_reason="sku_resolves_multiple_products"), None
        if len(rows) == 1:
            matched[key] = str(rows[0])
    pids_from_skus = {v for v in matched.values()}
    if len(pids_from_skus) > 1:
        return replace(parsed, ambiguous=True, ambiguous_reason="multiple_sku_map_to_different_products"), None

    u_prod = set(pids_from_skus)
    s_asin, ea = _pim_pm_singleton_pids_union_for_tokens(list(parsed.asin.accepted), asin_map)
    if ea:
        return replace(parsed, ambiguous=True, ambiguous_reason=ea), None
    u_prod |= s_asin
    s_fn, ef = _pim_pm_singleton_pids_union_for_tokens(list(parsed.fnsku.accepted), fnsku_map)
    if ef:
        return replace(parsed, ambiguous=True, ambiguous_reason=ef), None
    u_prod |= s_fn
    s_up, eu = _pim_pm_singleton_pids_union_for_tokens(list(parsed.upc.accepted), upc_map)
    if eu:
        return replace(parsed, ambiguous=True, ambiguous_reason=eu), None
    u_prod |= s_up

    if len(u_prod) > 1:
        if acc is not None:
            acc["multi_identifier_rows_conflicting"] = int(acc.get("multi_identifier_rows_conflicting") or 0) + 1
            acc["conflict_rows_blocked"] = int(acc.get("conflict_rows_blocked") or 0) + 1
        return replace(
            parsed,
            ambiguous=True,
            ambiguous_reason="product_master_identifiers_resolve_to_different_products",
        ), None

    if len(u_prod) == 1:
        pid = next(iter(u_prod))
        pref = next((str(s) for s in sk if str(s).strip() in matched and matched[str(s).strip()] == pid), None)
        if not pref:
            try:
                pr = (
                    db.table("products")
                    .select("sku")
                    .eq("id", pid)
                    .eq("organization_id", organization_id)
                    .eq("store_id", store_id)
                    .limit(1)
                    .execute()
                )
                db_sku = str(pr.data[0].get("sku") or "").strip() if pr.data else ""
                pref = next((str(s) for s in sk if str(s) == db_sku), None)
            except Exception:
                log.exception("pim_product_master product lookup for unify failed pid=%s", pid)
                pref = None
        if not pref and sk:
            pref = str(sk[0])
        if acc is not None:
            acc["multi_identifier_rows_allowed"] = int(acc.get("multi_identifier_rows_allowed") or 0) + 1
            acc["canonical_products_from_multi_id_rows"] = int(acc.get("canonical_products_from_multi_id_rows") or 0) + 1
            extra = max(0, len(sk) + len(parsed.asin.accepted) + len(parsed.fnsku.accepted) + len(parsed.upc.accepted) - 1)
            acc["identifier_tokens_attached"] = int(acc.get("identifier_tokens_attached") or 0) + extra
        return replace(parsed, ambiguous=False, ambiguous_reason=None), pref

    imap_pids = _pim_imap_distinct_product_ids_for_tokens(
        db,
        organization_id,
        store_id,
        list(sk),
        list(parsed.asin.accepted),
        list(parsed.fnsku.accepted),
        list(parsed.upc.accepted),
    )
    if len(imap_pids) > 1:
        if acc is not None:
            acc["multi_identifier_rows_conflicting"] = int(acc.get("multi_identifier_rows_conflicting") or 0) + 1
            acc["conflict_rows_blocked"] = int(acc.get("conflict_rows_blocked") or 0) + 1
        return replace(
            parsed,
            ambiguous=True,
            ambiguous_reason="product_master_identifier_map_resolves_to_multiple_products",
        ), None
    if len(imap_pids) == 1:
        pid = next(iter(imap_pids))
        pref: str | None = None
        try:
            pr = (
                db.table("products")
                .select("sku")
                .eq("id", pid)
                .eq("organization_id", organization_id)
                .eq("store_id", store_id)
                .limit(1)
                .execute()
            )
            db_sku = str(pr.data[0].get("sku") or "").strip() if pr.data else ""
            pref = next((str(s) for s in sk if str(s) == db_sku), str(sk[0]) if sk else None)
        except Exception:
            log.exception("pim_product_master product lookup for imap unify failed pid=%s", pid)
            pref = str(sk[0]) if sk else None
        if acc is not None:
            acc["multi_identifier_rows_allowed"] = int(acc.get("multi_identifier_rows_allowed") or 0) + 1
            acc["canonical_products_from_multi_id_rows"] = int(acc.get("canonical_products_from_multi_id_rows") or 0) + 1
            extra = max(0, len(sk) + len(parsed.asin.accepted) + len(parsed.fnsku.accepted) + len(parsed.upc.accepted) - 1)
            acc["identifier_tokens_attached"] = int(acc.get("identifier_tokens_attached") or 0) + extra
        return replace(parsed, ambiguous=False, ambiguous_reason=None), pref

    if acc is not None:
        acc["multi_identifier_rows_conflicting"] = int(acc.get("multi_identifier_rows_conflicting") or 0) + 1
        acc["conflict_rows_blocked"] = int(acc.get("conflict_rows_blocked") or 0) + 1
    return replace(
        parsed,
        ambiguous=True,
        ambiguous_reason="product_master_multi_sku_no_unified_db_identity",
    ), None


def _pim_resolve_product_from_cache(
    seller_sku: str | None,
    fnsku: str | None,
    asin: str | None,
    upc: str | None,
    by_sku: dict[str, list[str]],
    by_fnsku: dict[str, list[str]],
    by_asin: dict[str, list[str]],
    by_upc: dict[str, list[str]],
) -> tuple[str | None, str, str | None]:
    """Mirror _pim_resolve_product using batched lookup dicts (values are product id lists)."""
    if seller_sku:
        rows = by_sku.get(seller_sku, [])
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]), "update", None
        return None, "insert", seller_sku
    if fnsku:
        rows = by_fnsku.get(fnsku, [])
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]), "update", None
        return None, "insert", fnsku
    if asin:
        rows = by_asin.get(asin, [])
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]), "update", None
        return None, "insert", asin
    if upc:
        rows = by_upc.get(upc, [])
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]), "update", None
        return None, "insert", upc
    return None, "no_identity", None


def _pim_imap_conflicts_other_cached(
    imap_rows: list[dict[str, Any]],
    preview_pid: str,
    seller_sku: str | None,
    asin: str | None,
    fnsku: str | None,
    upc: str | None,
) -> bool:
    checks: list[tuple[str, str]] = []
    if seller_sku and str(seller_sku).strip():
        checks.append(("seller_sku", str(seller_sku).strip()))
    if asin and str(asin).strip():
        checks.append(("asin", str(asin).strip()))
    if fnsku and str(fnsku).strip():
        checks.append(("fnsku", str(fnsku).strip()))
    if upc and str(upc).strip():
        checks.append(("upc_code", str(upc).strip()))
    pid_s = str(preview_pid)
    for col, val in checks:
        for row in imap_rows:
            if str(row.get(col) or "").strip() != val:
                continue
            rpid = str(row.get("product_id") or "")
            if rpid and rpid != pid_s:
                return True
    return False


def _pim_imap_row_fate_cached(
    rows_for_pid: list[dict[str, Any]],
    product_id: str | None,
    seller_sku: str | None,
    asin: str | None,
    fnsku: str | None,
    upc: str | None,
) -> str:
    """Mirror _pim_identifier_map_row_fate using rows already scoped to org+store (global chunk list ok)."""
    if not (seller_sku or asin or fnsku or upc):
        return "noop"
    if not product_id:
        return "insert"
    cands = rows_for_pid
    rec: dict[str, Any] | None = None
    if seller_sku:
        rec = next((x for x in cands if str(x.get("seller_sku") or "") == str(seller_sku)), None)
    elif fnsku:
        rec = next((x for x in cands if str(x.get("fnsku") or "") == str(fnsku)), None)
    elif asin:
        rec = next((x for x in cands if str(x.get("asin") or "") == str(asin)), None)
    elif upc:
        rec = next((x for x in cands if str(x.get("upc_code") or "") == str(upc)), None)
    if not rec:
        return "insert"
    payload: dict[str, Any] = {}
    if asin and not rec.get("asin"):
        payload["asin"] = asin
    if fnsku and not rec.get("fnsku"):
        payload["fnsku"] = fnsku
    if seller_sku and not rec.get("seller_sku"):
        payload["seller_sku"] = seller_sku
    if upc and not rec.get("upc_code"):
        payload["upc_code"] = upc
    return "update" if payload else "noop"


@dataclass
class PimPreviewChunkDbCache:
    """Batched DB lookups for one CSV preview chunk (avoids per-row PostgREST storms)."""

    by_sku: dict[str, list[str]]
    by_fnsku: dict[str, list[str]]
    by_asin: dict[str, list[str]]
    by_upc: dict[str, list[str]]
    sku_multi: dict[str, list[str]]
    imap_rows: list[dict[str, Any]]
    imap_by_pid: dict[str, list[dict[str, Any]]]


def _pim_preview_load_imap_touching(
    db: Any,
    organization_id: str,
    store_id: str,
    skus: set[str],
    fnskus: set[str],
    asins: set[str],
    upcs: set[str],
) -> list[dict[str, Any]]:
    """Union of identifier_map rows matching any chunk identifier (deduped by id)."""
    seen: set[str] = set()
    out: list[dict[str, Any]] = []
    batch_n = 120
    for col, vals in (
        ("seller_sku", skus),
        ("fnsku", fnskus),
        ("asin", asins),
        ("upc_code", upcs),
    ):
        uniq = [str(v).strip() for v in vals if str(v).strip()]
        if not uniq:
            continue
        for i in range(0, len(uniq), batch_n):
            chunk = uniq[i : i + batch_n]
            try:
                r = (
                    db.table("product_identifier_map")
                    .select("id,product_id,seller_sku,asin,fnsku,upc_code")
                    .eq("organization_id", organization_id)
                    .eq("store_id", store_id)
                    .in_(col, chunk)
                    .limit(8000)
                    .execute()
                )
                for row in r.data or []:
                    rid = str(row.get("id") or "")
                    if rid and rid not in seen:
                        seen.add(rid)
                        out.append(row)
            except Exception:
                log.exception("pim preview imap batch failed col=%s", col)
    return out


def _pim_preview_prepare_chunk_db_cache(
    db: Any,
    organization_id: str,
    store_id: str,
    rows_chunk: list[tuple[int, dict[str, Any], int]],
    handles: dict[str, str | None],
    headers: list[str],
    asin_prefetch: dict[str, list[str]],
    upc_prefetch: dict[str, list[str]],
) -> PimPreviewChunkDbCache:
    """Parse chunk once, batch-fetch products + identifier_map, return caches for per-row analysis."""
    _ = headers  # reserved — row_cells already keyed by header names
    multi_skus: set[str] = set()
    staged: list[tuple[int, dict[str, Any], int, ParsedIdentifierRow]] = []
    for idx, row_cells, row_trim in rows_chunk:
        parsed = parse_identifier_row(row_cells, handles)
        by_asin = {a: asin_prefetch.get(a, []) for a in parsed.asin.accepted}
        by_upc = {u: upc_prefetch.get(u, []) for u in parsed.upc.accepted}
        parsed = finalize_ambiguity_with_db(parsed, product_ids_by_asin=by_asin, product_ids_by_upc=by_upc)
        if len(parsed.seller_sku.accepted) > 1:
            for s in parsed.seller_sku.accepted:
                if str(s).strip():
                    multi_skus.add(str(s).strip())
        staged.append((idx, row_cells, row_trim, parsed))

    sku_multi = (
        _pim_product_ids_for_values_batch(db, organization_id, store_id, list(multi_skus), "sku")
        if multi_skus
        else {}
    )

    skus: set[str] = set()
    fnskus: set[str] = set()
    asins: set[str] = set()
    upcs: set[str] = set()
    for _idx, _row_cells, _row_trim, parsed in staged:
        parsed, sku_resolve_pref = _pim_finalize_multi_sku_batched(parsed, sku_multi)
        if parsed.ambiguous:
            continue
        seller_sku = sku_resolve_pref or parsed.primary_sku()
        asin = parsed.primary_asin()
        fnsku = parsed.primary_fnsku()
        upc = parsed.primary_upc()
        if seller_sku:
            skus.add(str(seller_sku).strip())
        if fnsku:
            fnskus.add(str(fnsku).strip())
        if asin:
            asins.add(str(asin).strip())
        if upc:
            upcs.add(str(upc).strip())

    by_sku = _pim_product_ids_for_values_batch(db, organization_id, store_id, list(skus), "sku") if skus else {}
    by_fnsku = _pim_product_ids_for_values_batch(db, organization_id, store_id, list(fnskus), "fnsku") if fnskus else {}
    by_asin = _pim_product_ids_for_values_batch(db, organization_id, store_id, list(asins), "asin") if asins else {}
    by_upc = _pim_product_ids_for_values_batch(db, organization_id, store_id, list(upcs), "upc_code") if upcs else {}

    imap_skus: set[str] = set()
    imap_fns: set[str] = set()
    imap_asins: set[str] = set()
    imap_upcs: set[str] = set()
    for _idx, _row_cells, _row_trim, parsed in staged:
        parsed, sku_resolve_pref = _pim_finalize_multi_sku_batched(parsed, sku_multi)
        if parsed.ambiguous:
            continue
        seller_sku = sku_resolve_pref or parsed.primary_sku()
        asin = parsed.primary_asin()
        fnsku = parsed.primary_fnsku()
        upc = parsed.primary_upc()
        if not seller_sku and not fnsku and not asin and not upc:
            continue
        _prod_id, resolution, sku_insert = _pim_resolve_product_from_cache(
            seller_sku, fnsku, asin, upc, by_sku, by_fnsku, by_asin, by_upc
        )
        if resolution in ("ambiguous", "no_identity"):
            continue
        if resolution == "insert" and not sku_insert:
            continue
        map_sku = sku_resolve_pref or parsed.primary_sku() or seller_sku or sku_insert
        for v in build_identifier_map_variants(parsed, map_sku):
            if v.get("seller_sku"):
                imap_skus.add(str(v["seller_sku"]).strip())
            if v.get("fnsku"):
                imap_fns.add(str(v["fnsku"]).strip())
            if v.get("asin"):
                imap_asins.add(str(v["asin"]).strip())
            if v.get("upc_code"):
                imap_upcs.add(str(v["upc_code"]).strip())

    imap_rows = _pim_preview_load_imap_touching(
        db, organization_id, store_id, imap_skus, imap_fns, imap_asins, imap_upcs
    )
    imap_by_pid: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in imap_rows:
        pid = str(row.get("product_id") or "")
        if pid:
            imap_by_pid[pid].append(row)

    return PimPreviewChunkDbCache(
        by_sku=by_sku,
        by_fnsku=by_fnsku,
        by_asin=by_asin,
        by_upc=by_upc,
        sku_multi=sku_multi,
        imap_rows=imap_rows,
        imap_by_pid=dict(imap_by_pid),
    )


def _pim_simulate_vendor_id(raw_label: str | None, index: dict[str, str]) -> tuple[str | None, bool]:
    """Return (existing_id_or_none, would_create_new). Mutates index with a placeholder id so preview dedupes rows."""
    name = _pim_normalize_vendor_name(raw_label)
    key = _pim_vendor_index_key(name)
    if key in index:
        return index[key], False
    index[key] = "__pim_preview_vendor__"
    return None, True


def _pim_clean_product_master_row_variants(
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
    *,
    max_variants: int = 48,
) -> list[dict[str, Any]]:
    """Split multi-token identifier columns into one dict per combination (cartesian, capped)."""
    from itertools import product

    parsed = parse_identifier_row(row_cells, handles)
    axes: list[list[tuple[str, str]]] = []
    for attr, hk in (
        ("seller_sku", "seller_sku"),
        ("asin", "asin"),
        ("fnsku", "fnsku"),
        ("upc", "upc"),
    ):
        col = handles.get(hk)
        if not col:
            continue
        part = getattr(parsed, attr)
        acc_tok = [str(x).strip() for x in part.accepted if str(x).strip()]
        if len(acc_tok) <= 1:
            continue
        axes.append([(str(col), a) for a in acc_tok])
    if not axes:
        return [row_cells]
    out: list[dict[str, Any]] = []
    for combo in product(*axes):
        if len(out) >= int(max_variants):
            break
        rc = dict(row_cells)
        for cname, val in combo:
            rc[cname] = val
        out.append(rc)
    return out if out else [row_cells]


def _pim_simulate_category_id(raw_label: str | None, index: dict[str, str]) -> tuple[str | None, bool]:
    """Match `_pim_ensure_category` / DB index: canonical display + normalized key; preview dedupes new keys."""
    disp = canonical_category_display_name(raw_label)
    if not disp:
        return None, False
    key = normalize_category_label_for_key(raw_label)
    if not key:
        return None, False
    if key in index:
        return index[key], False
    index[key] = "__pim_preview_category__"
    return None, True


def _pim_analyze_seed_row_for_quality(
    db: Any,
    organization_id: str,
    store_id: str,
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
    headers: list[str],
    vendor_index: dict[str, str],
    category_index: dict[str, str],
    acc: dict[str, Any],
    row_index: int | str,
    row_trim: int,
    *,
    amazon_creds: Any,
    skip_amazon: bool,
    asin_prefetch: dict[str, list[str]] | None = None,
    upc_prefetch: dict[str, list[str]] | None = None,
    chunk_db_cache: PimPreviewChunkDbCache | None = None,
    pim_import_mode: str | None = None,
    clean_pm_depth: int = 0,
) -> None:
    eff_mode = (pim_import_mode or "generic_raw").strip().lower()

    if clean_pm_depth == 0 and eff_mode == "clean_product_master":
        variants = _pim_clean_product_master_row_variants(row_cells, handles)
        if len(variants) > 1:
            acc["rows_total"] = int(acc["rows_total"]) + 1
            acc["clean_pm_physical_rows_split"] = int(acc.get("clean_pm_physical_rows_split") or 0) + 1
            acc["clean_pm_synthetic_rows_emitted"] = int(acc.get("clean_pm_synthetic_rows_emitted") or 0) + len(variants)
            for i, rc in enumerate(variants):
                sub_lbl = f"{row_index}×{i + 1}"
                _pim_analyze_seed_row_for_quality(
                    db,
                    organization_id,
                    store_id,
                    rc,
                    handles,
                    headers,
                    vendor_index,
                    category_index,
                    acc,
                    sub_lbl,
                    row_trim,
                    amazon_creds=amazon_creds,
                    skip_amazon=skip_amazon,
                    asin_prefetch=asin_prefetch if clean_pm_depth == 0 else None,
                    upc_prefetch=upc_prefetch if clean_pm_depth == 0 else None,
                    chunk_db_cache=chunk_db_cache if clean_pm_depth == 0 else None,
                    pim_import_mode=pim_import_mode,
                    clean_pm_depth=clean_pm_depth + 1,
                )
            return

    if clean_pm_depth == 0:
        acc["rows_total"] = int(acc["rows_total"]) + 1

    parsed = parse_identifier_row(row_cells, handles)
    if clean_pm_depth == 0:
        if len(parsed.seller_sku.accepted) > 1 or int(parsed.multi_identifier_cells_split or 0) > 0:
            acc["preview_rows_multi_identifier"] = int(acc.get("preview_rows_multi_identifier") or 0) + 1
    if asin_prefetch is not None:
        by_asin = {a: asin_prefetch.get(a, []) for a in parsed.asin.accepted}
    else:
        by_asin = _pim_product_ids_for_values(db, organization_id, store_id, parsed.asin.accepted, "asin")
    if upc_prefetch is not None:
        by_upc = {u: upc_prefetch.get(u, []) for u in parsed.upc.accepted}
    else:
        by_upc = _pim_product_ids_for_values(db, organization_id, store_id, parsed.upc.accepted, "upc_code")
    parsed = finalize_ambiguity_with_db(parsed, product_ids_by_asin=by_asin, product_ids_by_upc=by_upc)
    if chunk_db_cache is not None:
        parsed, sku_resolve_pref = _pim_finalize_multi_sku_batched(parsed, chunk_db_cache.sku_multi)
    else:
        parsed, sku_resolve_pref = _pim_finalize_multi_sku(db, organization_id, store_id, parsed)
    if eff_mode == "product_master" and parsed.ambiguous and parsed.ambiguous_reason == "multiple_seller_sku_requires_single_resolved_product":
        sku_map = (
            chunk_db_cache.sku_multi
            if chunk_db_cache is not None
            else _pim_product_ids_for_values_batch(
                db, organization_id, store_id, list(parsed.seller_sku.accepted), "sku"
            )
        )
        if asin_prefetch is not None:
            asin_map = {a: asin_prefetch.get(a, []) for a in parsed.asin.accepted}
        else:
            asin_map = (
                _pim_product_ids_for_values_batch(db, organization_id, store_id, list(parsed.asin.accepted), "asin")
                if parsed.asin.accepted
                else {}
            )
        if upc_prefetch is not None:
            upc_map = {u: upc_prefetch.get(u, []) for u in parsed.upc.accepted}
        else:
            upc_map = (
                _pim_product_ids_for_values_batch(
                    db, organization_id, store_id, list(parsed.upc.accepted), "upc_code"
                )
                if parsed.upc.accepted
                else {}
            )
        if chunk_db_cache is not None:
            fnsku_map = {str(f).strip(): chunk_db_cache.by_fnsku.get(str(f).strip(), []) for f in parsed.fnsku.accepted}
        else:
            fnsku_map = (
                _pim_product_ids_for_values_batch(db, organization_id, store_id, list(parsed.fnsku.accepted), "fnsku")
                if parsed.fnsku.accepted
                else {}
            )
        parsed, sku_resolve_pref = _pim_product_master_resolve_multi_sku_row(
            db,
            organization_id,
            store_id,
            parsed,
            sku_map,
            row_cells,
            handles,
            acc,
            asin_to_pids=asin_map,
            fnsku_to_pids=fnsku_map,
            upc_to_pids=upc_map,
        )
    _pim_bump_cleaning_counters(acc, parsed, row_trim, include_accepted=not parsed.ambiguous)

    if parsed.ambiguous:
        acc["dirty_rows"] = int(acc["dirty_rows"]) + 1
        acc["skipped_ambiguous"] = int(acc["skipped_ambiguous"]) + 1
        acc["ambiguous_multi_identifier_rows"] = int(acc.get("ambiguous_multi_identifier_rows") or 0) + 1
        _pim_seed_preview_append(
            acc,
            row_index,
            f"multi-identifier ambiguous ({parsed.ambiguous_reason}): {_pim_identifier_preview_detail(parsed)}",
        )
        return

    seller_sku = sku_resolve_pref or parsed.primary_sku()
    asin = parsed.primary_asin()
    fnsku = parsed.primary_fnsku()
    upc = parsed.primary_upc()
    cost_raw, _price_handle_key, _price_pick_meta = _pim_pick_row_unit_price(row_cells, handles, eff_mode)

    if not seller_sku and not fnsku and not asin and not upc:
        acc["dirty_rows"] = int(acc["dirty_rows"]) + 1
        acc["skipped_no_identity"] = int(acc["skipped_no_identity"]) + 1
        # Include raw mapped values in the message for debug visibility.
        raw_sku = str(_pim_cell(row_cells, handles.get("seller_sku")) or "").strip()
        raw_asin = str(_pim_cell(row_cells, handles.get("asin")) or "").strip()
        raw_fnsku = str(_pim_cell(row_cells, handles.get("fnsku")) or "").strip()
        raw_upc = str(_pim_cell(row_cells, handles.get("upc")) or "").strip()
        raw_parts = []
        if raw_sku: raw_parts.append(f"raw_sku={raw_sku[:30]!r}")
        if raw_asin: raw_parts.append(f"raw_asin={raw_asin[:30]!r}")
        if raw_fnsku: raw_parts.append(f"raw_fnsku={raw_fnsku[:30]!r}")
        if raw_upc: raw_parts.append(f"raw_upc={raw_upc[:30]!r}")
        detail = ("; ".join(raw_parts)) if raw_parts else _pim_identifier_preview_detail(parsed)
        _pim_seed_preview_append(
            acc,
            row_index,
            f"no_identity_after_cleaning: {detail}",
        )
        return

    if asin and amazon_creds and not skip_amazon:
        _fetch_amazon_catalog_data(asin, amazon_creds)
        time.sleep(0.2)

    if chunk_db_cache is not None:
        prod_id, resolution, sku_insert = _pim_resolve_product_from_cache(
            seller_sku,
            fnsku,
            asin,
            upc,
            chunk_db_cache.by_sku,
            chunk_db_cache.by_fnsku,
            chunk_db_cache.by_asin,
            chunk_db_cache.by_upc,
        )
    else:
        prod_id, resolution, sku_insert = _pim_resolve_product(
            db, organization_id, store_id, seller_sku, fnsku, asin, upc
        )
    if resolution == "ambiguous":
        acc["dirty_rows"] = int(acc["dirty_rows"]) + 1
        acc["skipped_ambiguous"] = int(acc["skipped_ambiguous"]) + 1
        _pim_seed_preview_append(acc, row_index, "ambiguous product match for org+store+identifiers")
        return
    if resolution == "no_identity":
        acc["dirty_rows"] = int(acc["dirty_rows"]) + 1
        acc["skipped_no_identity"] = int(acc["skipped_no_identity"]) + 1
        _pim_seed_preview_append(acc, row_index, "no_identity after resolve")
        return

    if resolution == "insert" and not sku_insert:
        acc["dirty_rows"] = int(acc["dirty_rows"]) + 1
        acc["blocked_new_without_seller_sku"] = int(acc["blocked_new_without_seller_sku"]) + 1
        _pim_seed_preview_append(
            acc,
            row_index,
            "blocked: could not determine primary sku/upc/asin for new product insert",
        )
        return

    if resolution == "insert":
        acc["products_would_create"] = int(acc["products_would_create"]) + 1
    elif resolution == "update":
        acc["products_would_update"] = int(acc["products_would_update"]) + 1
        acc["existing_products_matched"] = int(acc.get("existing_products_matched") or 0) + 1

    map_sku = sku_resolve_pref or parsed.primary_sku() or seller_sku or sku_insert
    variants = build_identifier_map_variants(parsed, map_sku)
    preview_pid = prod_id if resolution == "update" else None
    ident_conflict_dirty = False

    # Holistic row-level conflict analysis:
    # Collect ALL product IDs that this row's identifiers are linked to (excluding preview_pid).
    # - 0 other PIDs  → already linked to same product OR new links → safe (not dirty)
    # - 1 other PID   → all identifiers consistently link to ONE alternative product → safe update (task D)
    # - 2+ other PIDs → true conflict, identifiers point to multiple products → dirty
    if preview_pid:
        pid_s = str(preview_pid)
        other_pids_set: set[str] = set()
        imap_source = chunk_db_cache.imap_rows if chunk_db_cache is not None else []
        for v in variants:
            for _col, _val in (
                ("seller_sku", str(v.get("seller_sku") or "").strip()),
                ("asin", str(v.get("asin") or "").strip()),
                ("fnsku", str(v.get("fnsku") or "").strip()),
                ("upc_code", str(v.get("upc_code") or "").strip()),
            ):
                if not _val:
                    continue
                for imap_row in imap_source:
                    if str(imap_row.get(_col) or "").strip() == _val:
                        rpid = str(imap_row.get("product_id") or "")
                        if rpid and rpid != pid_s:
                            other_pids_set.add(rpid)

        if len(other_pids_set) == 0:
            # All identifiers are either new (not in imap yet) or already linked to preview_pid.
            # Nothing to do — handled by the fate loop below.
            pass
        elif len(other_pids_set) == 1:
            # Task item D: all conflicting identifiers consistently point to ONE other product.
            # Treat as safe update / existing_product_matched — NOT dirty.
            acc["already_complete"] = int(acc.get("already_complete") or 0) + 1
            # Still fall through to the fate loop so insert/update/noop are counted.
        else:
            # Task item E: identifiers point to MULTIPLE different products — true conflict.
            acc["identifier_map_conflicts_preview"] = int(acc.get("identifier_map_conflicts_preview") or 0) + 1
            acc["conflict_rows_blocked"] = int(acc.get("conflict_rows_blocked") or 0) + 1
            ident_conflict_dirty = True
            acc["dirty_rows"] = int(acc["dirty_rows"]) + 1
            other_sorted = sorted(other_pids_set)
            first_two = ", ".join(p[:8] for p in other_sorted[:2])
            _pim_seed_preview_append(
                acc,
                row_index,
                f"true conflict: row identifiers linked to {len(other_pids_set)} different products "
                f"(e.g. {first_two}…) — sku={seller_sku or '—'} asin={asin or '—'}",
            )
            # Structured conflict record for operator UI panel
            try:
                if len(acc.get("conflict_detail", [])) < 500:
                    acc.setdefault("conflict_detail", []).append({
                        "row": str(row_index),
                        "sku": str(seller_sku or ""),
                        "asin": str(asin or ""),
                        "fnsku": str(fnsku or ""),
                        "upc": str(upc or ""),
                        "product_name": str(_pim_cell(row_cells, handles.get("product_name")) or "")[:120],
                        "resolved_pid": str(preview_pid) if preview_pid else None,
                        "conflict_pids": [str(p) for p in other_sorted[:5]],
                        "reason_source": "imap",
                        "recommended": "detach_wrong_imap" if preview_pid else "review",
                    })
            except Exception:
                log.debug("conflict_detail append skipped", exc_info=True)
    elif not preview_pid and variants:
        # Insert row: no conflict check needed, but scan imap for identifiers already in use
        # by other products — this is informational only, not dirty.
        imap_source_ins = chunk_db_cache.imap_rows if chunk_db_cache is not None else []
        ins_other_set: set[str] = set()
        for v in variants:
            for _col, _val in (
                ("seller_sku", str(v.get("seller_sku") or "").strip()),
                ("asin", str(v.get("asin") or "").strip()),
                ("fnsku", str(v.get("fnsku") or "").strip()),
                ("upc_code", str(v.get("upc_code") or "").strip()),
            ):
                if not _val:
                    continue
                for imap_row in imap_source_ins:
                    if str(imap_row.get(_col) or "").strip() == _val:
                        rpid = str(imap_row.get("product_id") or "")
                        if rpid:
                            ins_other_set.add(rpid)
        if len(ins_other_set) > 1:
            # Identifiers for a NEW product point to multiple existing products — true conflict.
            acc["identifier_map_conflicts_preview"] = int(acc.get("identifier_map_conflicts_preview") or 0) + 1
            acc["conflict_rows_blocked"] = int(acc.get("conflict_rows_blocked") or 0) + 1
            ident_conflict_dirty = True
            acc["dirty_rows"] = int(acc["dirty_rows"]) + 1
            other_sorted_ins = sorted(ins_other_set)
            first_two_ins = ", ".join(p[:8] for p in other_sorted_ins[:2])
            _pim_seed_preview_append(
                acc,
                row_index,
                f"true conflict: new-product row identifiers already claim "
                f"{len(ins_other_set)} different products ({first_two_ins}…) — "
                f"sku={seller_sku or '—'} asin={asin or '—'}",
            )
            try:
                if len(acc.get("conflict_detail", [])) < 500:
                    acc.setdefault("conflict_detail", []).append({
                        "row": str(row_index),
                        "sku": str(seller_sku or ""),
                        "asin": str(asin or ""),
                        "fnsku": str(fnsku or ""),
                        "upc": str(upc or ""),
                        "product_name": str(_pim_cell(row_cells, handles.get("product_name")) or "")[:120],
                        "resolved_pid": None,
                        "conflict_pids": [str(p) for p in other_sorted_ins[:5]],
                        "reason_source": "imap",
                        "recommended": "review",
                    })
            except Exception:
                log.debug("conflict_detail insert-row append skipped", exc_info=True)

    for v in variants:
        if ident_conflict_dirty:
            continue
        if chunk_db_cache is not None:
            rows_for = chunk_db_cache.imap_by_pid.get(str(preview_pid), []) if preview_pid else []
            fate = _pim_imap_row_fate_cached(
                rows_for,
                preview_pid,
                v.get("seller_sku"),
                v.get("asin"),
                v.get("fnsku"),
                v.get("upc_code"),
            )
        else:
            fate = _pim_identifier_map_row_fate(
                db,
                organization_id,
                store_id,
                preview_pid,
                v.get("seller_sku"),
                v.get("asin"),
                v.get("fnsku"),
                v.get("upc_code"),
            )
        if fate == "insert":
            acc["identifier_map_rows_would_insert"] = int(acc.get("identifier_map_rows_would_insert") or 0) + 1
        elif fate == "update":
            acc["identifier_map_rows_would_update"] = int(acc.get("identifier_map_rows_would_update") or 0) + 1
        elif fate == "noop":
            acc["duplicates_reused"] = int(acc.get("duplicates_reused") or 0) + 1

    if not ident_conflict_dirty:
        mapped_headers = {str(v).strip() for v in handles.values() if v}
        extra_attrs = collect_product_attributes_from_row(
            row_cells, mapped_headers=mapped_headers, all_headers=headers
        )
        if extra_attrs:
            acc["metadata_attributes_detected"] = int(acc.get("metadata_attributes_detected") or 0) + 1

        v_raw = _clean_identifier(_pim_cell(row_cells, handles.get("vendor")))
        _, wv = _pim_simulate_vendor_id(v_raw, vendor_index)
        if wv:
            acc["vendors_would_create"] = int(acc["vendors_would_create"]) + 1
        elif v_raw:
            acc["vendors_reused"] = int(acc.get("vendors_reused") or 0) + 1

        c_raw = _clean_identifier(_pim_cell(row_cells, handles.get("category")))
        if c_raw:
            acc["preview_category_nonempty_rows"] = int(acc.get("preview_category_nonempty_rows") or 0) + 1
        _, wc = _pim_simulate_category_id(c_raw, category_index)
        acc["preview_category_simulated_rows"] = int(acc.get("preview_category_simulated_rows") or 0) + 1
        if wc:
            acc["categories_would_create"] = int(acc["categories_would_create"]) + 1
        elif c_raw:
            acc["categories_reused"] = int(acc.get("categories_reused") or 0) + 1

        dbg = acc.get("pim_category_debug_samples")
        if isinstance(dbg, list) and c_raw and len(dbg) < 18:
            dbg.append(
                {
                    "raw": str(c_raw)[:200],
                    "normalized_key": normalize_category_label_for_key(c_raw),
                    "would_create_preview": bool(wc),
                }
            )

    if cost_raw is not None:
        s = str(cost_raw).strip()
        if s and s.lower() not in ("x", "n/a", "na", "-", "none", "null"):
            try:
                float(s.replace("$", "").replace(",", ""))
                acc["prices_would_insert"] = int(acc["prices_would_insert"]) + 1
            except (TypeError, ValueError):
                # Invalid price/cost does NOT mark the row dirty — identifiers and product data
                # are still valid; only the price column is unparseable. Track separately.
                acc["invalid_price_rows"] = int(acc["invalid_price_rows"]) + 1
                _pim_seed_preview_append(acc, row_index, f"invalid price/cost value (skipped): {s[:40]!r}")


def _pim_seed_quality_finalize(acc: dict[str, Any]) -> dict[str, Any]:
    total = max(int(acc.get("rows_total") or 0), 1)
    dirty = int(acc.get("dirty_rows") or 0)
    dr = round(float(dirty) / float(total), 4)
    pe = acc.get("preview_errors") or []
    rejected_sample: list[dict[str, str]] = []
    if isinstance(pe, list):
        for msg in pe[:200]:
            if isinstance(msg, str):
                rejected_sample.append({"message": msg})
    rows_total = int(acc.get("rows_total") or 0)
    skipped = int(acc.get("skipped_no_identity") or 0) + int(acc.get("skipped_ambiguous") or 0)
    ins_map = int(acc.get("identifier_map_rows_would_insert") or 0)
    upd_map = int(acc.get("identifier_map_rows_would_update") or 0)
    cat_samples = acc.get("pim_category_debug_samples")
    cat_dbg_out: dict[str, Any] | None = None
    if isinstance(cat_samples, list) and cat_samples:
        cat_dbg_out = {"samples": cat_samples[:25]}
    conflict_blocked = int(acc.get("conflict_rows_blocked") or 0)
    acc_out = {
        k: v
        for k, v in acc.items()
        if k != "pim_category_debug_samples" and not str(k).startswith("_pim_")
    }
    rows_accepted_est = max(rows_total - dirty, 0)
    cat_wc = int(acc.get("categories_would_create") or 0)
    cat_ru = int(acc.get("categories_reused") or 0)
    if cat_wc == 0 and cat_ru == 0:
        if cat_dbg_out is None:
            cat_dbg_out = {}
        reasons: list[str] = []
        if not bool(acc.get("_pim_scan_has_category_handle")):
            reasons.append("no_category_column_mapped")
        elif rows_accepted_est <= 0:
            reasons.append("no_rows_accepted_past_identity_and_resolve")
        elif int(acc.get("preview_category_nonempty_rows") or 0) == 0:
            reasons.append("category_cells_empty_on_accepted_rows")
        elif int(acc.get("preview_category_simulated_rows") or 0) == 0:
            reasons.append("category_metrics_skipped_other_row_issues")
        cat_dbg_out["why_categories_zero"] = reasons
    return {
        **acc_out,
        "identifier_rows_would_insert": ins_map,
        "identifier_rows_would_update": upd_map,
        "rows_skipped": skipped,
        "ambiguous_rows": int(acc.get("skipped_ambiguous") or 0),
        "dirty_rate": dr,
        "apply_blocked_by_dirty_rate": dr > PIM_SEED_MAX_DIRTY_RATE,
        "apply_blocked_by_conflicts": conflict_blocked > 0,
        "max_dirty_rate": PIM_SEED_MAX_DIRTY_RATE,
        "category_import_debug": cat_dbg_out,
        "rejected_sample": rejected_sample,
        "accepted_sample": [],
        "rows_accepted_estimate": rows_accepted_est,
        "skipped_dirty_row": dirty,
        "would_create_products": int(acc.get("products_would_create") or 0),
        "would_update_products": int(acc.get("products_would_update") or 0),
        "ambiguous_matches": int(acc.get("skipped_ambiguous") or 0),
        "multi_identifier_rows_allowed": int(acc.get("multi_identifier_rows_allowed") or 0),
        "multi_identifier_rows_conflicting": int(acc.get("multi_identifier_rows_conflicting") or 0),
        "canonical_products_from_multi_id_rows": int(acc.get("canonical_products_from_multi_id_rows") or 0),
        "identifier_tokens_attached": int(acc.get("identifier_tokens_attached") or 0),
        "conflict_rows_blocked": int(acc.get("conflict_rows_blocked") or 0),
        "invalid_price_rows": int(acc.get("invalid_price_rows") or 0),
        "already_complete": int(acc.get("already_complete") or 0),
        "conflict_detail": acc.get("conflict_detail", [])[:500],
    }


def _pim_seed_scan_prepared_quality(
    db: Any,
    organization_id: str,
    store_id: str,
    prepared: list[dict[str, Any]],
    amazon_creds: Any,
    *,
    skip_amazon: bool,
) -> dict[str, Any]:
    vendor_index = _pim_load_vendor_index(db, organization_id)
    category_index = _pim_load_category_index(db, organization_id)
    acc = _pim_seed_quality_acc()
    for block in prepared:
        df = block["df"]
        headers = block["headers"]
        handles = block["handles"]
        sl = block.get("sheet_label")
        sl_str = str(sl).strip() if sl is not None else ""
        for i, (_, row) in enumerate(df.iterrows()):
            row_cells, row_trim = _pim_row_cells_from_series(row, headers)
            row_lbl: int | str = f"{sl_str}!{i + 2}" if sl_str else i + 2
            _pim_analyze_seed_row_for_quality(
                db,
                organization_id,
                store_id,
                row_cells,
                handles,
                headers,
                vendor_index,
                category_index,
                acc,
                row_lbl,
                row_trim,
                amazon_creds=amazon_creds,
                skip_amazon=skip_amazon,
            )
    return _pim_seed_quality_finalize(acc)


def _pim_collect_prepared_from_upload(
    raw_bytes: bytes, fname: str | None, openai_key: str | None
) -> list[dict[str, Any]]:
    prepared: list[dict[str, Any]] = []
    column_map: dict[str, Any] | None = None
    reference_headers: list[str] | None = None
    last_mapping_src = "unknown"
    saw_any_yield = False
    saw_frame = False
    for sheet_label, df in _iter_seed_product_frames(raw_bytes, fname):
        saw_any_yield = True
        original_headers = [str(c).strip() for c in df.columns]
        if not any(h for h in original_headers if h):
            continue
        saw_frame = True
        if column_map is None or reference_headers != original_headers:
            column_map, last_mapping_src = _pim_resolve_column_map(original_headers, openai_key)
            _validate_gpt_column_map(column_map, original_headers)
            reference_headers = list(original_headers)
            mapping_src = last_mapping_src
        else:
            mapping_src = "reused_header_match"
        assert column_map is not None
        handles = _pim_augment_handles_from_deterministic(
            original_headers, _pim_handles_from_map(column_map, original_headers)
        )
        df = _pim_coerce_mapped_columns_string(df, handles)
        sheet_key = sheet_label or "(csv)"
        match_src = f"etl_seed_products:{sheet_label}" if sheet_label else "etl_seed_products_csv"
        attrs = getattr(df, "attrs", {}) if isinstance(getattr(df, "attrs", None), dict) else {}
        prepared.append(
            {
                "sheet_key": sheet_key,
                "sheet_label": sheet_label or "",
                "df": df,
                "headers": original_headers,
                "handles": handles,
                "match_src": match_src,
                "pim_column_map": _pim_public_column_map(column_map),
                "mapping_source": mapping_src,
                "pim_delimiter_detected": attrs.get("pim_delimiter_detected"),
                "pim_delimiter_uncertain": bool(attrs.get("pim_delimiter_uncertain")),
                "import_file_name": fname,
            }
        )
    if not saw_any_yield:
        raise HTTPException(status_code=400, detail="Empty file.")
    if not saw_frame:
        raise HTTPException(
            status_code=400,
            detail="No readable tabular data: for Excel, add at least one sheet with a header row and data.",
        )
    return prepared


def _pim_google_build_prepared_frames(
    db: Any, organization_id: str, sheet_id: str, creds: Any, openai_key: str | None
) -> list[dict[str, Any]]:
    try:
        from googleapiclient.discovery import build
    except ModuleNotFoundError as e:
        raise HTTPException(
            status_code=500,
            detail=(
                "Google Sheets dependencies are not installed. Run: "
                "python -m pip install google-api-python-client google-auth"
            ),
        ) from e

    service = build("sheets", "v4", credentials=creds)
    sheet_metadata = service.spreadsheets().get(spreadsheetId=sheet_id).execute()
    sheets = sheet_metadata.get("sheets", [])
    prepared: list[dict[str, Any]] = []
    for sheet in sheets:
        sheet_name = sheet["properties"]["title"]
        if sheet_name != PRODUCT_MASTER_SHEET_NAME:
            log.info("Skipping Google Sheet tab tab=%r (only %r is imported)", sheet_name, PRODUCT_MASTER_SHEET_NAME)
            continue
        result = service.spreadsheets().values().get(spreadsheetId=sheet_id, range=f"'{sheet_name}'!A:Z").execute()
        values = result.get("values", [])
        if not values or len(values) < 2:
            log.info("Skipping empty Google Sheet tab tab=%r spreadsheet=%r", sheet_name, sheet_id)
            continue
        headers = [str(c).strip() for c in values[0]]
        rows = [row + [""] * (len(headers) - len(row)) for row in values[1:]]
        df = pd.DataFrame(rows, columns=headers)
        original_headers = headers
        column_map, mapping_src = _pim_resolve_column_map(original_headers, openai_key)
        _validate_gpt_column_map(column_map, original_headers)
        handles = _pim_augment_handles_from_deterministic(
            original_headers, _pim_handles_from_map(column_map, original_headers)
        )
        df = _pim_coerce_mapped_columns_string(df, handles)
        prepared.append(
            {
                "sheet_key": sheet_name,
                "sheet_label": sheet_name,
                "df": df,
                "headers": original_headers,
                "handles": handles,
                "match_src": "etl_google_sheets",
                "pim_column_map": _pim_public_column_map(column_map),
                "mapping_source": mapping_src,
                "import_file_name": f"google_sheet:{sheet_id}",
            }
        )
    if not prepared:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Google Sheet must contain a non-empty tab named {PRODUCT_MASTER_SHEET_NAME!r} "
                "(header row + data). Other tabs are ignored."
            ),
        )
    return prepared


def _pim_append_error(errors: list[str], row_index: int | str | None, message: str, cap: int = 200) -> None:
    if len(errors) >= cap:
        return
    prefix = f"row {row_index}: " if row_index is not None else ""
    errors.append(f"{prefix}{message}")


def _gpt_map_catalog_columns_or_raise(headers: list[str], api_key: str) -> dict[str, Any]:
    if not api_key or not str(api_key).strip():
        raise HTTPException(
            status_code=400,
            detail=(
                "OpenAI API key is required for column mapping. In Settings, save an LLM provider key "
                "(organization_api_keys: role='llm_provider', name like 'OpenAI', api_key=your sk-… secret), "
                "or set OPENAI_API_KEY on the ETL server."
            ),
        )
    try:
        from openai import OpenAI

        client = OpenAI(api_key=api_key)
        prompt = (
            "You are a data mapping AI. Map these CSV/Sheet headers to standard keys: "
            "'vendor', 'category', 'mfg_part', 'product_name', 'seller_sku', 'asin', 'fnsku', 'upc', 'cost', 'status'. "
            "Values must be EXACT header strings from the list (character-for-character match). "
            "Omit a key if no column applies. "
            "Reply ONLY with a valid JSON object mapping standard keys to exact header names. "
            f"Headers: {json.dumps(headers)}"
        )
        res = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.0,
            response_format={"type": "json_object"},
        )
        raw = res.choices[0].message.content
        parsed = json.loads(raw) if raw else {}
        if not isinstance(parsed, dict):
            raise ValueError("OpenAI returned non-object JSON")
        return parsed
    except HTTPException:
        raise
    except Exception as e:
        log.exception("OpenAI column mapping failed")
        raise HTTPException(
            status_code=502,
            detail={
                "error": "column_mapping_failed",
                "message": str(e),
                "headers": headers,
            },
        ) from e


def _validate_gpt_column_map(column_map: dict[str, Any], original_headers: list[str]) -> None:
    inv = set(original_headers)
    bad: list[str] = []
    for std in _PIM_MAP_STANDARD_KEYS:
        v = column_map.get(std)
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        h = str(v).strip()
        if h not in inv:
            bad.append(f"{std} -> {h!r} (not in headers)")
    if bad:
        raise HTTPException(
            status_code=400,
            detail={
                "error": "invalid_column_mapping",
                "message": "Mapped column names must exist exactly in the file headers.",
                "invalid_mappings": bad,
                "headers": original_headers,
                "mapping": {k: column_map.get(k) for k in _PIM_MAP_STANDARD_KEYS if column_map.get(k) is not None},
            },
        )


def _pim_handles_from_map(column_map: dict[str, Any], original_headers: list[str]) -> dict[str, str | None]:
    inv = set(original_headers)
    out: dict[str, str | None] = {k: None for k in _PIM_MAP_STANDARD_KEYS}
    for std in _PIM_MAP_STANDARD_KEYS:
        v = column_map.get(std)
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        h = str(v).strip()
        if h in inv:
            out[std] = h
    return out


def _pim_augment_handles_from_deterministic(
    original_headers: list[str], handles: dict[str, str | None]
) -> dict[str, str | None]:
    """
    Fill missing standard handles (e.g. `category` from a `Category` column) when GPT/JSON
    mapping omitted them but `_pim_try_deterministic_column_map` can infer them.
    """
    inv = set(original_headers)
    det = _pim_try_deterministic_column_map(original_headers)
    if not det:
        return handles
    out = dict(handles)
    for std in _PIM_MAP_STANDARD_KEYS:
        if out.get(std):
            continue
        v = det.get(std)
        if v is None or (isinstance(v, str) and not str(v).strip()):
            continue
        h = str(v).strip()
        if h in inv:
            out[std] = h
    return out


_PIM_VENDOR_WS = re.compile(r"\s+")


def _pim_normalize_vendor_name(raw: str | None) -> str:
    """Trim, collapse internal whitespace; empty -> Unknown Vendor (canonical display name)."""
    s = (raw or "").strip()
    if not s:
        return "Unknown Vendor"
    collapsed = _PIM_VENDOR_WS.sub(" ", s).strip()
    return collapsed or "Unknown Vendor"


def _pim_vendor_index_key(display: str) -> str:
    return display.lower()


def _pim_load_vendor_index(db: Any, organization_id: str) -> dict[str, str]:
    idx: dict[str, str] = {}
    try:
        r = db.table("vendors").select("id,name").eq("organization_id", organization_id).execute()
        for row in r.data or []:
            n = str(row.get("name") or "").strip()
            if not n:
                continue
            display = _pim_normalize_vendor_name(n)
            key = _pim_vendor_index_key(display)
            if key not in idx:
                idx[key] = str(row["id"])
    except Exception as e:
        log.exception("Failed to load vendors for org=%s", organization_id)
        raise HTTPException(status_code=500, detail=f"Failed to load vendors: {e}") from e
    return idx


def _pim_ensure_vendor(
    db: Any,
    organization_id: str,
    raw_label: str | None,
    index: dict[str, str],
    metrics: dict[str, Any],
    errors: list[str],
    row_index: int | str | None,
) -> str | None:
    name = _pim_normalize_vendor_name(raw_label)
    key = _pim_vendor_index_key(name)
    if key in index:
        metrics["vendors_reused"] = int(metrics.get("vendors_reused") or 0) + 1
        return index[key]
    try:
        ins = db.table("vendors").insert({"organization_id": organization_id, "name": name}).execute()
        if not ins.data:
            _pim_append_error(errors, row_index, "vendor insert returned no row")
            return None
        vid = str(ins.data[0]["id"])
        index[key] = vid
        metrics["vendors_created"] += 1
        return vid
    except Exception as e:
        log.warning("vendor insert failed (possible duplicate) org=%s name=%s: %s", organization_id, name, e)
        try:
            r = db.table("vendors").select("id,name").eq("organization_id", organization_id).execute()
            for row in r.data or []:
                n = str(row.get("name") or "").strip()
                if not n:
                    continue
                display = _pim_normalize_vendor_name(n)
                k = _pim_vendor_index_key(display)
                if k not in index:
                    index[k] = str(row["id"])
            if key in index:
                return index[key]
        except Exception as e2:
            log.exception("vendor refetch after insert failure")
            _pim_append_error(errors, row_index, f"vendor resolution failed: {e2}")
            return None
        _pim_append_error(errors, row_index, f"vendor upsert failed: {name}: {e}")
        return None


def _pim_load_category_index(db: Any, organization_id: str) -> dict[str, str]:
    idx: dict[str, str] = {}
    try:
        r = db.table("product_categories").select("id,name").eq("organization_id", organization_id).execute()
        for row in r.data or []:
            n = str(row.get("name") or "").strip()
            if n:
                k = normalize_category_label_for_key(n)
                if k and k not in idx:
                    idx[k] = str(row["id"])
    except Exception as e:
        log.exception("Failed to load product_categories for org=%s", organization_id)
        raise HTTPException(status_code=500, detail=f"Failed to load categories: {e}") from e
    return idx


def _pim_ensure_category(
    db: Any,
    organization_id: str,
    raw_label: str | None,
    index: dict[str, str],
    metrics: dict[str, Any],
    errors: list[str],
    row_index: int | str | None,
) -> str | None:
    name = canonical_category_display_name(raw_label)
    if not name:
        return None
    key = normalize_category_label_for_key(raw_label)
    if key in index:
        metrics["categories_reused"] = int(metrics.get("categories_reused") or 0) + 1
        return index[key]
    try:
        ins = db.table("product_categories").insert({"organization_id": organization_id, "name": name}).execute()
        if not ins.data:
            _pim_append_error(errors, row_index, "category insert returned no row")
            return None
        cid = str(ins.data[0]["id"])
        index[key] = cid
        metrics["categories_created"] += 1
        return cid
    except Exception as e:
        log.warning("category insert failed org=%s name=%s: %s", organization_id, name, e)
        try:
            r = db.table("product_categories").select("id,name").eq("organization_id", organization_id).execute()
            for row in r.data or []:
                n = str(row.get("name") or "").strip()
                if n:
                    nk = normalize_category_label_for_key(n)
                    if nk and nk not in index:
                        index[nk] = str(row["id"])
            if key in index:
                return index[key]
        except Exception as e2:
            log.exception("category refetch failed")
            _pim_append_error(errors, row_index, f"category resolution failed: {e2}")
            return None
        _pim_append_error(errors, row_index, f"category upsert failed: {name}: {e}")
        return None


def _clean_identifier(val: Any) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    if s.lower() in ["x", "n/a", "na", "-", "none", "null", ""] or (len(s) < 2 and s.lower() == "x"):
        return None
    return s


def _pim_cell(row_cells: dict[str, Any], header: str | None) -> Any:
    if not header:
        return None
    return row_cells.get(header)


def _pim_resolve_product(
    db: Any,
    organization_id: str,
    store_id: str,
    seller_sku: str | None,
    fnsku: str | None,
    asin: str | None,
    upc: str | None = None,
) -> tuple[str | None, str, str | None]:
    """Returns (product_id_or_none, resolution, sku_for_insert_or_none).
    resolution: update | insert | ambiguous | no_identity
    Match priority: seller_sku, fnsku, asin, upc (store-scoped).
    """
    # NEXT-02b: defense-in-depth scope guard. Any caller that bypasses
    # `_process_pim_seed_row` (e.g. preview-time helpers in pim_import_async)
    # would otherwise issue store-less SELECTs that miss every existing product
    # and report "insert" — the seed row caller would then INSERT a duplicate.
    # Reuse `_pim_is_valid_store_uuid` (introduced in PATCH-01).
    if not _pim_is_valid_store_uuid(organization_id) or not _pim_is_valid_store_uuid(store_id):
        return None, "no_identity", None

    if seller_sku:
        r = (
            db.table("products")
            .select("id")
            .eq("organization_id", organization_id)
            .eq("store_id", store_id)
            .eq("sku", seller_sku)
            .limit(10)
            .execute()
        )
        rows = r.data or []
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]["id"]), "update", None
        return None, "insert", seller_sku
    if fnsku:
        r = (
            db.table("products")
            .select("id")
            .eq("organization_id", organization_id)
            .eq("store_id", store_id)
            .eq("fnsku", fnsku)
            .limit(10)
            .execute()
        )
        rows = r.data or []
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]["id"]), "update", None
        return None, "insert", fnsku
    if asin:
        r = (
            db.table("products")
            .select("id")
            .eq("organization_id", organization_id)
            .eq("store_id", store_id)
            .eq("asin", asin)
            .limit(20)
            .execute()
        )
        rows = r.data or []
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]["id"]), "update", None
        return None, "insert", asin
    if upc:
        r = (
            db.table("products")
            .select("id")
            .eq("organization_id", organization_id)
            .eq("store_id", store_id)
            .eq("upc_code", upc)
            .limit(25)
            .execute()
        )
        rows = r.data or []
        if len(rows) > 1:
            return None, "ambiguous", None
        if len(rows) == 1:
            return str(rows[0]["id"]), "update", None
        return None, "insert", upc
    return None, "no_identity", None


def _pim_identifier_map_conflicts_other_product(
    db: Any,
    organization_id: str,
    store_id: str,
    product_id: str,
    seller_sku: str | None,
    asin: str | None,
    fnsku: str | None,
    upc: str | None,
) -> bool:
    """True if any non-empty identifier is already linked to a different product in this store."""
    try:
        checks: list[tuple[str, str]] = []
        if seller_sku and str(seller_sku).strip():
            checks.append(("seller_sku", str(seller_sku).strip()))
        if asin and str(asin).strip():
            checks.append(("asin", str(asin).strip()))
        if fnsku and str(fnsku).strip():
            checks.append(("fnsku", str(fnsku).strip()))
        if upc and str(upc).strip():
            checks.append(("upc_code", str(upc).strip()))
        for col, val in checks:
            r = (
                db.table("product_identifier_map")
                .select("id,product_id")
                .eq("organization_id", organization_id)
                .eq("store_id", store_id)
                .eq(col, val)
                .limit(5)
                .execute()
            )
            for row in r.data or []:
                if str(row.get("product_id") or "") != str(product_id):
                    return True
    except Exception:
        return True
    return False


def _pim_identifier_map_row_fate(
    db: Any,
    organization_id: str,
    store_id: str,
    product_id: str | None,
    seller_sku: str | None,
    asin: str | None,
    fnsku: str | None,
    upc: str | None,
) -> str:
    """Preview projection: insert | update | noop (mirrors _pim_upsert_identifier_map lookup/update rules)."""
    if not (seller_sku or asin or fnsku or upc):
        return "noop"
    if not product_id:
        return "insert"
    try:
        q = (
            db.table("product_identifier_map")
            .select("id", "seller_sku", "asin", "fnsku", "upc_code")
            .eq("organization_id", organization_id)
            .eq("store_id", store_id)
            .eq("product_id", product_id)
        )
        if seller_sku:
            q = q.eq("seller_sku", seller_sku)
        elif fnsku:
            q = q.eq("fnsku", fnsku)
        elif asin:
            q = q.eq("asin", asin)
        elif upc:
            q = q.eq("upc_code", upc)
        existing = q.limit(5).execute()
        rows = existing.data or []
        if not rows:
            return "insert"
        rec = rows[0]
        payload: dict[str, Any] = {}
        if asin and not rec.get("asin"):
            payload["asin"] = asin
        if fnsku and not rec.get("fnsku"):
            payload["fnsku"] = fnsku
        if seller_sku and not rec.get("seller_sku"):
            payload["seller_sku"] = seller_sku
        if upc and not rec.get("upc_code"):
            payload["upc_code"] = upc
        return "update" if payload else "noop"
    except Exception:
        return "insert"


def _pim_upsert_identifier_map(
    db: Any,
    organization_id: str,
    store_id: str,
    product_id: str,
    seller_sku: str | None,
    asin: str | None,
    fnsku: str | None,
    upc: str | None,
    match_source: str,
    metrics: dict[str, Any],
    errors: list[str],
    row_index: int | str,
) -> None:
    if not (seller_sku or asin or fnsku or upc):
        return
    # Holistic conflict check: collect all OTHER product_ids linked to any of this row's identifiers.
    # Mirrors the preview-path check: only true conflicts (>1 other product) block insertion.
    try:
        other_pids: set[str] = set()
        pid_s = str(product_id)
        for _col, _val in (
            ("seller_sku", str(seller_sku or "").strip()),
            ("asin", str(asin or "").strip()),
            ("fnsku", str(fnsku or "").strip()),
            ("upc_code", str(upc or "").strip()),
        ):
            if not _val:
                continue
            r = (
                db.table("product_identifier_map")
                .select("product_id")
                .eq("organization_id", organization_id)
                .eq("store_id", store_id)
                .eq(_col, _val)
                .is_("deleted_at", "null")
                .limit(10)
                .execute()
            )
            for row in r.data or []:
                rpid = str(row.get("product_id") or "")
                if rpid and rpid != pid_s:
                    other_pids.add(rpid)
        if len(other_pids) > 1:
            # True conflict: identifiers point to 2+ different products — skip
            _pim_append_error(
                errors,
                row_index,
                f"identifier map skipped: identifiers claim {len(other_pids)} different products in this store",
            )
            return
        # 0 or 1 other_pids: safe to upsert (either no conflict or single consistent alternative)
    except Exception as e:
        log.warning("pim imap holistic conflict check failed product_id=%s: %s", product_id, e)
    try:
        q = (
            db.table("product_identifier_map")
            .select("*")
            .eq("organization_id", organization_id)
            .eq("store_id", store_id)
            .eq("product_id", product_id)
        )
        if seller_sku:
            q = q.eq("seller_sku", seller_sku)
        elif fnsku:
            q = q.eq("fnsku", fnsku)
        elif asin:
            q = q.eq("asin", asin)
        elif upc:
            q = q.eq("upc_code", upc)
        existing = q.limit(5).execute()
        rows = existing.data or []
        if not rows:
            db.table("product_identifier_map").insert(
                {
                    "organization_id": organization_id,
                    "store_id": store_id,
                    "product_id": product_id,
                    "seller_sku": seller_sku,
                    "asin": asin,
                    "fnsku": fnsku,
                    "upc_code": upc,
                    "match_source": match_source,
                }
            ).execute()
            metrics["identifiers_created"] += 1
            return
        rec = rows[0]
        mid = str(rec["id"])
        payload: dict[str, Any] = {}
        if asin and not rec.get("asin"):
            payload["asin"] = asin
        if fnsku and not rec.get("fnsku"):
            payload["fnsku"] = fnsku
        if seller_sku and not rec.get("seller_sku"):
            payload["seller_sku"] = seller_sku
        if upc and not rec.get("upc_code"):
            payload["upc_code"] = upc
        if payload:
            db.table("product_identifier_map").update(payload).eq("id", mid).eq("organization_id", organization_id).eq(
                "store_id", store_id
            ).execute()
            metrics["identifiers_updated"] += 1
    except Exception as e:
        log.exception("product_identifier_map upsert failed product_id=%s", product_id)
        _pim_append_error(errors, row_index, f"identifier map: {e}")


def _pim_norm_header_key(key: str) -> str:
    return str(key).strip().lower().replace(" ", "_").replace("-", "_").replace(".", "_")


def _pim_try_parse_positive_number(cost_raw: Any) -> float | None:
    if cost_raw is None:
        return None
    s = str(cost_raw).strip()
    if not s or s.lower() in ("x", "n/a", "na", "-", "none", "null"):
        return None
    try:
        v = float(s.replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        return None
    if v <= 0 or v != v:  # NaN
        return None
    return v


def _pim_currency_from_row(row_cells: dict[str, Any], handles: dict[str, str | None]) -> str:
    for logical in ("currency", "iso_currency", "price_currency"):
        h = handles.get(logical)
        if not h:
            continue
        raw = _pim_cell(row_cells, h)
        if raw is None:
            continue
        c = str(raw).strip().upper()
        if len(c) == 3 and c.isalpha():
            return c
    for col, raw in row_cells.items():
        nh = _pim_norm_header_key(col)
        if "currency" not in nh and nh not in ("iso", "ccy"):
            continue
        c = str(raw or "").strip().upper()
        if len(c) == 3 and c.isalpha():
            return c
    return "USD"


def _pim_find_pack_units_per_case(row_cells: dict[str, Any]) -> float | None:
    """Units per case / pack size from common Product Master column names (optional case→unit cost)."""
    for col, raw in row_cells.items():
        nh = _pim_norm_header_key(col)
        if "case" not in nh:
            continue
        if "cost" in nh or "price" in nh:
            continue
        if not any(x in nh for x in ("unit", "qty", "quantity", "pack", "count", "size", "item")):
            continue
        v = _pim_try_parse_positive_number(raw)
        if v is not None and v > 1:
            return v
    return None


def _pim_pick_product_master_unit_price(
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
) -> tuple[Any | None, str | None, dict[str, Any]]:
    """Product Master: priority unit-level cost; optional case÷pack. Returns (raw_for_insert, key, meta_extra)."""
    meta_extra: dict[str, Any] = {}
    logical_priority: list[tuple[str, tuple[str, ...]]] = [
        ("selling_unit_cost_without_freight", ("selling_unit_cost_without_freight", "selling_unit_cost_wo_freight")),
        ("selling_unit_cost", ("selling_unit_cost", "unit_selling_cost")),
        ("case_cost_without_freight", ("case_cost_without_freight", "case_cost_wo_freight")),
        ("case_cost", ("case_cost", "casecost")),
    ]
    for label, logical_keys in logical_priority:
        for lk in logical_keys:
            h = handles.get(lk)
            if not h:
                continue
            raw = _pim_cell(row_cells, h)
            if raw is None:
                continue
            s = str(raw).strip()
            if not s or s.lower() in ("x", "n/a", "na", "-", "none", "null"):
                continue
            meta_extra["matched_via"] = "column_map"
            if label in ("case_cost", "case_cost_without_freight"):
                pack = _pim_find_pack_units_per_case(row_cells)
                case_amt = _pim_try_parse_positive_number(raw)
                if pack and pack > 1 and case_amt is not None:
                    unit = case_amt / pack
                    if unit > 0:
                        meta_extra["unit_derived_from_case"] = True
                        meta_extra["case_pack_units"] = pack
                        meta_extra["case_cost_value"] = case_amt
                        return (f"{unit:.10f}".rstrip("0").rstrip("."), f"{label}/per_unit_from_case", meta_extra)
            return raw, label, meta_extra

    header_rules: list[tuple[str, Any]] = [
        (
            "selling_unit_cost_without_freight",
            lambda nh: all(x in nh for x in ("selling", "unit", "cost", "without", "freight")),
        ),
        (
            "selling_unit_cost",
            lambda nh: "selling" in nh and "unit" in nh and "cost" in nh and not ("without" in nh and "freight" in nh),
        ),
        (
            "case_cost_without_freight",
            lambda nh: "case" in nh and "cost" in nh and "without" in nh and "freight" in nh,
        ),
        (
            "case_cost",
            lambda nh: "case" in nh and "cost" in nh and not ("without" in nh and "freight" in nh),
        ),
    ]
    for label, pred in header_rules:
        for col_key, raw in row_cells.items():
            nh = _pim_norm_header_key(col_key)
            if not pred(nh):
                continue
            if raw is None:
                continue
            s = str(raw).strip()
            if not s or s.lower() in ("x", "n/a", "na", "-", "none", "null"):
                continue
            meta_extra["matched_via"] = "header"
            meta_extra["source_header"] = str(col_key).strip()
            if label in ("case_cost", "case_cost_without_freight"):
                pack = _pim_find_pack_units_per_case(row_cells)
                case_amt = _pim_try_parse_positive_number(raw)
                if pack and pack > 1 and case_amt is not None:
                    unit = case_amt / pack
                    if unit > 0:
                        meta_extra["unit_derived_from_case"] = True
                        meta_extra["case_pack_units"] = pack
                        meta_extra["case_cost_value"] = case_amt
                        return (f"{unit:.10f}".rstrip("0").rstrip("."), f"{label}/per_unit_from_case", meta_extra)
            return raw, label, meta_extra

    fb_raw, fb_key = _pim_pick_price_raw_from_handles(row_cells, handles)
    if fb_raw is not None:
        meta_extra["matched_via"] = "fallback_generic_cost"
    return fb_raw, fb_key, meta_extra


def _pim_pick_row_unit_price(
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
    eff_mode: str,
) -> tuple[Any | None, str | None, dict[str, Any]]:
    em = eff_mode.strip().lower()
    if em in ("product_master", "clean_product_master"):
        raw, key, extra = _pim_pick_product_master_unit_price(row_cells, handles)
        if raw is not None:
            return raw, key, extra
    a, b = _pim_pick_price_raw_from_handles(row_cells, handles)
    return a, b, {}


def _pim_pick_price_raw_from_handles(
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
) -> tuple[Any | None, str | None]:
    """First non-empty mapped price/cost column (Product Master files vary: cost, case_cost, unit_cost, …)."""
    for logical in ("cost", "price", "unit_cost", "case_cost", "last_cost", "list_price", "msrp"):
        h = handles.get(logical)
        if not h:
            continue
        raw = _pim_cell(row_cells, h)
        if raw is None:
            continue
        s = str(raw).strip()
        if not s or s.lower() in ("x", "n/a", "na", "-", "none", "null"):
            continue
        return raw, logical
    # Unmapped columns: match common header tokens (preview parity with apply when AI mapping misses a cost column).
    skip_substr = ("asin", "sku", "upc", "fnsku", "ean", "gtin", "mpn", "note", "desc", "url", "image", "photo")
    for key, raw in row_cells.items():
        if raw is None:
            continue
        kn = str(key).strip().lower().replace(" ", "_").replace("-", "_")
        if not kn:
            continue
        if any(x in kn for x in skip_substr):
            continue
        hit = None
        for token in (
            "case_cost",
            "casecost",
            "unit_cost",
            "unitcost",
            "last_cost",
            "lastcost",
            "list_price",
            "listprice",
            "dealer_price",
            "dealerprice",
            "msrp",
            "map_price",
            "your_price",
            "sell_price",
            "selling_price",
            "retail",
            "cost",
            "price",
        ):
            if token in kn or kn == token:
                hit = token
                break
        if not hit:
            continue
        s = str(raw).strip()
        if not s or s.lower() in ("x", "n/a", "na", "-", "none", "null"):
            continue
        return raw, f"header:{kn[:72]}"
    return None, None


def _pim_row_brand_from_handles(row_cells: dict[str, Any], handles: dict[str, str | None]) -> str | None:
    """Explicit brand columns only — never treat manufacturer/vendor as brand without Brand_By_MFG mapping."""
    for logical in ("brand", "brand_name"):
        h = handles.get(logical)
        if not h:
            continue
        v = _clean_identifier(_pim_cell(row_cells, h))
        if v:
            return v
    return None


def _pim_brand_from_mfg_map_lookup(
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
    brand_by_mfg_map: dict[str, str] | None,
) -> str | None:
    """Resolve brand using Brand_By_MFG sheet (mfg / manufacturer -> brand)."""
    if not brand_by_mfg_map:
        return None
    for logical in ("mfg", "manufacturer", "mfg_name"):
        h = handles.get(logical)
        if not h:
            continue
        raw = _clean_identifier(_pim_cell(row_cells, h))
        if not raw:
            continue
        k = raw.strip().lower()
        hit = brand_by_mfg_map.get(k)
        if hit and str(hit).strip():
            return str(hit).strip()
    return None


def _pim_load_brand_by_mfg_from_xlsx(path: str) -> dict[str, str]:
    """Read optional ``Brand_By_MFG`` sheet: manufacturer key -> brand label (lowercase keys)."""
    import openpyxl

    out: dict[str, str] = {}
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return out
    try:
        sheet_name = None
        for sn in wb.sheetnames:
            if str(sn).strip().lower() == "brand_by_mfg":
                sheet_name = sn
                break
        if sheet_name is None:
            return out
        ws = wb[sheet_name]
        rows_it = ws.iter_rows(min_row=1, values_only=True)
        header_row = next(rows_it, None)
        if not header_row:
            return out
        headers = [str(c).strip().lower() if c is not None else "" for c in header_row]
        mfg_i = None
        brand_i = None
        for i, h in enumerate(headers):
            if not h:
                continue
            if mfg_i is None and any(x in h for x in ("mfg", "manufacturer", "vendor", "company")):
                mfg_i = i
            if brand_i is None and "brand" in h:
                brand_i = i
        if mfg_i is None:
            mfg_i = 0
        if brand_i is None:
            brand_i = 1 if len(headers) > 1 else 0
        if mfg_i == brand_i:
            return out
        for raw in rows_it:
            if not raw:
                continue
            cells = list(raw)
            if mfg_i >= len(cells) or brand_i >= len(cells):
                continue
            mk = cells[mfg_i]
            bv = cells[brand_i]
            if mk is None or bv is None:
                continue
            ms = str(mk).strip()
            bs = str(bv).strip()
            if not ms or not bs:
                continue
            out[ms.lower()] = bs
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def _pim_reconcile_apply_vs_preview(metrics: dict[str, Any], pq: dict[str, Any] | None) -> dict[str, Any]:
    if not pq or not isinstance(pq, dict):
        return {}
    preview_rows = int(pq.get("rows_accepted_estimate") or pq.get("rows_total") or 0)
    proc = int(metrics.get("rows_processed") or 0)
    return {
        "preview_accepted_rows_estimate": preview_rows,
        "apply_rows_processed": proc,
        "products_created": int(metrics.get("products_created") or 0),
        "products_updated": int(metrics.get("products_updated") or 0),
        "skipped_no_identity": int(metrics.get("skipped_no_identity") or 0),
        "skipped_ambiguous": int(metrics.get("skipped_ambiguous") or 0),
        "vendors_created": int(metrics.get("vendors_created") or 0),
        "vendors_reused": int(metrics.get("vendors_reused") or 0),
        "categories_created": int(metrics.get("categories_created") or 0),
        "categories_reused": int(metrics.get("categories_reused") or 0),
        "identifiers_created": int(metrics.get("identifiers_created") or 0),
        "prices_inserted": int(metrics.get("prices_inserted") or 0),
        "apply_error_events": len(metrics.get("errors") or []),
        "delta_apply_rows_minus_preview_accepted": proc - preview_rows if preview_rows else None,
    }


def _pim_insert_price_if_present(
    db: Any,
    organization_id: str,
    store_id: str,
    product_id: str,
    cost_raw: Any,
    price_source: str,
    metrics: dict[str, Any],
    errors: list[str],
    row_index: int | str,
    product_sku: str | None = None,
    *,
    asin_value: str | None = None,
    source_column: str | None = None,
    pim_upload_id: str | None = None,
    import_file_name: str | None = None,
    import_observed_iso: str | None = None,
    currency: str | None = None,
    row_cells: dict[str, Any] | None = None,
    handles: dict[str, str | None] | None = None,
    price_pick_meta: dict[str, Any] | None = None,
) -> None:
    if cost_raw is None:
        return
    s = str(cost_raw).strip()
    if not s or s.lower() in ("x", "n/a", "na", "-", "none", "null"):
        return
    try:
        amt = float(s.replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        _pim_append_error(errors, row_index, f"invalid price/cost value: {s[:40]!r}")
        return
    if amt <= 0 or amt != amt:
        return
    try:
        src_label = "product_master_import" if str(price_source).strip() == "pim_import_async" else (price_source or "import")
        cur = (currency or "").strip().upper() if currency else ""
        if not cur and row_cells is not None and handles is not None:
            cur = _pim_currency_from_row(row_cells, handles)
        if not cur:
            cur = "USD"

        meta_blob: dict[str, Any] = {
            "price_source": price_source,
            "source_column": source_column,
            "source_column_name": source_column,
            "raw_value": s[:200],
            "seller_sku": str(product_sku).strip() if product_sku and str(product_sku).strip() else None,
            "asin": str(asin_value).strip() if asin_value and str(asin_value).strip() else None,
            "import_file_name": str(import_file_name).strip() if import_file_name and str(import_file_name).strip() else None,
        }
        if price_pick_meta:
            for k, v in price_pick_meta.items():
                if v is not None:
                    meta_blob[k] = v
        if pim_upload_id and str(pim_upload_id).strip():
            uid_s = str(pim_upload_id).strip()
            meta_blob["import_upload_id"] = uid_s
            meta_blob["import_job_id"] = uid_s
        if asin_value and str(asin_value).strip():
            meta_blob["asin"] = str(asin_value).strip()

        obs_iso = import_observed_iso or datetime.now(timezone.utc).isoformat()

        def _price_dup_exists() -> bool:
            from datetime import timedelta

            uid = str(pim_upload_id).strip() if pim_upload_id else ""
            q = (
                db.table("product_prices")
                .select("id,amount,price")
                .eq("organization_id", organization_id)
                .eq("store_id", store_id)
                .eq("product_id", product_id)
                .eq("currency", cur)
                .eq("source", src_label)
            )
            if uid:
                q = q.eq("source_upload_id", uid)
            else:
                cutoff = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()
                q = q.gte("observed_at", cutoff)
            dup_check = q.limit(40).execute()
            for rec in dup_check.data or []:
                for colname in ("amount", "price"):
                    v = rec.get(colname)
                    if v is None:
                        continue
                    try:
                        if abs(float(v) - float(amt)) < 1e-5:
                            return True
                    except (TypeError, ValueError):
                        continue
            return False

        try:
            if _price_dup_exists():
                metrics["prices_skipped_duplicate"] = int(metrics.get("prices_skipped_duplicate") or 0) + 1
                return
        except Exception as dup_err:
            log.debug("price dedup check skipped: %s", dup_err)

        row: dict[str, Any] = {
            "organization_id": organization_id,
            "store_id": store_id,
            "product_id": product_id,
            "amount": amt,
            "price": amt,
            "currency": cur,
            "source": src_label,
            "observed_at": obs_iso,
            "metadata": meta_blob,
        }
        if product_sku and str(product_sku).strip():
            row["sku"] = str(product_sku).strip()
        uid = str(pim_upload_id).strip() if pim_upload_id else ""
        if uid:
            row["source_upload_id"] = uid
        db.table("product_prices").insert(row).execute()
        metrics["prices_inserted"] += 1
    except Exception as e:
        log.exception("product_prices insert failed product_id=%s", product_id)
        _pim_append_error(errors, row_index, f"price insert: {e}")


def _pim_seed_apply_prepared_writes(
    db: Any,
    organization_id: str,
    store_id: str,
    prepared: list[dict[str, Any]],
    amazon_creds: Any,
    price_source: str,
) -> dict[str, Any]:
    metrics = _empty_pim_seed_metrics()
    errors: list[str] = metrics["errors"]
    vendor_index = _pim_load_vendor_index(db, organization_id)
    category_index = _pim_load_category_index(db, organization_id)
    for block in prepared:
        df = block["df"]
        headers = block["headers"]
        handles = block["handles"]
        match_src = str(block["match_src"])
        sheet_key = str(block.get("sheet_key") or "")
        rows_before = int(metrics["rows_processed"])
        metrics["sheets_processed"] = int(metrics["sheets_processed"]) + 1
        sl = block.get("sheet_label")
        sl_str = str(sl).strip() if sl is not None else ""
        fn = str(block.get("import_file_name") or "").strip() or None
        for i, (_, row) in enumerate(df.iterrows()):
            row_cells, row_trim = _pim_row_cells_from_series(row, headers)
            row_lbl: int | str = f"{sl_str}!{i + 2}" if sl_str else i + 2
            _process_pim_seed_row(
                db,
                organization_id,
                store_id,
                row_cells,
                handles,
                headers,
                amazon_creds,
                price_source,
                match_src,
                metrics,
                errors,
                row_lbl,
                vendor_index,
                category_index,
                row_trim,
                import_file_name=fn,
            )
        metrics["rows_per_sheet"][sheet_key] = int(metrics["rows_processed"]) - rows_before
    return metrics


def _process_pim_seed_row(
    db: Any,
    organization_id: str,
    store_id: str,
    row_cells: dict[str, Any],
    handles: dict[str, str | None],
    headers: list[str],
    amazon_creds: Any,
    price_source: str,
    match_source: str,
    metrics: dict[str, Any],
    errors: list[str],
    row_index: int | str,
    vendor_index: dict[str, str],
    category_index: dict[str, str],
    row_trim: int,
    pim_import_mode: str | None = None,
    clean_pm_depth: int = 0,
    *,
    skip_amazon_enrichment: bool = False,
    pim_upload_id: str | None = None,
    brand_by_mfg_map: dict[str, str] | None = None,
    prices_only: bool = False,
    import_file_name: str | None = None,
) -> None:
    # PATCH-01: hard scope guard — refuse rows where org/store scope is missing or
    # malformed. Without this, `_pim_resolve_product` would run with empty filters,
    # miss every existing product (priority 1..4 incl. UPC), fall to "insert", and
    # create duplicate `products` rows under an unscoped (org, store) pair. The
    # guard runs BEFORE clean_pm variant expansion, parse_identifier_row, and any
    # DB read or write — including price/identifier-map writes.
    if not _pim_is_valid_store_uuid(organization_id) or not _pim_is_valid_store_uuid(store_id):
        metrics["skipped_invalid_store_scope"] = int(metrics.get("skipped_invalid_store_scope") or 0) + 1
        metrics["skipped_no_identity"] = int(metrics.get("skipped_no_identity") or 0) + 1
        _pim_append_error(
            errors,
            row_index,
            (
                "invalid_store_scope: row skipped (organization_id="
                f"{str(organization_id)[:48]!r} store_id={str(store_id)[:48]!r} "
                "are not both valid UUIDs)"
            ),
        )
        return

    eff_mode = (pim_import_mode or ("product_master" if "pim_async" in str(match_source) else "generic_raw")).strip().lower()

    if clean_pm_depth == 0 and eff_mode == "clean_product_master":
        variants = _pim_clean_product_master_row_variants(row_cells, handles)
        if len(variants) > 1:
            for i, rc in enumerate(variants):
                _process_pim_seed_row(
                    db,
                    organization_id,
                    store_id,
                    rc,
                    handles,
                    headers,
                    amazon_creds,
                    price_source,
                    match_source,
                    metrics,
                    errors,
                    f"{row_index}×{i + 1}",
                    vendor_index,
                    category_index,
                    row_trim,
                    pim_import_mode=pim_import_mode,
                    clean_pm_depth=clean_pm_depth + 1,
                    skip_amazon_enrichment=skip_amazon_enrichment,
                    pim_upload_id=pim_upload_id,
                    brand_by_mfg_map=brand_by_mfg_map,
                    prices_only=prices_only,
                    import_file_name=import_file_name,
                )
            return

    metrics["rows_processed"] += 1

    parsed = parse_identifier_row(row_cells, handles)
    by_asin = _pim_product_ids_for_values(db, organization_id, store_id, parsed.asin.accepted, "asin")
    by_upc = _pim_product_ids_for_values(db, organization_id, store_id, parsed.upc.accepted, "upc_code")
    parsed = finalize_ambiguity_with_db(parsed, product_ids_by_asin=by_asin, product_ids_by_upc=by_upc)
    parsed, sku_resolve_pref = _pim_finalize_multi_sku(db, organization_id, store_id, parsed)
    if eff_mode == "product_master" and parsed.ambiguous and parsed.ambiguous_reason == "multiple_seller_sku_requires_single_resolved_product":
        sku_map = _pim_product_ids_for_values_batch(
            db, organization_id, store_id, list(parsed.seller_sku.accepted), "sku"
        )
        asin_map = _pim_product_ids_for_values_batch(
            db, organization_id, store_id, list(parsed.asin.accepted), "asin"
        )
        fnsku_map = _pim_product_ids_for_values_batch(
            db, organization_id, store_id, list(parsed.fnsku.accepted), "fnsku"
        )
        upc_map = _pim_product_ids_for_values_batch(
            db, organization_id, store_id, list(parsed.upc.accepted), "upc_code"
        )
        parsed, sku_resolve_pref = _pim_product_master_resolve_multi_sku_row(
            db,
            organization_id,
            store_id,
            parsed,
            sku_map,
            row_cells,
            handles,
            metrics,
            asin_to_pids=asin_map,
            fnsku_to_pids=fnsku_map,
            upc_to_pids=upc_map,
        )
    _pim_bump_cleaning_counters(metrics, parsed, row_trim, include_accepted=not parsed.ambiguous)

    if parsed.ambiguous:
        metrics["skipped_ambiguous"] += 1
        metrics["ambiguous_multi_identifier_rows"] = int(metrics.get("ambiguous_multi_identifier_rows") or 0) + 1
        _pim_append_error(
            errors,
            row_index,
            f"multi-identifier ambiguous ({parsed.ambiguous_reason}): {_pim_identifier_preview_detail(parsed)}",
        )
        return

    seller_sku = sku_resolve_pref or parsed.primary_sku()
    asin = parsed.primary_asin()
    fnsku = parsed.primary_fnsku()
    upc = parsed.primary_upc()
    sheet_product_name = _clean_identifier(_pim_cell(row_cells, handles.get("product_name")))
    mfg_part = _clean_identifier(_pim_cell(row_cells, handles.get("mfg_part")))
    price_raw, price_col_key, price_pick_meta = _pim_pick_row_unit_price(row_cells, handles, eff_mode)
    status_val = normalize_pim_status(_pim_cell(row_cells, handles.get("status")))
    mapped_headers = {str(v).strip() for v in handles.values() if v}
    extra_attrs = collect_product_attributes_from_row(
        row_cells, mapped_headers=mapped_headers, all_headers=headers
    )

    vendor_cell = _clean_identifier(_pim_cell(row_cells, handles.get("vendor")))
    vendor_id = None
    vendor_display_name = _pim_normalize_vendor_name(vendor_cell) if vendor_cell else None
    category_id = None
    if not prices_only:
        vendor_id = _pim_ensure_vendor(
            db,
            organization_id,
            vendor_cell,
            vendor_index,
            metrics,
            errors,
            row_index,
        )
        vendor_display_name = _pim_normalize_vendor_name(vendor_cell) if vendor_cell else None
        category_id = _pim_ensure_category(
            db,
            organization_id,
            _clean_identifier(_pim_cell(row_cells, handles.get("category"))),
            category_index,
            metrics,
            errors,
            row_index,
        )

    if not seller_sku and not fnsku and not asin and not upc:
        metrics["skipped_no_identity"] += 1
        _pim_append_error(
            errors,
            row_index,
            f"skipped: no valid identifiers after split/validate: {_pim_identifier_preview_detail(parsed)}",
        )
        return

    if prices_only:
        prod_id, resolution, sku_insert = _pim_resolve_product(
            db, organization_id, store_id, seller_sku, fnsku, asin, upc
        )
        if resolution != "update" or not prod_id:
            metrics["prices_backfill_skipped_no_product"] = int(metrics.get("prices_backfill_skipped_no_product") or 0) + 1
            return
        if price_raw is None:
            metrics["prices_skipped_no_valid_price"] = int(metrics.get("prices_skipped_no_valid_price") or 0) + 1
            return
        map_sku = sku_resolve_pref or parsed.primary_sku() or seller_sku or (sku_insert or "")
        cur_po = _pim_currency_from_row(row_cells, handles)
        _pim_insert_price_if_present(
            db,
            organization_id,
            store_id,
            str(prod_id),
            price_raw,
            price_source,
            metrics,
            errors,
            row_index,
            product_sku=str(map_sku).strip() if map_sku else None,
            asin_value=asin,
            source_column=price_col_key,
            pim_upload_id=pim_upload_id,
            import_file_name=import_file_name,
            currency=cur_po,
            row_cells=row_cells,
            handles=handles,
            price_pick_meta=price_pick_meta,
        )
        return

    amazon_data: dict[str, Any] = {}
    if not skip_amazon_enrichment and asin and amazon_creds:
        amazon_data = _fetch_amazon_catalog_data(asin, amazon_creds)
        time.sleep(0.2)
        if amazon_data:
            metrics["products_enriched_by_amazon"] += 1
        elif asin:
            _pim_append_error(errors, row_index, f"Amazon enrichment returned no data for ASIN {asin}")

    sheet_brand = _pim_row_brand_from_handles(row_cells, handles)
    mfg_mapped_brand = _pim_brand_from_mfg_map_lookup(row_cells, handles, brand_by_mfg_map)
    effective_brand = sheet_brand or mfg_mapped_brand
    final_name = (
        (amazon_data.get("product_name") if amazon_data else None)
        or sheet_product_name
        or f"Pending Details ({asin or seller_sku or fnsku or upc})"
    )
    final_brand = effective_brand or (amazon_data.get("brand") if amazon_data else None)
    main_image = amazon_data.get("main_image_url")
    amazon_raw = amazon_data.get("amazon_raw") if amazon_data else None

    prod_id, resolution, sku_insert = _pim_resolve_product(
        db, organization_id, store_id, seller_sku, fnsku, asin, upc
    )
    if resolution == "ambiguous":
        metrics["skipped_ambiguous"] += 1
        _pim_append_error(
            errors,
            row_index,
            f"ambiguous product match org+store+identifiers (sku={seller_sku!r} fnsku={fnsku!r} asin={asin!r} upc={upc!r})",
        )
        return
    if resolution == "no_identity":
        metrics["skipped_no_identity"] += 1
        _pim_append_error(errors, row_index, "no_identity: missing resolvable seller_sku, fnsku, asin, and upc")
        return

    sync_iso = datetime.now(timezone.utc).isoformat()

    try:
        if resolution == "update" and prod_id:
            prev_meta: Any = {}
            try:
                prm = (
                    db.table("products")
                    .select("metadata")
                    .eq("id", prod_id)
                    .eq("organization_id", organization_id)
                    .eq("store_id", store_id)
                    .limit(1)
                    .execute()
                )
                if prm.data and isinstance(prm.data[0].get("metadata"), dict):
                    prev_meta = prm.data[0]["metadata"]
            except Exception:
                prev_meta = {}
            merged_meta = merge_product_attributes_into_metadata(prev_meta, extra_attrs)
            merged_meta = {
                **merged_meta,
                "pim_seed": {"source": match_source, "last_import_at": sync_iso},
            }
            upd: dict[str, Any] = {
                "vendor_id": vendor_id,
                "category_id": category_id,
                "mfg_part_number": mfg_part,
                "product_name": final_name,
                "brand": final_brand,
                "main_image_url": main_image,
                "amazon_raw": amazon_raw or {},
                "status": status_val,
                "metadata": merged_meta,
                "last_catalog_sync_at": sync_iso,
            }
            if vendor_display_name:
                upd["vendor_name"] = vendor_display_name
            if seller_sku:
                upd["sku"] = seller_sku
            if asin:
                upd["asin"] = asin
            if fnsku:
                upd["fnsku"] = fnsku
            if upc:
                upd["upc_code"] = upc
            db.table("products").update(upd).eq("id", prod_id).eq("organization_id", organization_id).eq(
                "store_id", store_id
            ).execute()
            metrics["products_updated"] += 1
        elif resolution == "insert" and sku_insert:
            seed_meta = merge_product_attributes_into_metadata({}, extra_attrs)
            seed_meta["pim_seed"] = {"source": match_source, "imported_at": sync_iso}
            ins = (
                db.table("products")
                .insert(
                    {
                        "organization_id": organization_id,
                        "store_id": store_id,
                        "sku": sku_insert,
                        "vendor_id": vendor_id,
                        "vendor_name": vendor_display_name,
                        "category_id": category_id,
                        "mfg_part_number": mfg_part,
                        "product_name": final_name,
                        "brand": final_brand,
                        "main_image_url": main_image,
                        "amazon_raw": amazon_raw if amazon_raw is not None else {},
                        "asin": asin,
                        "fnsku": fnsku,
                        "upc_code": upc,
                        "status": status_val,
                        "metadata": seed_meta,
                        "last_catalog_sync_at": sync_iso,
                    }
                )
                .execute()
            )
            if not ins.data:
                _pim_append_error(errors, row_index, "product insert returned no row")
                return
            prod_id = str(ins.data[0]["id"])
            metrics["products_created"] += 1
        else:
            metrics["skipped_no_identity"] += 1
            metrics["blocked_new_without_seller_sku"] = int(metrics.get("blocked_new_without_seller_sku") or 0) + 1
            _pim_append_error(errors, row_index, "could not determine primary sku for new product insert")
            return
    except Exception as e:
        log.exception("product upsert failed row=%s", row_index)
        _pim_append_error(errors, row_index, f"product upsert: {e}")
        return

    map_sku = sku_resolve_pref or parsed.primary_sku() or seller_sku or sku_insert
    payloads = build_identifier_map_variants(parsed, map_sku)
    if not payloads and (seller_sku or asin or fnsku or upc):
        payloads = [
            {
                "seller_sku": seller_sku or map_sku,
                "asin": asin,
                "fnsku": fnsku,
                "upc_code": upc,
            }
        ]
    for payload in payloads:
        _pim_upsert_identifier_map(
            db,
            organization_id,
            store_id,
            prod_id,
            payload.get("seller_sku"),
            payload.get("asin"),
            payload.get("fnsku"),
            payload.get("upc_code"),
            match_source,
            metrics,
            errors,
            row_index,
        )
    cur_ins = _pim_currency_from_row(row_cells, handles)
    _pim_insert_price_if_present(
        db,
        organization_id,
        store_id,
        prod_id,
        price_raw,
        price_source,
        metrics,
        errors,
        row_index,
        product_sku=map_sku,
        asin_value=asin,
        source_column=price_col_key,
        pim_upload_id=pim_upload_id,
        import_file_name=import_file_name,
        currency=cur_ins,
        row_cells=row_cells,
        handles=handles,
        price_pick_meta=price_pick_meta,
    )


def _pim_row_cells_from_series(row: Any, headers: list[str]) -> tuple[dict[str, Any], int]:
    """Build row dict with every cell trimmed (empty-after-trim -> None). Returns (cells, trim_count)."""
    out: dict[str, Any] = {}
    trimmed = 0
    for h in headers:
        try:
            if h in row.index:
                raw = _cell_to_json_safe(row[h])
            else:
                raw = None
        except Exception:
            raw = None
        nv, inc = trim_cell_value(raw)
        out[h] = nv
        trimmed += int(inc)
    return out, trimmed


# --- Google Sheets Sync helpers (used by etl_sync_google_sheets) ---
def _extract_google_sheet_id_from_module_configs(module_configs: Any) -> str | None:
    if not module_configs or not isinstance(module_configs, dict):
        return None
    cat = module_configs.get("catalog")
    if isinstance(cat, dict):
        sid = cat.get("google_sheet_id")
        if isinstance(sid, str) and sid.strip():
            return sid.strip()
    return None


def _get_google_sheet_id_from_workspace(db: Any, organization_id: str) -> str | None:
    """Resolve google_sheet_id from workspace_settings.module_configs.catalog (Next app shape), with fallbacks."""
    try:
        res = (
            db.table("workspace_settings")
            .select("module_configs")
            .eq("organization_id", organization_id)
            .limit(1)
            .execute()
        )
        if res.data:
            found = _extract_google_sheet_id_from_module_configs(res.data[0].get("module_configs"))
            if found:
                return found
    except Exception as e:
        log.debug("workspace_settings by organization_id: %s", e)
    try:
        res = db.table("workspace_settings").select("module_configs").limit(1).execute()
        if res.data:
            found = _extract_google_sheet_id_from_module_configs(res.data[0].get("module_configs"))
            if found:
                return found
    except Exception as e:
        log.debug("workspace_settings singleton: %s", e)
    try:
        res = (
            db.table("workspace_settings")
            .select("value")
            .eq("organization_id", organization_id)
            .eq("key", "google_sheet_id")
            .limit(1)
            .execute()
        )
        if res.data:
            v = res.data[0].get("value")
            if isinstance(v, str) and v.strip():
                return v.strip()
    except Exception as e:
        log.debug("workspace_settings legacy key/value: %s", e)
    return None


def _enrich_products_with_vendor_names(db: Any, rows: List[dict[str, Any]]) -> List[dict[str, Any]]:
    if not rows:
        return rows
    seen: set[str] = set()
    vendor_ids: list[str] = []
    for r in rows:
        vid = r.get("vendor_id")
        if vid is None:
            continue
        sid = str(vid)
        if sid not in seen:
            seen.add(sid)
            vendor_ids.append(sid)
    id_to_name: dict[str, str] = {}
    if vendor_ids:
        try:
            vr = db.table("vendors").select("id, name").in_("id", vendor_ids).execute()
            for v in vr.data or []:
                if v.get("id") is not None:
                    id_to_name[str(v["id"])] = str(v.get("name") or "").strip()
        except Exception as e:
            log.warning("Vendor name enrichment failed: %s", e)
    for r in rows:
        if r.get("vendor_name"):
            continue
        vid = r.get("vendor_id")
        r["vendor_name"] = id_to_name.get(str(vid)) if vid else None
    return rows


def _get_google_creds(db: Any, org_id: str):
    """Fetches Google Service Account JSON from organization_api_keys."""
    try:
        from google.oauth2 import service_account
    except ModuleNotFoundError as e:
        raise HTTPException(
            status_code=500,
            detail=(
                "Google Sheets dependencies are not installed. Run: "
                "python -m pip install google-api-python-client google-auth"
            ),
        ) from e

    try:
        res = (
            db.table("organization_api_keys")
            .select("api_key")
            .eq("organization_id", org_id)
            .eq("name", "google_sheets_api")
            .execute()
        )
        if res.data:
            creds_info = json.loads(res.data[0]["api_key"])
            return service_account.Credentials.from_service_account_info(creds_info)
    except HTTPException:
        raise
    except Exception as e:
        log.error("Failed to load Google credentials: %s", e)
    return None


@app.post("/etl/seed-products")
async def etl_seed_products(
    file: UploadFile = File(...),
    organization_id: str = Form(...),
    store_id: str = Form(...),
    mode: str = Form("preview"),
    confirm: str = Form("false"),
    seed_session_id: str | None = Form(None),
):
    """Catalog seed: default ``mode=preview`` (no writes). Use ``mode=apply`` + ``confirm=true`` or ``confirm=1`` after review."""
    organization_id, store_id = _validate_pim_org_store(organization_id, store_id)
    db = _require_supabase()
    try:
        raw_bytes = await file.read()
        if not raw_bytes:
            raise HTTPException(status_code=400, detail="Empty file.")

        openai_key = _get_api_credentials(db, organization_id, "openai_api_key")
        amazon_creds = _get_api_credentials(db, organization_id, "amazon_sp_api")

        fname = file.filename
        prepared = _pim_collect_prepared_from_upload(raw_bytes, fname, openai_key)
        quality = _pim_seed_scan_prepared_quality(
            db, organization_id, store_id, prepared, amazon_creds, skip_amazon=True
        )

        first = prepared[0] if prepared else {}
        column_mapping = first.get("pim_column_map") if isinstance(first, dict) else {}
        if not isinstance(column_mapping, dict):
            column_mapping = {}
        mapping_source = str(first.get("mapping_source") or "unknown")

        m = (mode or "preview").strip().lower()
        if m in ("preview", "dry_run", "simulate"):
            hist_id = _pim_seed_history_insert_preview(
                db, organization_id, store_id, fname, column_mapping, mapping_source, quality
            )
            pe = quality.get("preview_errors") or []
            err_first = pe[:200] if isinstance(pe, list) else []
            return {
                "status": "preview",
                "stage": "preview_ready",
                "mode": "preview",
                "message": "Quality scan only — no writes. Review mapping and quality, then POST mode=apply with confirm=true or confirm=1.",
                "mapping": column_mapping,
                "mapping_source": mapping_source,
                "quality": quality,
                "delimiter_detected": first.get("pim_delimiter_detected") if isinstance(first, dict) else None,
                "delimiter_uncertain": bool(first.get("pim_delimiter_uncertain")) if isinstance(first, dict) else False,
                "metrics": {
                    "rows_total": quality.get("rows_total"),
                    "dirty_rows": quality.get("dirty_rows"),
                    "products_would_create": quality.get("products_would_create"),
                    "products_would_update": quality.get("products_would_update"),
                },
                "accepted_sample": quality.get("accepted_sample") or [],
                "rejected_sample": quality.get("rejected_sample") or [],
                "errors": err_first,
                "seed_session_id": hist_id,
            }

        if m != "apply":
            raise HTTPException(
                status_code=400,
                detail={"error": "invalid_mode", "message": 'mode must be "preview" or "apply".', "got": mode},
            )

        if not _pim_catalog_seed_confirm_true(confirm):
            raise HTTPException(
                status_code=400,
                detail={
                    "error": "confirm_required",
                    "message": "Import writes are disabled until you pass confirm=true or confirm=1 (after reviewing preview).",
                    "quality": quality,
                    "mapping": column_mapping,
                },
            )

        if quality.get("apply_blocked_by_dirty_rate"):
            _pim_seed_history_finalize_apply(
                db,
                seed_session_id,
                organization_id,
                store_id,
                success=False,
                metrics=None,
                quality=quality,
                error_message="apply_blocked_by_dirty_rate",
            )
            raise HTTPException(
                status_code=422,
                detail={
                    "error": "dirty_rate_too_high",
                    "message": (
                        f"dirty_rate {quality.get('dirty_rate')} exceeds configured maximum "
                        f"{quality.get('max_dirty_rate')} — fix rows or raise env PIM_SEED_MAX_DIRTY_RATE."
                    ),
                    "quality": quality,
                    "mapping": column_mapping,
                },
            )

        try:
            metrics = _pim_seed_apply_prepared_writes(
                db, organization_id, store_id, prepared, amazon_creds, "etl_seed_products"
            )
        except Exception as apply_exc:
            _pim_seed_history_finalize_apply(
                db,
                seed_session_id,
                organization_id,
                store_id,
                success=False,
                metrics=None,
                quality=quality,
                error_message=str(apply_exc),
            )
            raise
        _pim_seed_history_finalize_apply(
            db,
            seed_session_id,
            organization_id,
            store_id,
            success=True,
            metrics=metrics,
            quality=quality,
            error_message=None,
        )
        pe = metrics.get("errors") or []
        err_first = pe[:200] if isinstance(pe, list) else []
        return {
            "status": "success",
            "stage": "complete",
            "mode": "apply",
            "message": "Catalog import finished.",
            "mapping": column_mapping,
            "mapping_source": mapping_source,
            "metrics": metrics,
            "quality": quality,
            "accepted_sample": quality.get("accepted_sample") or [],
            "rejected_sample": quality.get("rejected_sample") or [],
            "errors": err_first,
            "seed_session_id": seed_session_id,
        }
    except HTTPException:
        raise
    except Exception as e:
        log.exception("etl_seed_products failed")
        raise HTTPException(status_code=500, detail=str(e)) from e


@app.post("/etl/pim-import/retry-preview")
async def etl_pim_import_retry_preview_route(body: PimImportPreviewStepBody):
    """Reset persisted preview cursor and quality accumulators; reuse same upload row."""
    organization_id, store_id = _validate_pim_org_store(body.organization_id, body.store_id)
    uid = (body.upload_id or "").strip()
    if not uid:
        raise HTTPException(status_code=400, detail="upload_id is required.")
    db = _require_supabase()
    try:
        return run_pim_import_retry_preview(db, organization_id, store_id, uid)
    except HTTPException:
        raise
    except Exception as e:
        log.exception("pim-import retry-preview failed")
        raise HTTPException(status_code=500, detail=str(e)) from e


@app.post("/etl/pim-import/preview-status")
async def etl_pim_import_preview_status(body: PimImportPreviewStepBody):
    """Poll persisted preview result (same payload as preview-step when complete)."""
    organization_id, store_id = _validate_pim_org_store(body.organization_id, body.store_id)
    uid = (body.upload_id or "").strip()
    if not uid:
        raise HTTPException(status_code=400, detail="upload_id is required.")
    db = _require_supabase()
    try:
        return get_pim_import_preview_status(db, organization_id, store_id, uid)
    except HTTPException:
        raise
    except Exception as e:
        log.exception("pim-import preview-status failed")
        raise HTTPException(status_code=500, detail=str(e)) from e


@app.post("/etl/pim-import/preview-step")
async def etl_pim_import_preview_step(body: PimImportPreviewStepBody):
    """One bounded chunk of PIM Product Master preview (poll until done)."""
    organization_id, store_id = _validate_pim_org_store(body.organization_id, body.store_id)
    uid = (body.upload_id or "").strip()
    if not uid:
        raise HTTPException(status_code=400, detail="upload_id is required.")
    db = _require_supabase()
    openai_key = _get_api_credentials(db, organization_id, "openai_api_key")
    amazon_creds = _get_api_credentials(db, organization_id, "amazon_sp_api")
    try:
        return run_pim_import_preview_step(
            db,
            organization_id,
            store_id,
            uid,
            openai_key=openai_key,
            amazon_creds=amazon_creds,
            row_chunk=body.row_chunk,
            scan_data_row_hint=body.scan_data_row_hint,
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        err_detail = f"{type(e).__name__}: {e}"
        log.exception("pim-import preview-step failed: %s", err_detail)
        raise HTTPException(
            status_code=500,
            detail={"error": "preview_step_exception", "message": err_detail,
                    "traceback": traceback.format_exc()[-2000:]},
        ) from e


@app.post("/etl/pim-import/apply-step")
async def etl_pim_import_apply_step(body: PimImportApplyStepBody):
    """One bounded chunk of PIM Product Master apply (poll until done)."""
    organization_id, store_id = _validate_pim_org_store(body.organization_id, body.store_id)
    uid = (body.upload_id or "").strip()
    if not uid:
        raise HTTPException(status_code=400, detail="upload_id is required.")
    db = _require_supabase()
    res = (
        db.table("raw_report_uploads")
        .select("metadata")
        .eq("id", uid)
        .eq("organization_id", organization_id)
        .limit(1)
        .execute()
    )
    rows = res.data or []
    if not rows:
        raise HTTPException(status_code=404, detail="Upload not found.")
    meta = rows[0].get("metadata") if isinstance(rows[0].get("metadata"), dict) else {}
    job = meta.get("pim_import_job") if isinstance(meta.get("pim_import_job"), dict) else {}
    q = job.get("preview_quality") if isinstance(job.get("preview_quality"), dict) else {}

    # Canonical safe-rows-only flag — accepts both import_safe_rows_only and legacy skip_conflicts.
    # ALSO auto-reads from metadata so resume/retry inherits the original run's safe-only setting
    # even when the frontend doesn't re-send the flag (e.g. after ConnectionTerminated).
    safe_rows_only = body.safe_rows_only or bool(meta.get("pim_import_safe_rows_only"))
    log.info(
        "pim apply-step upload_id=%s import_safe_rows_only=%s skip_conflicts=%s "
        "safe_rows_only=%s apply_blocked_by_conflicts=%s",
        uid, body.import_safe_rows_only, body.skip_conflicts,
        safe_rows_only, q.get("apply_blocked_by_conflicts"),
    )

    if q.get("apply_blocked_by_dirty_rate"):
        raise HTTPException(
            status_code=422,
            detail={
                "error": "dirty_rate_too_high",
                "message": f"dirty_rate {q.get('dirty_rate')} exceeds max {q.get('max_dirty_rate')}",
            },
        )
    if q.get("apply_blocked_by_conflicts") and not safe_rows_only:
        raise HTTPException(
            status_code=422,
            detail={
                "error": "conflicts_block_import",
                "message": "Preview has blocking identifier conflicts. Use 'Import safe rows only' to import accepted rows and skip the conflicting rows.",
                "conflict_count": int(q.get("conflict_rows_blocked") or 0),
                "can_skip_conflicts": True,
                "import_safe_rows_only": True,
            },
        )
    confirm_true_in_body = _pim_catalog_seed_confirm_true(body.confirm)
    confirmed = confirm_true_in_body or bool(meta.get("pim_import_confirmed"))
    if not confirmed:
        raise HTTPException(
            status_code=400,
            detail={"error": "confirm_required", "message": "Pass confirm=true after reviewing preview."},
        )

    # Persist the safe-rows-only flag in metadata so the apply loop has visibility
    if safe_rows_only and not meta.get("pim_import_safe_rows_only"):
        try:
            import datetime as _dt
            meta2 = {**meta, "pim_import_safe_rows_only": True}
            db.table("raw_report_uploads").update(
                {"metadata": meta2, "updated_at": _dt.datetime.now(_dt.timezone.utc).isoformat()}
            ).eq("id", uid).eq("organization_id", organization_id).execute()
        except Exception as _persist_err:
            log.warning("could not persist pim_import_safe_rows_only flag: %s", _persist_err)

    amazon_creds = _get_api_credentials(db, organization_id, "amazon_sp_api")
    try:
        return run_pim_import_apply_step(
            db,
            organization_id,
            store_id,
            uid,
            amazon_creds=amazon_creds,
            price_source="pim_import_async",
            row_chunk=body.row_chunk,
            confirmed_via_api=confirm_true_in_body,
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback as _tb
        err_detail = f"{type(e).__name__}: {e}"
        log.exception("pim-import apply-step failed: %s", err_detail)
        raise HTTPException(
            status_code=500,
            detail={"error": "apply_step_exception", "message": err_detail,
                    "traceback": _tb.format_exc()[-2000:]},
        ) from e


@app.post("/etl/pim-import/backfill-prices-step")
async def etl_pim_import_backfill_prices_step_route(body: PimImportPriceBackfillStepBody):
    """One bounded chunk of Product Master price backfill; checkpoints in file_processing_status.import_metrics."""
    organization_id = body.organization_id.strip()
    upload_id = body.upload_id.strip()
    if not organization_id or not upload_id:
        raise HTTPException(status_code=400, detail="organization_id and upload_id are required.")
    db = _require_supabase()
    try:
        return run_pim_import_price_backfill_step(
            db,
            organization_id,
            upload_id,
            row_chunk=body.row_chunk,
            restart=bool(body.restart),
            cancel=bool(body.cancel),
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback as _tb

        err_detail = f"{type(e).__name__}: {e}"
        log.exception("pim-import backfill-prices-step failed: %s", err_detail)
        raise HTTPException(
            status_code=500,
            detail={"error": "backfill_prices_step_exception", "message": err_detail, "traceback": _tb.format_exc()[-2000:]},
        ) from e


@app.post("/etl/sync-google-sheets")
async def etl_sync_google_sheets(
    organization_id: str = Form(...),
    store_id: str = Form(...),
    mode: str = Form("preview"),
    confirm: str = Form("false"),
    seed_session_id: str | None = Form(None),
):
    """Google Sheets catalog sync: default ``mode=preview`` (no writes). ``mode=apply`` + ``confirm=true`` or ``confirm=1`` to write."""
    organization_id, store_id = _validate_pim_org_store(organization_id, store_id)
    db = _require_supabase()

    sheet_id = _get_google_sheet_id_from_workspace(db, organization_id)
    if not sheet_id:
        raise HTTPException(
            status_code=400,
            detail=(
                "Google Sheet ID not found. Configure catalog.google_sheet_id in workspace module_configs "
                "or legacy workspace_settings key google_sheet_id."
            ),
        )

    creds = _get_google_creds(db, organization_id)
    if not creds:
        raise HTTPException(status_code=401, detail="Google API Key not configured in Organization Keys.")

    try:
        openai_key = _get_api_credentials(db, organization_id, "openai_api_key")
        amazon_creds = _get_api_credentials(db, organization_id, "amazon_sp_api")
        prepared = _pim_google_build_prepared_frames(db, organization_id, sheet_id, creds, openai_key)
        quality = _pim_seed_scan_prepared_quality(
            db, organization_id, store_id, prepared, amazon_creds, skip_amazon=True
        )

        first = prepared[0] if prepared else {}
        column_mapping = first.get("pim_column_map") if isinstance(first, dict) else {}
        if not isinstance(column_mapping, dict):
            column_mapping = {}
        mapping_source = str(first.get("mapping_source") or "unknown")
        virtual_fname = f"google_sheet_{sheet_id}.xlsx"

        m = (mode or "preview").strip().lower()
        if m in ("preview", "dry_run", "simulate"):
            hist_id = _pim_seed_history_insert_preview(
                db, organization_id, store_id, virtual_fname, column_mapping, mapping_source, quality
            )
            pe = quality.get("preview_errors") or []
            err_first = pe[:200] if isinstance(pe, list) else []
            return {
                "status": "preview",
                "stage": "preview_ready",
                "mode": "preview",
                "message": "Quality scan only — no writes. Review mapping and quality, then POST mode=apply with confirm=true or confirm=1.",
                "mapping": column_mapping,
                "mapping_source": mapping_source,
                "quality": quality,
                "delimiter_detected": None,
                "delimiter_uncertain": False,
                "metrics": {
                    "rows_total": quality.get("rows_total"),
                    "dirty_rows": quality.get("dirty_rows"),
                    "products_would_create": quality.get("products_would_create"),
                    "products_would_update": quality.get("products_would_update"),
                },
                "accepted_sample": quality.get("accepted_sample") or [],
                "rejected_sample": quality.get("rejected_sample") or [],
                "errors": err_first,
                "seed_session_id": hist_id,
            }

        if m != "apply":
            raise HTTPException(
                status_code=400,
                detail={"error": "invalid_mode", "message": 'mode must be "preview" or "apply".', "got": mode},
            )

        if not _pim_catalog_seed_confirm_true(confirm):
            raise HTTPException(
                status_code=400,
                detail={
                    "error": "confirm_required",
                    "message": "Sheet sync writes are disabled until you pass confirm=true or confirm=1 (after reviewing preview).",
                    "quality": quality,
                    "mapping": column_mapping,
                },
            )

        if quality.get("apply_blocked_by_dirty_rate"):
            _pim_seed_history_finalize_apply(
                db,
                seed_session_id,
                organization_id,
                store_id,
                success=False,
                metrics=None,
                quality=quality,
                error_message="apply_blocked_by_dirty_rate",
            )
            raise HTTPException(
                status_code=422,
                detail={
                    "error": "dirty_rate_too_high",
                    "message": (
                        f"dirty_rate {quality.get('dirty_rate')} exceeds configured maximum "
                        f"{quality.get('max_dirty_rate')} — fix sheet rows or raise env PIM_SEED_MAX_DIRTY_RATE."
                    ),
                    "quality": quality,
                    "mapping": column_mapping,
                },
            )

        try:
            metrics = _pim_seed_apply_prepared_writes(
                db, organization_id, store_id, prepared, amazon_creds, "etl_google_sheets"
            )
        except Exception as apply_exc:
            _pim_seed_history_finalize_apply(
                db,
                seed_session_id,
                organization_id,
                store_id,
                success=False,
                metrics=None,
                quality=quality,
                error_message=str(apply_exc),
            )
            raise
        _pim_seed_history_finalize_apply(
            db,
            seed_session_id,
            organization_id,
            store_id,
            success=True,
            metrics=metrics,
            quality=quality,
            error_message=None,
        )
        pe = metrics.get("errors") or []
        err_first = pe[:200] if isinstance(pe, list) else []
        return {
            "status": "success",
            "stage": "complete",
            "mode": "apply",
            "message": f"Google Sheets sync finished ({metrics['sheets_processed']} non-empty tabs).",
            "mapping": column_mapping,
            "mapping_source": mapping_source,
            "metrics": metrics,
            "quality": quality,
            "accepted_sample": quality.get("accepted_sample") or [],
            "rejected_sample": quality.get("rejected_sample") or [],
            "errors": err_first,
            "seed_session_id": seed_session_id,
        }
    except HTTPException:
        raise
    except Exception as e:
        log.exception("etl_sync_google_sheets failed")
        raise HTTPException(status_code=500, detail=f"Google Sync Failed: {e}") from e


# --- 2. Real Product List Endpoint ---
@app.get("/etl/products/{organization_id}")
async def get_real_products(organization_id: str):
    db = _require_supabase()
    # Fetch from products table with their current prices and identifiers
    res = db.table("products").select("*, product_identifier_map(*), product_prices(*)").eq("organization_id", organization_id).execute()
    rows = list(res.data or [])
    rows = _enrich_products_with_vendor_names(db, rows)
    return {"status": "success", "products": rows}