"""Canonical mapping plan from product dimensions xlsx; prints JSON to stdout."""
import json
import re
import sys
from collections import defaultdict

import openpyxl

xlsx = sys.argv[1]
intake_run_id = sys.argv[2] if len(sys.argv) > 2 else "20260527T210000Z"


def norm_header(h):
    if h is None:
        return ""
    return str(h).replace("\r\n", "\n").strip()


def norm_val(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def asin_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^B0[A-Z0-9]{8}$", s))


def fnsku_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^X0[A-Z0-9]{8,}$", s))


def parse_dims(s):
    if not s:
        return None
    t = str(s).strip()
    unit = "inch"
    if re.search(r"\bCM\b", t, re.I):
        unit = "unknown"
    nums = []
    parts = re.split(r"\s*[xX×]\s*", t)
    alt = []
    for p in parts:
        m = re.search(r"([\d.]+)", p)
        if m:
            alt.append(float(m.group(1)))
    if len(alt) >= 3:
        nums = alt[:3]
    if len(nums) >= 3:
        return {"l": nums[0], "w": nums[1], "h": nums[2], "unit": unit, "raw": t}
    if len(nums) > 0:
        return {"partial": nums, "unit": unit, "raw": t}
    return {"unit": "unknown", "raw": t}


def map_fba_fbm(val):
    if not val:
        return "unknown"
    t = str(val).strip().upper()
    if "FBA" in t and "FBM" not in t:
        return "fba"
    if "FBM" in t or "MFN" in t:
        return "mfn"
    if "FBA" in t and "FBM" in t:
        return "unknown"
    return "unknown"


wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
wb.close()

headers = [norm_header(h) for h in rows[0]]
data_rows = []
for i, row in enumerate(rows[1:], start=2):
    if not any(row):
        continue
    rec = {headers[j]: norm_val(row[j]) if j < len(row) else None for j in range(len(headers))}
    rec["_row"] = i
    data_rows.append(rec)

dim_col = next((h for h in headers if "dimension" in h.lower()), None)
wt_col = next((h for h in headers if "wt" in h.lower() or "weight" in h.lower()), None)
asin_col = next((h for h in headers if "asin" in h.lower()), None)
fnsku_col = next((h for h in headers if "fnsku" in h.lower()), None)
sku_col = next((h for h in headers if "seller-sku" in h.lower() or h.lower() == "sku"), None)
unit_upc_col = next((h for h in headers if h and "unit upc" in h.lower()), None)
case_upc_col = next((h for h in headers if h and "case upc" in h.lower()), None)
mfg_col = next((h for h in headers if "mfg" in h.lower()), None)
brand_col = next((h for h in headers if h and h.lower() == "brand"), None)
fba_col = next((h for h in headers if "fba" in h.lower()), None)
desc_col = next((h for h in headers if h and h.lower() == "description"), None)
sell_pack_col = next((h for h in headers if "selling pack" in h.lower()), None)
case_pack_col = next((h for h in headers if "case pack" in h.lower()), None)
ti_col = next((h for h in headers if h and h.lower() == "ti"), None)
hi_col = next((h for h in headers if h and h.lower() == "hi"), None)
pallet_col = next((h for h in headers if "pallet" in h.lower()), None)

dup_asin = defaultdict(list)
dup_sku = defaultdict(list)
asin_dims = defaultdict(set)

for rec in data_rows:
    row = rec["_row"]
    asin = str(rec.get(asin_col) or "").strip().upper() if asin_col else ""
    sku = str(rec.get(sku_col) or "").strip() if sku_col else ""
    if asin:
        dup_asin[asin].append(row)
    if sku:
        dup_sku[sku].append(row)
    parsed = parse_dims(rec.get(dim_col) if dim_col else None)
    if asin_ok(asin) and parsed and parsed.get("l"):
        key = (round(parsed["l"], 4), round(parsed["w"], 4), round(parsed["h"], 4))
        asin_dims[asin].add(key)

dup_asin_multi = {k for k, v in dup_asin.items() if len(v) > 1}
dup_sku_multi = {k for k, v in dup_sku.items() if len(v) > 1}
conflict_asins = [a for a, keys in asin_dims.items() if len(keys) > 1]

queue = []
counts = defaultdict(int)

for rec in data_rows:
    row = rec["_row"]
    asin = str(rec.get(asin_col) or "").strip().upper() if asin_col else ""
    sku = str(rec.get(sku_col) or "").strip() if sku_col else ""
    fnsku = str(rec.get(fnsku_col) or "").strip().upper() if fnsku_col else ""
    parsed = parse_dims(rec.get(dim_col) if dim_col else None)
    has_lwh = bool(parsed and parsed.get("l"))
    fulfillment = map_fba_fbm(rec.get(fba_col) if fba_col else None)

    reasons = []
    merge_class = "merge-safe"

    if not has_lwh:
        reasons.append("missing_case_dimensions")
        merge_class = "review-required"
    if asin in conflict_asins:
        reasons.append("asin_conflicting_dimensions")
        merge_class = "conflict"
    elif asin in dup_asin_multi:
        reasons.append("duplicate_asin_variant_rows")
        if merge_class != "conflict":
            merge_class = "review-required"
    if sku in dup_sku_multi:
        reasons.append("duplicate_seller_sku")
        if merge_class == "merge-safe":
            merge_class = "review-required"
    if not fnsku_ok(fnsku):
        reasons.append("fnsku_absent_or_invalid")
    if has_lwh and asin_ok(asin) and sku and merge_class == "merge-safe":
        reasons.append("ready_for_staging_match_census")

    counts[merge_class] += 1
    queue.append(
        {
            "row": row,
            "seller_sku": sku,
            "asin": asin,
            "fnsku": fnsku if fnsku_ok(fnsku) else None,
            "fulfillment_hint": fulfillment,
            "has_parseable_lwh": has_lwh,
            "merge_class": merge_class,
            "reasons": reasons,
        }
    )

# Summarize queue tiers for markdown (cap samples)
review_rows = [q for q in queue if q["merge_class"] in ("review-required", "conflict")]
importable_rows = [q for q in queue if q["has_parseable_lwh"] and q["merge_class"] == "merge-safe"]

summary = {
    "intake_run_id": intake_run_id,
    "total_rows": len(data_rows),
    "merge_safe": counts["merge-safe"],
    "review_required": counts["review-required"],
    "conflict": counts["conflict"],
    "parseable_lwh": sum(1 for q in queue if q["has_parseable_lwh"]),
    "duplicate_asin_keys": len(dup_asin_multi),
    "conflicting_asin_keys": len(conflict_asins),
}

print(
    json.dumps(
        {
            "summary": summary,
            "headers": headers,
            "column_keys": {
                "asin": asin_col,
                "fnsku": fnsku_col,
                "sku": sku_col,
                "unit_upc": unit_upc_col,
                "case_upc": case_upc_col,
                "mfg": mfg_col,
                "brand": brand_col,
                "dimensions": dim_col,
                "weight": wt_col,
                "fba_fbm": fba_col,
                "description": desc_col,
                "selling_pack": sell_pack_col,
                "case_pack": case_pack_col,
                "ti": ti_col,
                "hi": hi_col,
                "cases_per_pallet": pallet_col,
            },
            "queue_sample_review": review_rows[:40],
            "queue_sample_importable": importable_rows[:25],
            "review_required_count": len(review_rows),
            "importable_merge_safe_count": len(importable_rows),
        }
    )
)
