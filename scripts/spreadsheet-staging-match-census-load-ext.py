"""Extended spreadsheet loader for staging match census dryrun."""
import json
import re
import sys

import openpyxl

xlsx = sys.argv[1]


def norm_header(h):
    if h is None:
        return ""
    return str(h).replace("\r\n", "\n").strip()


def norm_val(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def asin_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^B0[A-Z0-9]{8}$", s))


def fnsku_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^X0[A-Z0-9]{8,}$", s))


def parse_dims(s):
    if not s:
        return None
    t = str(s).strip()
    unit = "in"
    if re.search(r"\bCM\b", t, re.I):
        unit = "cm"
    nums = []
    parts = re.split(r"\s*[xX×]\s*", t)
    alt = []
    for p in parts:
        m = re.search(r"([\d.]+)", p)
        if m:
            alt.append(float(m.group(1)))
    if len(alt) >= 3:
        nums = alt[:3]
    if len(nums) >= 3:
        return {
            "length": round(nums[0], 4),
            "width": round(nums[1], 4),
            "height": round(nums[2], 4),
            "dimension_unit": unit,
            "raw": t,
        }
    return None


def parse_weight(v):
    if v is None:
        return None
    try:
        n = float(str(v).replace(",", ""))
        if n > 0:
            return {"weight": round(n, 4), "weight_unit": "lb"}
    except ValueError:
        pass
    return None


def parse_int(v):
    if v is None:
        return None
    try:
        n = int(float(str(v).replace(",", "")))
        return n if n > 0 else None
    except ValueError:
        return None


def map_fba_fbm(val):
    if not val:
        return "unknown"
    t = str(val).strip().upper()
    if "FBM" in t or "MFN" in t:
        return "mfn"
    if "FBA" in t:
        return "fba"
    return "unknown"


def brand_starts_1883(brand):
    if not brand:
        return False
    return str(brand).strip().startswith("1883")


wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
wb.close()

headers = [norm_header(h) for h in rows[0]]
dim_col = next((h for h in headers if "dimension" in h.lower()), None)
wt_col = next((h for h in headers if "wt" in h.lower() or "weight" in h.lower()), None)
asin_col = next((h for h in headers if "asin" in h.lower()), None)
fnsku_col = next((h for h in headers if "fnsku" in h.lower()), None)
sku_col = next((h for h in headers if "seller-sku" in h.lower() or h.lower() == "sku"), None)
brand_col = next((h for h in headers if h and h.lower() == "brand"), None)
fba_col = next((h for h in headers if "fba" in h.lower()), None)
case_pack_col = next((h for h in headers if "case pack" in h.lower()), None)
sell_pack_col = next((h for h in headers if "selling pack" in h.lower()), None)
ti_col = next((h for h in headers if h and h.strip().lower() in ("ti", "t.i.", "ti ")), None)
if not ti_col:
    ti_col = next((h for h in headers if h and re.search(r"\bti\b", h.lower())), None)
hi_col = next((h for h in headers if h and h.strip().lower() in ("hi", "h.i.", "hi ")), None)
if not hi_col:
    hi_col = next((h for h in headers if h and re.search(r"\bhi\b", h.lower())), None)
cpp_col = next(
    (
        h
        for h in headers
        if h and ("cases per pallet" in h.lower() or "case per pallet" in h.lower() or "cases/pallet" in h.lower())
    ),
    None,
)

dup_asin = {}
for i, row in enumerate(rows[1:], start=2):
    if not any(row):
        continue
    rec = {headers[j]: norm_val(row[j]) if j < len(row) else None for j in range(len(headers))}
    asin = str(rec.get(asin_col) or "").strip().upper() if asin_col else ""
    if asin:
        dup_asin.setdefault(asin, []).append(i)
dup_asin_multi = {k for k, v in dup_asin.items() if len(v) > 1}

out_rows = []
for i, row in enumerate(rows[1:], start=2):
    if not any(row):
        continue
    rec = {headers[j]: norm_val(row[j]) if j < len(row) else None for j in range(len(headers))}
    asin = str(rec.get(asin_col) or "").strip().upper() if asin_col else ""
    sku = str(rec.get(sku_col) or "").strip() if sku_col else ""
    fnsku = str(rec.get(fnsku_col) or "").strip().upper() if fnsku_col else ""
    dims = parse_dims(rec.get(dim_col) if dim_col else None)
    wt = parse_weight(rec.get(wt_col) if wt_col else None)
    brand = rec.get(brand_col) if brand_col else None
    fulfillment = map_fba_fbm(rec.get(fba_col) if fba_col else None)
    has_lwh = dims is not None
    merge_class_sheet = "merge-safe"
    if not has_lwh:
        merge_class_sheet = "review-required"
    elif asin in dup_asin_multi:
        merge_class_sheet = "review-required"

    out_rows.append(
        {
            "row": i,
            "seller_sku": sku,
            "asin": asin if asin_ok(asin) else None,
            "fnsku": fnsku if fnsku_ok(fnsku) else None,
            "brand": brand,
            "brand_starts_1883": brand_starts_1883(brand),
            "fulfillment_context": fulfillment,
            "packaging_level": "case",
            "dimensions": dims,
            "weight": wt,
            "case_pack": parse_int(rec.get(case_pack_col) if case_pack_col else None),
            "selling_pack": parse_int(rec.get(sell_pack_col) if sell_pack_col else None),
            "ti": parse_int(rec.get(ti_col) if ti_col else None),
            "hi": parse_int(rec.get(hi_col) if hi_col else None),
            "cases_per_pallet": parse_int(rec.get(cpp_col) if cpp_col else None),
            "sheet_merge_class": merge_class_sheet,
            "has_parseable_lwh": has_lwh,
        }
    )

merge_safe = [r for r in out_rows if r["sheet_merge_class"] == "merge-safe"]

print(
    json.dumps(
        {
            "headers": headers,
            "total_rows": len(out_rows),
            "merge_safe_count": len(merge_safe),
            "parseable_lwh_count": sum(1 for r in out_rows if r["has_parseable_lwh"]),
            "duplicate_asin_keys": len(dup_asin_multi),
            "rows": out_rows,
            "merge_safe_rows": merge_safe,
        }
    )
)
