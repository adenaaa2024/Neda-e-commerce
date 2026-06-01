"""Load spreadsheet rows with full identifier columns for missing-product link plan."""
import json
import re
import sys

import openpyxl

xlsx = sys.argv[1]
missing_rows_arg = sys.argv[2] if len(sys.argv) > 2 else ""


def norm_header(h):
    if h is None:
        return ""
    return str(h).replace("\r\n", "\n").strip()


def norm_val(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def asin_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^B0[A-Z0-9]{8}$", s))


def fnsku_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^X0[A-Z0-9]{8,}$", s))


def upc_ok(v):
    if not v:
        return False
    s = re.sub(r"[^0-9]", "", str(v).strip())
    return 8 <= len(s) <= 14


def norm_upc(v):
    if not v:
        return None
    s = re.sub(r"[^0-9]", "", str(v).strip())
    return s if upc_ok(v) else None


missing_rows = set()
if missing_rows_arg:
    if missing_rows_arg.startswith("@"):
        with open(missing_rows_arg[1:], encoding="utf-8") as f:
            for line in f:
                part = line.strip()
                if part.isdigit():
                    missing_rows.add(int(part))
    else:
        for part in missing_rows_arg.split(","):
            part = part.strip()
            if part.isdigit():
                missing_rows.add(int(part))


wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
wb.close()

headers = [norm_header(h) for h in rows[0]]
asin_col = next((h for h in headers if "asin" in h.lower()), None)
fnsku_col = next((h for h in headers if "fnsku" in h.lower()), None)
sku_col = next((h for h in headers if "seller-sku" in h.lower() or h.lower() == "sku"), None)
brand_col = next((h for h in headers if h and h.lower() == "brand"), None)
unit_upc_col = next((h for h in headers if h and "unit upc" in h.lower()), None)
case_upc_col = next((h for h in headers if h and "case upc" in h.lower()), None)
mfg_col = next((h for h in headers if h and "mfg" in h.lower()), None)
fba_col = next((h for h in headers if "fba" in h.lower()), None)

out_rows = []
for i, row in enumerate(rows[1:], start=2):
    if not any(row):
        continue
    if missing_rows and i not in missing_rows:
        continue
    rec = {headers[j]: norm_val(row[j]) if j < len(row) else None for j in range(len(headers))}
    asin = str(rec.get(asin_col) or "").strip().upper() if asin_col else ""
    fnsku = str(rec.get(fnsku_col) or "").strip().upper() if fnsku_col else ""
    sku = str(rec.get(sku_col) or "").strip() if sku_col else ""
    unit_upc = norm_upc(rec.get(unit_upc_col) if unit_upc_col else None)
    case_upc = norm_upc(rec.get(case_upc_col) if case_upc_col else None)
    mfg = str(rec.get(mfg_col) or "").strip() if mfg_col else ""
    brand = rec.get(brand_col) if brand_col else None
    fba = str(rec.get(fba_col) or "").strip().upper() if fba_col else ""

    out_rows.append(
        {
            "row": i,
            "seller_sku": sku,
            "asin": asin if asin_ok(asin) else None,
            "fnsku": fnsku if fnsku_ok(fnsku) else None,
            "unit_upc": unit_upc,
            "case_upc": case_upc,
            "mfg_number": mfg or None,
            "brand": brand,
            "fulfillment_context": "mfn" if "FBM" in fba or "MFN" in fba else ("fba" if "FBA" in fba else "unknown"),
        }
    )

print(json.dumps({"headers": headers, "rows": out_rows, "count": len(out_rows)}))
