"""Load spreadsheet rows for staging match census; JSON to stdout."""
import json
import re
import sys
from collections import defaultdict

import openpyxl

xlsx = sys.argv[1]


def norm_header(h):
    if h is None:
        return ""
    return str(h).replace("\r\n", "\n").strip()


def norm_val(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def asin_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^B0[A-Z0-9]{8}$", s))


def fnsku_ok(v):
    if not v:
        return False
    s = str(v).strip().upper()
    return bool(re.match(r"^X0[A-Z0-9]{8,}$", s))


def parse_dims(s):
    if not s:
        return None
    t = str(s).strip()
    unit = "in"
    if re.search(r"\bCM\b", t, re.I):
        unit = "cm"
    nums = []
    parts = re.split(r"\s*[xX×]\s*", t)
    alt = []
    for p in parts:
        m = re.search(r"([\d.]+)", p)
        if m:
            alt.append(float(m.group(1)))
    if len(alt) >= 3:
        nums = alt[:3]
    if len(nums) >= 3:
        return {
            "length": round(nums[0], 4),
            "width": round(nums[1], 4),
            "height": round(nums[2], 4),
            "dimension_unit": unit,
            "raw": t,
        }
    return None


def parse_weight(v):
    if v is None:
        return None
    try:
        n = float(str(v).replace(",", ""))
        if n > 0:
            return {"weight": round(n, 4), "weight_unit": "lb"}
    except ValueError:
        pass
    return None


def map_fba_fbm(val):
    if not val:
        return "unknown"
    t = str(val).strip().upper()
    if "FBM" in t or "MFN" in t:
        return "mfn"
    if "FBA" in t:
        return "fba"
    return "unknown"


def brand_starts_1883(brand):
    if not brand:
        return False
    return str(brand).strip().startswith("1883")


wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
wb.close()

headers = [norm_header(h) for h in rows[0]]
dim_col = next((h for h in headers if "dimension" in h.lower()), None)
wt_col = next((h for h in headers if "wt" in h.lower() or "weight" in h.lower()), None)
asin_col = next((h for h in headers if "asin" in h.lower()), None)
fnsku_col = next((h for h in headers if "fnsku" in h.lower()), None)
sku_col = next((h for h in headers if "seller-sku" in h.lower() or h.lower() == "sku"), None)
brand_col = next((h for h in headers if h and h.lower() == "brand"), None)
fba_col = next((h for h in headers if "fba" in h.lower()), None)
case_pack_col = next((h for h in headers if "case pack" in h.lower()), None)
sell_pack_col = next((h for h in headers if "selling pack" in h.lower()), None)

dup_asin = defaultdict(list)
for i, row in enumerate(rows[1:], start=2):
    if not any(row):
        continue
    rec = {headers[j]: norm_val(row[j]) if j < len(row) else None for j in range(len(headers))}
    asin = str(rec.get(asin_col) or "").strip().upper() if asin_col else ""
    if asin:
        dup_asin[asin].append(i)

dup_asin_multi = {k for k, v in dup_asin.items() if len(v) > 1}
asin_dims = defaultdict(set)
out_rows = []

for i, row in enumerate(rows[1:], start=2):
    if not any(row):
        continue
    rec = {headers[j]: norm_val(row[j]) if j < len(row) else None for j in range(len(headers))}
    asin = str(rec.get(asin_col) or "").strip().upper() if asin_col else ""
    sku = str(rec.get(sku_col) or "").strip() if sku_col else ""
    fnsku = str(rec.get(fnsku_col) or "").strip().upper() if fnsku_col else ""
    dims = parse_dims(rec.get(dim_col) if dim_col else None)
    wt = parse_weight(rec.get(wt_col) if wt_col else None)
    brand = rec.get(brand_col) if brand_col else None
    fulfillment = map_fba_fbm(rec.get(fba_col) if fba_col else None)

    has_lwh = dims is not None
    merge_class_sheet = "merge-safe"
    if not has_lwh:
        merge_class_sheet = "review-required"
    elif asin in dup_asin_multi:
        merge_class_sheet = "review-required"
    if asin_ok(asin) and dims:
        key = (dims["length"], dims["width"], dims["height"])
        asin_dims[asin].add(key)

    case_pack = None
    if case_pack_col and rec.get(case_pack_col) is not None:
        try:
            case_pack = int(float(str(rec.get(case_pack_col))))
        except ValueError:
            case_pack = None

    selling_pack = None
    if sell_pack_col and rec.get(sell_pack_col) is not None:
        try:
            selling_pack = int(float(str(rec.get(sell_pack_col))))
        except ValueError:
            selling_pack = None

    out_rows.append(
        {
            "row": i,
            "seller_sku": sku,
            "asin": asin if asin_ok(asin) else None,
            "fnsku": fnsku if fnsku_ok(fnsku) else None,
            "brand": brand,
            "brand_starts_1883": brand_starts_1883(brand),
            "fulfillment_context": fulfillment,
            "packaging_level": "case",
            "dimensions": dims,
            "weight": wt,
            "case_pack": case_pack,
            "selling_pack": selling_pack,
            "sheet_merge_class": merge_class_sheet,
            "has_parseable_lwh": has_lwh,
        }
    )

conflict_asins = [a for a, keys in asin_dims.items() if len(keys) > 1]
merge_safe = [r for r in out_rows if r["sheet_merge_class"] == "merge-safe"]

print(
    json.dumps(
        {
            "headers": headers,
            "total_rows": len(out_rows),
            "merge_safe_count": len(merge_safe),
            "parseable_lwh_count": sum(1 for r in out_rows if r["has_parseable_lwh"]),
            "duplicate_asin_keys": len(dup_asin_multi),
            "conflicting_asin_keys": len(conflict_asins),
            "rows": out_rows,
            "merge_safe_rows": merge_safe,
        }
    )
)
